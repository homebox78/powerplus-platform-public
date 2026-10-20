# -*- coding: utf-8 -*-
"""올린 자료에서 '실제 데이터'를 뽑는다 — 표·시계열(datasets)과 낱개 수치(facts) (2026-10-18 와이어 고도화 B).

왜: 예전엔 자료를 글로만 읽어(엑셀 빈 칸이 빠져 열이 밀리고, hwpx 표 칸이 한 줄로 풀리고) 요약(digest)에 수치 몇 개만 남았다.
그래서 와이어는 기획 글 속 수치만 보고 그래프를 짰고, 자료에 있는 시계열·구성비는 그림으로 살아나지 못했다.

흐름: 자료 파일 → tables_of(표를 칸 그대로) → datasets_from_table(숫자 열·라벨 열·유형 짐작) + prose_facts(글 속 수치 문장)
     → facts.json({sig, datasets, facts}) → 기획(catalog_text: 장마다 data 에 id) → 와이어(slide_data_text: 그 장의 표·수치)
AI 없이 정해지는 부분만 여기 둔다. 글 속 수치에 이름·시점을 붙이는 AI 단계(label_prompt / apply_labels)는 값을 새로 만들지 못한다 —
후보(F 번호)에 이름표만 붙이고, 값은 늘 원문에서 뽑은 그대로다.
"""
import json, os, re, zipfile

MAX_ROWS = 200        # 표 하나에서 읽는 줄 상한(엑셀 수만 줄 — 기획이 다 쓸 수 없다)
MAX_COLS = 30
MAX_DATASETS = 40
MAX_FACTS = 120

# ── 숫자 칸 읽기 ──
_NUM_CELL = re.compile(r"^\s*([-+−]?)\s*([\d,]*\.?\d+)\s*(%|퍼센트|건|명|분|초|시간|일|개월|년|원|천원|만원|억원|억|조|점|개|곳|회|배|대|종|인|석|㎡|km)?\s*$")
UNIT_HEAD = re.compile(r"[(\[（]\s*([^)\]）]{1,8})\s*[)\]）]\s*$")   # '처리시간(분)' → 분


def num(s):
    """'1,820' → (1820.0, '') · '71.4%' → (71.4, '%') · 숫자가 아니면 (None, '')"""
    if isinstance(s, bool) or s is None:
        return None, ""
    if isinstance(s, (int, float)):
        return float(s), ""
    m = _NUM_CELL.match(str(s))
    if not m or not re.search(r"\d", m.group(2)):
        return None, ""
    raw = m.group(2).replace(",", "")
    try:
        v = float(raw)
    except ValueError:
        return None, ""
    return (-v if m.group(1) in ("-", "−") else v), (m.group(3) or "")


def fmt(v):
    """수치를 글로 — 정수면 소수점 없이, 아니면 원래 자릿수(최대 2)"""
    if v is None:
        return ""
    if abs(v - round(v)) < 1e-9:
        return f"{int(round(v)):,}"
    return f"{v:,.2f}".rstrip("0").rstrip(".")


def cell_str(c):
    """엑셀 칸 값 → 글. 날짜는 '2024-03', 정수 같은 실수는 정수로, 빈 칸은 ''"""
    import datetime
    if c is None:
        return ""
    if isinstance(c, datetime.datetime):
        return c.strftime("%Y-%m") if c.day == 1 and not (c.hour or c.minute) else c.strftime("%Y-%m-%d")
    if isinstance(c, datetime.date):
        return c.strftime("%Y-%m-%d")
    if isinstance(c, float) and abs(c - round(c)) < 1e-9:
        return str(int(round(c)))
    return str(c).replace("\n", " ").strip()


# ── 표 뽑기 — 파일 종류마다. 표 = {"file", "where", "title", "rows": [[칸 글]]} ──
def _pad(rows):
    rows = [[str(c or "").replace("\n", " ").replace("|", "/").strip() for c in r[:MAX_COLS]] for r in rows]
    w = max((len(r) for r in rows), default=0)
    return [r + [""] * (w - len(r)) for r in rows]


def xlsx_tables(path, max_rows=MAX_ROWS + 50):
    """시트마다 표 하나(빈 줄로 나뉜 여러 표는 datasets_from_table 이 가른다). 빈 칸도 자리를 지킨다 — 예전엔 빠져서 열이 밀렸다.
    max_rows — 데이터 추출은 250줄, 글(doc_text)은 더 넉넉히(맥락 자료 분량은 load_context 가 자료별로 고르게 자른다)"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([cell_str(c) for c in row])
            if len(rows) >= max_rows:
                break
        while rows and not any(rows[-1]):
            rows.pop()
        rows = _pad(rows)
        if rows:   # 오른쪽 끝 빈 열 떼기
            last = max((j for r in rows for j, c in enumerate(r) if c), default=-1)
            rows = [r[:last + 1] for r in rows]
        if any(any(r) for r in rows):
            out.append({"file": os.path.basename(path), "where": ws.title, "title": ws.title, "rows": rows})
    return out


def _docx_table_rows(t):
    rows = []
    for row in t.rows:
        try:
            rows.append([c.text for c in row.cells])
        except Exception:   # 깨진 병합 칸
            rows.append([c.text for c in row._tr.tc_lst] if hasattr(row, "_tr") else [])
    return _pad(rows)


def docx_blocks(path):
    """본문 순서 그대로 [('p', 글) | ('table', 줄들)] — 표가 글 사이 제자리에 온다"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(el, d).text
            if t.strip():
                out.append(("p", t))
        elif tag == "tbl":
            rows = _docx_table_rows(Table(el, d))
            if rows:
                out.append(("table", rows))
    return out


def _ln(tag):
    return tag.rsplit("}", 1)[-1]


def _hx_text(e):
    """hwpx 요소 안의 글(hp:t) 모두 — 표 안 표도 글로 풀어서"""
    return " ".join("".join(t.itertext()) for t in e.iter() if _ln(t.tag) == "t").strip()


def _hx_table(tbl):
    """hp:tbl → 줄들. 칸 주소(cellAddr colAddr·rowAddr)·병합(cellSpan)으로 격자에 놓아 열이 밀리지 않게"""
    cells, seq_r = [], 0
    for tr in tbl:
        if _ln(tr.tag) != "tr":
            continue
        seq_c = 0
        for tc in tr:
            if _ln(tc.tag) != "tc":
                continue
            addr = next((x for x in tc if _ln(x.tag) == "cellAddr"), None)
            span = next((x for x in tc if _ln(x.tag) == "cellSpan"), None)
            r = int(addr.get("rowAddr", seq_r)) if addr is not None else seq_r
            c = int(addr.get("colAddr", seq_c)) if addr is not None else seq_c
            cs = int(span.get("colSpan", 1)) if span is not None else 1
            txt = " ".join(_hx_text(sl) for sl in tc if _ln(sl.tag) == "subList") or _hx_text(tc)
            cells.append((r, c, txt))
            seq_c = c + max(1, cs)
        seq_r += 1
    if not cells:
        return []
    h, w = max(r for r, _, _ in cells) + 1, max(c for _, c, _ in cells) + 1
    grid = [[""] * w for _ in range(h)]
    for r, c, t in cells:
        if r < h and c < w:
            grid[r][c] = t
    return _pad(grid)


def hwpx_blocks(path):
    """한글(hwpx) — section XML 을 순서대로. 문단 글은 ('p', 글), 표는 ('table', 줄들). 예전엔 글(hp:t)만 이어 붙여 표 칸이 한 줄씩 풀렸다"""
    from xml.etree import ElementTree as ET
    out = []

    def para(p):
        buf = []

        def rec(e):
            for c in e:
                n = _ln(c.tag)
                if n == "tbl":
                    if buf:
                        out.append(("p", "".join(buf)))
                        buf.clear()
                    rows = _hx_table(c)
                    if rows:
                        out.append(("table", rows))
                elif n == "t":
                    buf.append("".join(c.itertext()))
                else:
                    rec(c)
        rec(p)
        if "".join(buf).strip():
            out.append(("p", "".join(buf)))

    with zipfile.ZipFile(path) as z:
        for n in sorted((x for x in z.namelist() if re.match(r"Contents/section\d+\.xml$", x)), key=lambda x: int(re.sub(r"\D", "", x) or 0)):
            root = ET.fromstring(z.read(n))
            for el in root:
                if _ln(el.tag) == "p":
                    para(el)
                elif _ln(el.tag) == "tbl":
                    rows = _hx_table(el)
                    if rows:
                        out.append(("table", rows))
    return out


def _pptx_shapes(shapes):
    for sh in shapes:
        if getattr(sh, "shape_type", None) == 6 or sh.__class__.__name__ == "GroupShape":   # 묶음 도형 안까지
            yield from _pptx_shapes(sh.shapes)
        else:
            yield sh


def chart_rows(chart):
    """pptx 차트 → 줄들(머리 줄 = 구분 + 계열 이름). 범주·값을 읽을 수 없으면 []"""
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = [(s.name or f"계열{i + 1}", list(s.values)) for i, s in enumerate(plot.series)]
    except Exception:
        return []
    if not cats or not series:
        return []
    rows = [["구분"] + [n for n, _ in series]]
    for i, c in enumerate(cats):
        rows.append([cell_str(c)] + [fmt(v[i]) if i < len(v) and v[i] is not None else "" for _, v in series])
    return rows


def pptx_blocks(path):
    """쪽마다 글·표·차트(범주·계열 값). 차트 값은 그림이 아니라 숫자로 읽힌다"""
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, s in enumerate(prs.slides, 1):
        out.append(("h", f"{i}쪽"))
        for sh in _pptx_shapes(s.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append(("p", ("  " * min(p.level, 3)) + "- " + t))
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(("table", _pad([[c.text for c in row.cells] for row in sh.table.rows])))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                rows = chart_rows(sh.chart)
                if rows:
                    title = ""
                    try:
                        if sh.chart.has_title:
                            title = sh.chart.chart_title.text_frame.text.strip()
                    except Exception:
                        pass
                    out.append(("chart", (title, _pad(rows))))   # 제목이 없으면 tables_of 가 쪽 제목(바로 앞 글)을 붙인다
    return out


PDF_TABLE_PAGES = 80   # 표는 앞 80쪽만(RFP 뒤쪽은 서식·별지)


def _pdf_cache(path, kind):
    """같은 파일(크기·수정 시각)을 다시 읽지 않게 — 자료 옆 .<kind>.json. 다시 돌리기·자료 수집이 같은 pdf 를 또 읽었다(R95)"""
    try:
        st = os.stat(path)
        return f"{path}.{kind}.json", f"{st.st_size}:{int(st.st_mtime)}"
    except OSError:
        return None, None


def _cache_get(path, kind):
    cp, sig = _pdf_cache(path, kind)
    try:
        with open(cp, encoding="utf-8") as f:
            d = json.load(f)
        return d["data"] if d.get("sig") == sig else None
    except Exception:
        return None


def _cache_put(path, kind, data):
    cp, sig = _pdf_cache(path, kind)
    if cp:
        try:
            with open(cp, "w", encoding="utf-8") as f:
                json.dump({"sig": sig, "data": data}, f, ensure_ascii=False)
        except Exception:
            pass


def _pdf_procs(n_pages):
    """쪽을 나눠 읽을 프로세스 수 — 코어 하나가 느린 서버(회사 서버 Xeon 2.0GHz)에서 129쪽 RFP 가 글 49초·표 127초 걸렸다(R95)"""
    return max(1, min(8, (os.cpu_count() or 2) // 3, n_pages // 8))


def _pdf_tables_part(path, start, end):
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for pi in range(start, end):
            for t in pdf.pages[pi].extract_tables() or []:
                rows = _pad([[c or "" for c in r] for r in t if r])
                if len(rows) >= 2 and any(any(r) for r in rows):
                    out.append({"file": os.path.basename(path), "where": f"{pi + 1}쪽", "title": "", "rows": rows})
    return out


def _pdf_text_part(path, start, end):
    from pypdf import PdfReader
    pages = PdfReader(path).pages
    return [(pages[i].extract_text() or "") for i in range(start, end)]


def _split_run(fn, path, n, procs):
    """0..n 쪽을 procs 개 묶음으로 나눠 프로세스로 돌리고 순서대로 합친다. 실패하면 한 프로세스로"""
    if procs <= 1:
        return [fn(path, 0, n)]
    step = -(-n // procs)
    spans = [(s, min(n, s + step)) for s in range(0, n, step)]
    try:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=len(spans)) as ex:
            return list(ex.map(fn, [path] * len(spans), [a for a, _ in spans], [b for _, b in spans]))
    except Exception:
        return [fn(path, 0, n)]


def pdf_page_count(path):
    from pypdf import PdfReader
    return len(PdfReader(path).pages)


def pdf_text(path):
    """pdf 글(쪽 순서대로) — 쪽을 나눠 여러 프로세스로, 같은 파일이면 저장해 둔 것"""
    got = _cache_get(path, "text")
    if got is not None:
        return got
    n = pdf_page_count(path)
    text = "\n".join(t for part in _split_run(_pdf_text_part, path, n, _pdf_procs(n)) for t in part)
    _cache_put(path, "text", text)
    return text


def pdf_tables(path):
    """pdf 표 — pdfplumber 가 있을 때만(엔진 서버: pip install pdfplumber). 없으면 [] (글은 pypdf 로 따로 읽는다)"""
    try:
        import pdfplumber  # noqa: F401
    except ImportError:
        return []
    got = _cache_get(path, "tables")
    if got is not None:
        return got
    try:
        n = min(PDF_TABLE_PAGES, pdf_page_count(path))
        out = [t for part in _split_run(_pdf_tables_part, path, n, _pdf_procs(n)) for t in part]
    except Exception:
        return []
    _cache_put(path, "tables", out)
    return out


def md_table(rows, max_rows=60):
    """줄들 → 마크다운 파이프 표(머리 줄 + 구분 줄). 빈 칸은 빈 칸으로 남아 열이 밀리지 않는다"""
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [list(r) + [""] * (w - len(r)) for r in rows]
    line = lambda r: "| " + " | ".join(c if c else " " for c in r) + " |"   # noqa: E731
    out = [line(rows[0]), "|" + "---|" * w] + [line(r) for r in rows[1:max_rows + 1]]
    if len(rows) - 1 > max_rows:
        out.append(f"(이하 {len(rows) - 1 - max_rows}줄 생략)")
    return "\n".join(out)


def tables_of(path):
    """파일 하나의 표 전부 — [{"file","where","title","rows"}]. 모르는 형식·실패면 []"""
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    try:
        if ext == ".xlsx":
            return xlsx_tables(path)
        if ext == ".pdf":
            return pdf_tables(path)
        if ext == ".csv":
            import csv
            with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
                rows = [r for r in csv.reader(f)][:MAX_ROWS + 1]
            return [{"file": name, "where": "", "title": os.path.splitext(name)[0], "rows": _pad(rows)}] if rows else []
        blocks = docx_blocks(path) if ext == ".docx" else hwpx_blocks(path) if ext == ".hwpx" else pptx_blocks(path) if ext == ".pptx" else []
    except Exception:
        return []
    out, last_p, where = [], "", ""
    for k, v in blocks:
        if k == "h":
            where = v
        elif k == "p":
            last_p = v.strip().lstrip("- ").strip()
        elif k == "table":
            out.append({"file": name, "where": where, "title": last_p[:60] if len(last_p) <= 60 else "", "rows": v})
        elif k == "chart":
            out.append({"file": name, "where": where, "title": v[0] or last_p[:60], "rows": v[1], "chart": True})
    return out


# ── 표 → 데이터셋 ──
TIME_LABEL = re.compile(r"^\s*(?:(?:19|20)\d{2}\s*(?:년|\.|-|/)?\s*(?:\d{1,2}\s*월?|[1-4]\s*분기|[상하]반기|Q[1-4])?|\d{1,2}\s*월|[1-4]\s*분기|Q[1-4]|[상하]반기|FY\s?\d{2,4})\s*$", re.I)
YM = re.compile(r"^\s*((?:19|20)\d{2})\s*(?:년|\.|-|/)\s*(\d{1,2})\s*월?\s*$")
BEFORE = re.compile(r"(도입|개선|적용|구축|시행)\s*(?:전|이전)|^\s*(?:이전|기존|현행|현재|as-?is|before)\b", re.I)
AFTER = re.compile(r"(도입|개선|적용|구축|시행)\s*(?:후|이후)|^\s*(?:이후|개선|목표|to-?be|after)\b", re.I)
SHARE_HEAD = re.compile(r"(비율|구성비|비중|점유|share|%)", re.I)
TARGET_HEAD = re.compile(r"(목표|target|달성)", re.I)
RANK_HEAD = re.compile(r"(순위|rank)", re.I)
STEP_HEAD = re.compile(r"^\s*(단계|절차|순서|step|공정)", re.I)


def _blocks(rows):
    """빈 줄로 나뉜 표 덩어리들(시트 하나에 표 여러 개)"""
    cur, out = [], []
    for r in rows:
        if any(c.strip() for c in r):
            cur.append(r)
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def _unit_of(head, vals_units):
    m = UNIT_HEAD.search(head or "")
    if m:
        return m.group(1).strip()
    us = [u for u in vals_units if u]
    return max(set(us), key=us.count) if us else ""


def _clean_head(h):
    return UNIT_HEAD.sub("", h or "").strip()


def datasets_from_table(t):
    """표 하나 → 데이터셋 0개 이상. 숫자 열이 하나 이상이고 라벨 열이 있는 덩어리만(머리 줄 1 + 자료 줄 2 이상).
    데이터셋 = {title, kind_hint, unit, series, rows:[[라벨, 값…]], source:{file, where}} — 값은 숫자(빈 칸은 None)"""
    out = []
    for bi, blk in enumerate(_blocks(t.get("rows") or [])):
        used = [j for j in range(max(len(r) for r in blk)) if any(j < len(r) and r[j].strip() for r in blk)]
        blk = [[r[j] if j < len(r) else "" for j in used] for r in blk]
        title = t.get("title") or ""
        if bi and title and title == t.get("where"):   # 시트의 둘째 표부터 — 시트 이름만으로는 어느 표인지 모른다
            title = ""
        while len(blk) > 3 and sum(1 for c in blk[0] if c.strip()) == 1 and len(blk[0]) > 1:   # 제목 줄(칸 하나만 찬 줄)
            title = next(c for c in blk[0] if c.strip())
            blk = blk[1:]
        if len(blk) < 3 or len(blk[0]) < 2:
            continue
        head, data = blk[0], blk[1:MAX_ROWS + 1]
        w = len(head)
        parsed = [[num(r[j]) for j in range(w)] for r in data]
        numeric = []
        for j in range(w):
            filled = [r[j] for r in data if r[j].strip()]
            ok = sum(1 for p in parsed if p[j][0] is not None)
            if ok >= 2 and ok >= 0.6 * max(1, len(filled)):
                numeric.append(j)
        # 첫 열이 연도(2023·2024…)면 숫자가 아니라 라벨이다 — 시계열의 시점
        if numeric and numeric[0] == 0 and all(re.match(r"^\s*(?:19|20)\d{2}\s*년?\s*$", r[0]) for r in data if r[0].strip()):
            numeric = numeric[1:]
        label = next((j for j in range(w) if j not in numeric), None)
        src = {"file": t.get("file", ""), "where": t.get("where", "")}
        if not numeric or label is None:
            if STEP_HEAD.match(head[0] or "") and len(data) >= 2:   # 단계·절차 표 — 숫자 없이도 흐름 재료
                out.append({"title": title or _clean_head(head[0]), "kind_hint": "process", "unit": "", "series": [_clean_head(h) for h in head[1:]],
                            "rows": [[r[0]] + list(r[1:]) for r in data if r[0].strip()][:20], "source": src})
            continue
        labels = [r[label].strip() for r in data]
        keep = [i for i, lab in enumerate(labels) if lab and any(parsed[i][j][0] is not None for j in numeric)]
        if len(keep) < 2:
            continue
        hunits = {m.group(1).strip() for m in (UNIT_HEAD.search(head[j] or "") for j in numeric) if m}
        if len(hunits) > 1:   # 열마다 단위가 다르면(명·건·분) 단위를 하나로 묶지 않고 계열 이름에 남긴다
            heads, unit = [head[j].strip() or f"값{k + 1}" for k, j in enumerate(numeric)], ""
        else:
            heads = [_clean_head(head[j]) or f"값{k + 1}" for k, j in enumerate(numeric)]
            unit = _unit_of(head[numeric[0]], [parsed[i][j][1] for i in keep for j in numeric])
        rows = [[labels[i]] + [parsed[i][j][0] for j in numeric] for i in keep]
        timeish = sum(1 for i in keep if TIME_LABEL.match(labels[i])) >= 0.7 * len(keep)
        hs = " ".join(head[j] for j in numeric)
        if TARGET_HEAD.search(hs):
            kind = "target"
        elif len(numeric) == 2 and BEFORE.search(head[numeric[0]]) and AFTER.search(head[numeric[1]]):
            kind = "before_after"
        elif timeish:
            kind = "trend"
        elif RANK_HEAD.search(hs) or RANK_HEAD.search(head[label]):
            kind = "rank"
        elif SHARE_HEAD.search(hs) or (unit == "%" and len(numeric) == 1 and abs(sum(r[1] or 0 for r in rows) - 100) <= 2):
            kind = "share"
        else:
            kind = "compare"
        if kind == "share" and not unit:
            unit = "%"
        ds = {"title": title or (_clean_head(head[label]) + "별 " + " · ".join(heads)),
              "kind_hint": kind, "unit": unit, "label": _clean_head(head[label]), "series": heads, "rows": rows, "source": src}
        if t.get("chart"):
            ds["from_chart"] = True
        out.append(ds)
        tot = yearly_totals(ds)
        if tot:
            out.append(tot)
    return out


def yearly_totals(ds):
    """월별 시계열(2023-01 … 2025-12)이 2년 이상 12달씩 다 있으면 연도별 합계 데이터셋 — 계산값임을 밝힌다(derived)"""
    if ds.get("kind_hint") != "trend":
        return None
    by = {}
    for r in ds["rows"]:
        m = YM.match(str(r[0]))
        if not m:
            return None
        by.setdefault(m.group(1), []).append(r[1:])
    full = {y: v for y, v in by.items() if len(v) == 12}
    if len(full) < 2:
        return None
    rows = []
    for y in sorted(full):
        sums = []
        for k in range(len(ds["series"])):
            col = [x[k] for x in full[y] if k < len(x)]
            sums.append(None if any(c is None for c in col) else round(sum(col), 4))
        rows.append([f"{y}년"] + sums)
    return {"title": ds["title"] + " — 연도별 합계", "kind_hint": "trend", "unit": ds.get("unit", ""), "label": "연도", "series": ds["series"],
            "rows": rows, "source": dict(ds["source"], note="계산: 월별 값의 합"), "derived": True}


# ── 데이터셋 → 대표 수치(facts) ──
def facts_from_dataset(ds, did):
    """데이터셋의 대표 수치 몇 개 — 처음·끝(추이), 짝(전·후), 각 몫(구성비), 최대·최소(비교)"""
    out, rows, unit = [], ds.get("rows") or [], ds.get("unit", "")
    if ds.get("kind_hint") == "process" or not rows:
        return out

    def add(r, k=0, metric=None):
        v = r[1 + k] if 1 + k < len(r) else None
        if v is None:
            return
        ser = (ds.get("series") or [""])[k] if k < len(ds.get("series") or []) else ""
        out.append({"metric": metric or f"{ds['title']} — {r[0]}" + (f" · {ser}" if ser and len(ds.get('series') or []) > 1 else ""),
                    "value": v, "unit": unit, "period": r[0] if TIME_LABEL.match(str(r[0])) else "", "group": ser if len(ds.get("series") or []) > 1 else "",
                    "source": ds["source"], "quote": (" | ".join([str(r[0])] + [fmt(x) for x in r[1:]]))[:80], "dataset": did,
                    **({"derived": True} if ds.get("derived") else {})})
    k = ds["kind_hint"]
    if k == "trend":
        add(rows[0])
        add(rows[-1])
    elif k == "before_after":
        for r in rows[:4]:
            add(r, 0)
            add(r, 1)
    elif k in ("share", "target"):
        for r in rows[:6]:
            add(r)
            if k == "target" and len(r) > 2:
                add(r, 1)
    else:
        vals = [(r[1], r) for r in rows if r[1] is not None]
        if vals:
            add(max(vals, key=lambda x: x[0])[1])
            add(min(vals, key=lambda x: x[0])[1])
    return out


# ── 글 속 수치 ──
# 단위 뒤에는 글 끝·기호·조사만 — '1분기'·'3분야' 같은 낱말 속 글자를 단위로 읽지 않게
_TAIL = r"(?=$|[^가-힣A-Za-z]|(?:이상|이하|미만|초과|이내|까지|부터|에서|으로|이며|이다|이고|이었|였|에서|대비|증가|감소|수준|안팎|가량|정도|에\s?그|로|에|이|가|을|를|은|는|의|와|과|도|씩|당|만|씩|째))"
PROSE_NUM = re.compile(r"(?<![\d.])(\d[\d,]*(?:\.\d+)?)\s*(%|퍼센트|천\s?건|만\s?건|건|명|분|초|시간|일|개월|천\s?원|만\s?원|억\s?원|원|억|조|점|개소|개|곳|회|배|대|종|석|㎡|km)" + _TAIL)
YEAR = re.compile(r"((?:19|20)\d{2})\s*년")


def prose_facts(text, file="", limit=80):
    """표가 아닌 글에서 '숫자+단위' 를 문장째 — 이름(metric)은 숫자 앞 낱말, quote 는 원문 80자. AI 없이"""
    out, seen = [], set()
    for sent in re.split(r"(?<=[.!?。])\s+|\n+", text or ""):
        s = sent.strip()
        if not s or s.startswith("|") or len(s) < 6:
            continue
        years = YEAR.findall(s)
        for m in PROSE_NUM.finditer(s):
            raw_v, unit = m.group(1), re.sub(r"\s", "", m.group(2))
            if unit in ("년",):
                continue
            try:
                v = float(raw_v.replace(",", ""))
            except ValueError:
                continue
            pre = re.sub(r"[\s·,:()\[\]「」'\"]+$", "", s[:m.start()])
            words = re.findall(r"[가-힣A-Za-z][가-힣A-Za-z0-9]*", pre)[-3:]
            metric = " ".join(re.sub(r"(은|는|이|가|을|를|의|에|로|으로|에서|와|과|도)$", "", w_) for w_ in words).strip() or "수치"
            key = (raw_v, unit, metric)
            if key in seen:
                continue
            seen.add(key)
            a = max(0, m.start() - 40)
            out.append({"metric": metric, "value": v, "unit": unit, "period": years[0] + "년" if years else "", "group": "",
                        "source": {"file": file, "where": ""}, "quote": s[a:a + 80]})
            if len(out) >= limit:
                return out
    return out


def facts_from_docs(docs):
    """[(파일 이름, 경로, 글)] → {datasets, facts}. AI 없이 정해진다(같은 자료면 같은 결과)"""
    datasets, facts = [], []
    for name, path, text in docs:
        for t in tables_of(path):
            t["file"] = name
            for ds in datasets_from_table(t):
                if len(datasets) < MAX_DATASETS:
                    datasets.append(ds)
        prose = "\n".join(l for l in (text or "").splitlines() if not l.lstrip().startswith("|"))
        facts += prose_facts(prose, name)
    for i, ds in enumerate(datasets, 1):
        ds["id"] = f"D{i}"
    ds_facts = [f for i, ds in enumerate(datasets, 1) for f in facts_from_dataset(ds, f"D{i}")]
    facts = (ds_facts + facts)[:MAX_FACTS]
    for i, f in enumerate(facts, 1):
        f["id"] = f"F{i}"
    return {"datasets": datasets, "facts": facts}


# ── AI 이름표(값은 못 바꾼다) ──
LABEL_SYSTEM = """너는 공공 사업 자료의 수치에 이름표를 붙이는 분석가다. 아래 후보는 원문에서 기계로 뽑은 '숫자+단위'와 그 둘레 원문이다.
후보마다 무엇을 잰 값인지(metric, 12~30자 명사구), 시점(period, 원문에 있을 때만), 묶음(group: 기관·채널·유형 등, 있을 때만)을 붙인다.
값(value)·단위는 바꾸거나 새로 만들 수 없다. 사업과 상관없는 수치(쪽 번호·법 조항·전화번호·서식 번호)는 keep:false.
JSON 하나만: {"labels":[{"id":"F12","metric":"…","period":"…","group":"…","keep":true}]}"""


def label_prompt(facts, limit=60):
    """글 속 수치 후보(데이터셋에서 온 것 제외) → AI 에게 줄 user 글. 후보가 없으면 ''"""
    rows = [f for f in facts if not f.get("dataset")][:limit]
    if not rows:
        return ""
    return "\n".join(f'{f["id"]} | {fmt(f["value"])}{f.get("unit", "")} | 원문: {f.get("quote", "")}' for f in rows)


def apply_labels(facts, answer):
    """AI 답(JSON 글이나 dict)을 후보에 입힌다 — id 가 맞는 것만, 값은 건드리지 않는다. 깨진 답이면 그대로 둔다. 반환: 바뀐 개수"""
    if isinstance(answer, str):
        try:
            m = re.search(r"\{.*\}", answer, re.S)
            answer = json.loads(m.group(0)) if m else {}
        except Exception:
            return 0
    labels = (answer or {}).get("labels") if isinstance(answer, dict) else None
    if not isinstance(labels, list):
        return 0
    by = {f["id"]: f for f in facts}
    n = 0
    for x in labels:
        if not isinstance(x, dict) or str(x.get("id")) not in by:
            continue
        f = by[str(x["id"])]
        if f.get("dataset"):
            continue
        if x.get("keep") is False:
            f["drop"] = True
            n += 1
            continue
        for k, lim in (("metric", 40), ("period", 20), ("group", 20)):
            v = str(x.get(k) or "").strip()
            if v:
                f[k] = v[:lim]
        n += 1
    facts[:] = [f for f in facts if not f.get("drop")]
    return n


# ── 기획·와이어에 넣는 글 ──
def _ds_line(ds):
    rows = ds.get("rows") or []
    span = f"{rows[0][0]} ~ {rows[-1][0]}" if ds.get("kind_hint") == "trend" and rows else ", ".join(str(r[0]) for r in rows[:5]) + (" …" if len(rows) > 5 else "")
    src = ds["source"].get("file", "") + (f" · {ds['source']['where']}" if ds["source"].get("where") else "")
    return (f"- {ds['id']} [{ds['kind_hint']}] {ds['title']}" + (f" ({ds['unit']})" if ds.get("unit") else "") + f" — {len(rows)}줄: {span}"
            + (f" · 계열: {'·'.join(ds.get('series') or [])}" if len(ds.get("series") or []) > 1 else "") + (" · 계산값" if ds.get("derived") else "") + f" · 출처 {src}")


def catalog_text(fx, limit=7000):
    """기획 AI 에 줄 '자료의 실제 데이터' 목록 — 장마다 data 에 id 를 적게 한다. 데이터가 없으면 ''"""
    if not fx or not (fx.get("datasets") or fx.get("facts")):
        return ""
    out = ["# 자료의 실제 데이터 — 올린 파일의 표·시계열(D)과 낱개 수치(F). 수치로 주장하는 장은 slide 의 \"data\" 에 쓸 id 를 적는다",
           "(그래프·표·KPI 는 와이어 단계에서 이 값 그대로 그린다. 여기 없는 수치를 만들지 않는다)"]
    if fx.get("datasets"):
        out.append("## 표·시계열")
        out += [_ds_line(ds) for ds in fx["datasets"]]
    if fx.get("facts"):
        out.append("## 낱개 수치")
        for f in fx["facts"]:
            if f.get("dataset"):
                continue
            out.append(f"- {f['id']} {f['metric']} {fmt(f['value'])}{f.get('unit', '')}" + (f" ({f['period']})" if f.get("period") else "")
                       + f" — {f['source'].get('file', '')}")
    t = "\n".join(out)
    return t[:limit] + ("\n…(이하 생략)" if len(t) > limit else "")


def _nums_in(text):
    return {re.sub(r"[,\s]", "", m.group(1)) for m in re.finditer(r"(?<![\d.])(\d[\d,]*(?:\.\d+)?)", text or "") if len(re.sub(r"\D", "", m.group(1))) >= 2}


def _ds_nums(ds):
    return {re.sub(r",", "", fmt(v)) for r in ds.get("rows") or [] for v in r[1:] if isinstance(v, (int, float))}


def slide_refs(s, fx):
    """이 장이 쓰는 데이터 — 기획이 적은 data id, 없으면 장 글의 수치와 2개 이상 겹치는 데이터셋·값이 같은 낱개 수치(자동 짝)"""
    if not fx:
        return [], []
    dsb = {d["id"]: d for d in fx.get("datasets") or []}
    fb = {f["id"]: f for f in fx.get("facts") or []}
    want = [str(x).strip().upper() for x in (s.get("data") or []) if str(x).strip()]
    ds = [dsb[x] for x in want if x in dsb]
    fs = [fb[x] for x in want if x in fb]
    if not ds and not fs:
        nums = _nums_in("\n".join(str(s.get(k) or "") for k in ("title", "message", "content")))
        if nums:
            ds = [d for d in dsb.values() if len(nums & _ds_nums(d)) >= 2][:2]
            fs = [f for f in fb.values() if not f.get("dataset") and re.sub(",", "", fmt(f["value"])) in nums][:6]
    return ds, fs


def ds_table(ds, max_rows=36):
    """데이터셋 → 와이어 AI 가 읽는 작은 표(값은 원래 숫자 그대로)"""
    head = [ds.get("label") or "구분"] + list(ds.get("series") or [])
    rows = ds.get("rows") or []
    lines = [" | ".join(head)] + [" | ".join([str(r[0])] + [fmt(v) if isinstance(v, (int, float)) else (str(v) if v else "") for v in r[1:]]) for r in rows[:max_rows]]
    if len(rows) > max_rows:
        lines.append(f"(이하 {len(rows) - max_rows}줄 생략 — 위 값 가운데서 고른다)")
    return "\n".join(lines)


def slide_data_text(s, fx):
    """와이어 AI 에 이 장과 함께 넣는 '실제 데이터' 덩어리. 쓸 데이터가 없으면 ''"""
    ds, fs = slide_refs(s, fx)
    if not ds and not fs:
        return ""
    out = ["데이터(올린 자료의 실제 값 — chart·kpi·table·gauge·progress 의 수치는 여기서만 가져온다):"]
    for d in ds:
        src = d["source"].get("file", "") + (f" · {d['source']['where']}" if d["source"].get("where") else "")
        out.append(f"[{d['id']} · 유형 {d['kind_hint']}" + (f" · 단위 {d['unit']}" if d.get("unit") else "") + f"] {d['title']} (출처: {src}"
                   + (", 계산값 — 월별 합" if d.get("derived") else "") + ")")
        out.append(ds_table(d))
    for f in fs:
        out.append(f"- {f['id']} {f['metric']}: {fmt(f['value'])}{f.get('unit', '')}" + (f" ({f['period']})" if f.get("period") else "") + f" — 원문 「{f.get('quote', '')}」")
    return "\n".join(out)


def sources_text(fx):
    """근거 확인(fact_check)에 더할 글 — 데이터셋 모든 값과 낱개 수치(계산값 포함). 단위를 붙여 '1,820건' 꼴도 맞게"""
    if not fx:
        return ""
    out = []
    for d in fx.get("datasets") or []:
        u = d.get("unit", "")
        out.append(d.get("title", ""))
        for r in d.get("rows") or []:
            # 단위가 행 이름 괄호에만 있는 표('상담 응답률(%) | 71.4 | 95') — 그 단위를 값에 붙인다(E2E 382: '95%' 를 자료에 없다고 오탐, 2026-10-18)
            m = re.search(r"\(\s*(%|[가-힣A-Za-z]{1,3})\s*\)\s*$", str(r[0])) if not u else None
            ru = u or (m.group(1) if m else "")
            out.append(" ".join([str(r[0])] + [f"{fmt(v)}{ru}" for v in r[1:] if isinstance(v, (int, float))]))
    for f in fx.get("facts") or []:
        out.append(f"{f.get('metric', '')} {fmt(f['value'])}{f.get('unit', '')}")
    return "\n".join(out)
