# -*- coding: utf-8 -*-
"""제작 뒤 마무리(결정적) — 에이전트가 놓쳐도 항상 같게.
1) 서체: 모든 글을 서체 한 벌(fontset — a시월구일1~4, 없으면 Asta Sans)로. 한글 칸(ea)까지 같은 서체를 넣고,
   bold 로 굵게 한 글은 굵은 서체(3·4)로 바꾼다. 키메시지는 G마켓 산스 TTF Bold 25pt, 기획의 강조어(keyword)는 강조색
2) 키메시지 자리·아래 여백·좌우 가이드를 전 장 같게 (조금 어긋난 것은 본문을 통째로 맞춘다)
3) 도형 안 글: 한 줄 글의 고정 줄간격을 풀고 세로 가운데
4) 장표에 쓴 서체 파일을 <fonts_dir> 에 모으고, 장마다 글을 <text.json> 으로 남긴다(발표 대본이 읽는다)
인자: <pptx> <fonts_dir> [--fonts-only] [--text <json>] [--plan <plan.json>] [--accent #RRGGBB]
고칠 수 없었던 것은 'WARN: N쪽|제목|설명' 줄로 알린다(워커가 확인 목록에 올린다).
"""
import copy, json, os, re, shutil, sys, zipfile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
import fontlicense, fontset, osdep

sys.stdout.reconfigure(encoding="utf-8")
src, fonts_dir = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
FONTS_ONLY = "--fonts-only" in sys.argv
TEXT_OUT = sys.argv[sys.argv.index("--text") + 1] if "--text" in sys.argv else None
PLAN = json.load(open(sys.argv[sys.argv.index("--plan") + 1], encoding="utf-8")) if "--plan" in sys.argv else {}
ACCENT = sys.argv[sys.argv.index("--accent") + 1] if "--accent" in sys.argv else "#FF2A6D"
ACCENT = ACCENT[:7] if re.match(r"^#[0-9A-Fa-f]{6}", ACCENT or "") else "#FF2A6D"
KEY_FONT, KEY_PT = fontset.KEY_FONT, 25
KEY_PT_EFF = round(KEY_PT * fontset.KEY_SCALE)   # 제목 서체 덩어리감 보정한 키메시지 크기
KEY_ALIGN = PP_ALIGN.LEFT if "--key-align" in sys.argv and sys.argv[sys.argv.index("--key-align") + 1] == "left" else PP_ALIGN.CENTER   # 디자인 타입(B·D 는 왼쪽)
KEEP = re.compile(r"^(G마켓|Gmarket|Wingdings|Webdings|Symbol|Segoe UI (Emoji|Symbol)|Marlett|\+m)", re.I)   # 서체 한 벌로 바꾸지 않는 것
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
IN = 914400
KEY_TOP, KEY_H, KEY_L, KEY_W, BODY_TOP, BODY_BOT, G_L, G_R = 1.03, 0.45, 0.20, 10.43, 1.65, 6.95, 0.30, 10.53
warns = []


def warn(no, title, detail):
    warns.append(f"WARN: {no}|{title}|{detail}")


def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


def set_font(run, face):
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(A + tag)
        if el is None:
            el = rPr.makeelement(A + tag, {})
            rPr.append(el)
        el.set("typeface", face)


def max_pt(s):
    v = [r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(v) if v else 0


def fit_pt(text, width_in):
    """키메시지가 한 줄에 들어가는 가장 큰 크기(23 → 17pt). 실제 서체로 잰다"""
    try:
        from PIL import ImageFont
        path = fontset.INSTALLED.get(KEY_FONT.lower()) or fontset.title_file()   # 설치 이름으로 못 찾으면 엔진 서체 파일로
        w100 = ImageFont.truetype(path, 100).getlength(text)   # 100pt 일 때 폭(pt)
    except Exception:
        w100 = len(text) * 88
    for pt in range(KEY_PT_EFF, 16, -1):
        if w100 * pt / 100 <= width_in * 72:
            return pt
    return 17


def unify_fonts(part_el):
    """글 런 하나하나를 서체 한 벌로. latin 만 적힌 런은 한글이 테마 기본 서체(보통 굵기)로 나오므로 ea·cs 에도 같은 서체를 넣는다.
    bold 로 굵게 한 글은 굵은 서체로 바꾸고 bold 를 지운다(가짜 굵게 대신 진짜 굵은 서체). 반환: 바꾼 런 수"""
    n = 0
    for r in part_el.iter(A + "r", A + "fld"):
        rPr = r.find(A + "rPr")
        if rPr is None:
            rPr = r.makeelement(A + "rPr", {})
            r.insert(0, rPr)
        n += fix_rpr(rPr)
    for tag in ("endParaRPr", "defRPr"):   # 빈 문단·목록 기본값 — 서체가 적힌 것만
        for rPr in part_el.iter(A + tag):
            if rPr.find(A + "latin") is not None or rPr.find(A + "ea") is not None:
                fix_rpr(rPr)
    return n


def fix_rpr(rPr):
    face = next((el.get("typeface") for el in (rPr.find(A + "latin"), rPr.find(A + "ea")) if el is not None and el.get("typeface")), None)
    if fontset.LANG_TAG != "ko-KR":   # 장표 언어(ja-JP·zh-CN …) — PowerPoint 가 일본어·중국어 글을 그 언어로 읽고 ea 서체로 그린다(2026-10-18)
        rPr.set("lang", fontset.LANG_TAG)
    if face and face != KEY_FONT and (face in fontset.TITLE_FACES or re.match(r"^(G마켓|Gmarket)", face, re.I)) and not fontset.face_usable(face):
        face = KEY_FONT   # 장표 언어 글자가 없는 제목 서체(일본어 + G마켓 산스 등) — 언어에 맞게 고른 제목 서체로(fontset.for_lang)
        for tag in ("latin", "ea", "cs"):
            el = rPr.find(A + tag)
            if el is not None:
                el.set("typeface", KEY_FONT)
    if face and (KEEP.match(face) or face in fontset.TITLE_FACES):   # 제목 서체(키메시지·표지 큰 글자)는 그대로
        lvl = None
    else:
        lvl = fontset.LEVEL.get(face, 1) if face else 1
        if rPr.get("b") == "1":
            lvl = 3 if lvl >= 2 else 2
        if "b" in rPr.attrib:
            del rPr.attrib["b"]
    new = face if lvl is None else fontset.face(lvl)
    if not new:
        return 0
    if lvl is not None and lvl >= 2 and int(rPr.get("sz") or 0) >= fontset.BODY_TRACK_MIN:   # 본문 굵게·크게 — 자간 조금 좁게
        rPr.set("spc", str(min(int(rPr.get("spc") or 0), round(fontset.BODY_TRACK * int(rPr.get("sz"))))))
    before = [A + t for t in ("ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill")]
    for tag in ("latin", "ea", "cs"):   # 스키마 순서: … 채우기·효과 뒤에 latin → ea → cs
        el = rPr.find(A + tag)
        if el is None:
            el = rPr.makeelement(A + tag, {})
            prev = [c for c in rPr if c.tag in before]
            if prev:
                prev[-1].addnext(el)
            else:
                rPr.insert(0, el)
        el.set("typeface", new)
        before.append(A + tag)
    for tag in ("latin", "ea", "cs"):   # panose·charset 등 예전 서체의 속성은 지운다
        for k in ("panose", "pitchFamily", "charset"):
            rPr.find(A + tag).attrib.pop(k, None)
    return 1


def keywords(no):
    """기획의 강조어 — plan.slides[no-1].keyword('A, B') 또는 title 의 **강조**"""
    sl = (PLAN.get("slides") or [])
    s = sl[no - 1] if 0 < no <= len(sl) else {}
    ks = [k.strip() for k in re.split(r"[,，/·]", str(s.get("keyword") or "")) if k.strip()]
    ks += re.findall(r"\*\*(.+?)\*\*", str(s.get("title") or ""))
    return [k for k in dict.fromkeys(ks) if len(k) >= 1]


def color_keywords(tf, words):
    """키메시지에 강조색 런이 하나도 없으면 강조어를 강조색으로 갈라 넣는다(시안의 분홍 글자를 장표에도)"""
    runs = [r for p in tf.paragraphs for r in p.runs]
    if not runs or not words:
        return False
    def col(r):
        try:
            return str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None
        except Exception:
            return None
    if len({col(r) for r in runs if r.text.strip()}) > 1:   # 이미 색이 갈려 있다 — 엔진이 시안대로 칠했다
        return False
    text = "".join(r.text for r in runs).replace("**", "")
    pat = "|".join(re.escape(w) for w in sorted(words, key=len, reverse=True))
    if not re.search(pat, text):
        return False
    p0 = tf.paragraphs[0]
    tmpl = copy.deepcopy(runs[0]._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        p0._p.remove(r._r)
    base = col(runs[0])
    end = p0._p.find(A + "endParaRPr")
    for i, part in enumerate(re.split(f"({pat})", text)):
        if not part:
            continue
        el = copy.deepcopy(tmpl)
        el.find(A + "t").text = part
        if end is not None:
            end.addprevious(el)
        else:
            p0._p.append(el)
    for r in p0.runs:
        hit = re.fullmatch(pat, r.text) is not None
        if hit:
            r.font.color.rgb = RGBColor.from_string(ACCENT[1:].upper())
        elif base:
            r.font.color.rgb = RGBColor.from_string(base)
    return True


def _lum(hexrgb):
    """상대 휘도(WCAG)"""
    c = [int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def _contrast(a, b):
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _rgb(colorfmt):
    try:
        return str(colorfmt.rgb) if colorfmt is not None and colorfmt.type is not None else None
    except Exception:
        return None


_FONTS = {}


def _measure(text, face, pt):
    """글 폭(pt). 설치된 서체 파일로 잰다 — 못 찾으면 글자 수로 어림(한글 1em, 그 밖 0.55em)"""
    from PIL import ImageFont
    path = fontset.INSTALLED.get((face or "").lower())
    if path and path not in _FONTS:
        try:
            _FONTS[path] = ImageFont.truetype(path, 200)
        except Exception:
            _FONTS[path] = None
    f = _FONTS.get(path) if path else None
    if f is not None:
        return f.getlength(text) / 200 * pt
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in text) * pt


def fit_one_line(no, slide, key):
    """줄바꿈 없이 놓인 글(줄바꿈 꺼짐 — 여러 문단이면 문단마다 한 줄, 또는 높이가 한 줄뿐인 상자)이 도형 안쪽 폭보다 길면
    그 상자의 글자를 같은 비율로 줄여 넣는다(크기 위계는 그대로). 쉐브론·오각형 화살표는 뾰족한 끝만큼 안쪽 폭을 덜 본다.
    원래 크기의 70% 밑으로는 줄이지 않고 경고를 남긴다"""
    A_ = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    n = 0
    for sh in walk(slide.shapes):
        if sh is key or not getattr(sh, "has_text_frame", False) or not sh.width:
            continue
        tf = sh.text_frame
        paras = [p for p in tf.paragraphs if p.text.strip() and p.runs]
        if not paras:
            continue
        sizes = [r.font.size.pt if r.font.size else 18.0 for p in paras for r in p.runs]
        # 줄바꿈이 켜진 짧은 제목(한 문단, 11pt 이상, 30자 안팎)이 끝 몇 글자만 다음 줄로 넘어가는 경우도 — 조금 줄여 한 줄로
        short_title = len(paras) == 1 and max(sizes) >= 11 and len(paras[0].text.strip()) <= 30
        if not (tf.word_wrap is False or (len(paras) == 1 and sh.height <= Pt(max(sizes) * 2.1)) or short_title):
            continue
        inner = (sh.width - (tf.margin_left or 0) - (tf.margin_right or 0)) / 12700   # EMU → pt
        try:
            ast_ = str(sh.auto_shape_type)
        except Exception:
            ast_ = ""
        if "CHEVRON" in ast_ or "PENTAGON" in ast_:
            inner -= sh.height / 12700 * (0.9 if "CHEVRON" in ast_ else 0.45)
        if inner <= 0:
            continue
        need = 0.0
        for p in paras:
            w = 0.0
            for r in p.runs:
                rp = r._r.find(A_ + "rPr")
                face = None
                if rp is not None:
                    for tag in ("ea", "latin"):
                        el = rp.find(A_ + tag)
                        if el is not None and el.get("typeface"):
                            face = el.get("typeface")
                            break
                w += _measure(r.text, face, r.font.size.pt if r.font.size else 18.0)
            need = max(need, w)
        if need <= inner * 0.98:
            continue
        if short_title and tf.word_wrap is not False and need > inner * 1.25:
            continue   # 많이 넘치는 제목은 두 줄이 맞다(시안도 두 줄일 것) — 줄이지 않는다
        k = max(0.7, inner * 0.96 / need)
        for p in paras:
            for r in p.runs:
                r.font.size = Pt(round((r.font.size.pt if r.font.size else 18.0) * k * 2) / 2)
        n += 1
        if k <= 0.7 and need * 0.7 > inner:
            warn(no, "도형 밖으로 넘치는 글", f"'{paras[0].text.strip()[:30]}' — 글자를 70%까지 줄여도 도형보다 깁니다. 글을 줄여 주세요")
    return n


def keep_aspect(slide):
    """잘라 넣은 그림이 원래 비율보다 늘어났으면(시안의 가로·세로 %를 서로 다른 배율로 옮기면 세로로 4~7% 길어진다)
    가운데를 지키며 원래 비율로 되돌린다. 크롭이 걸린 그림·틀 그림은 건드리지 않는다"""
    n = 0
    for sh in walk(slide.shapes):
        if sh.shape_type != 13 or (sh.name or "").startswith("pd_") or not sh.width or not sh.height:
            continue
        try:
            if sh.crop_left or sh.crop_right or sh.crop_top or sh.crop_bottom:
                continue
            iw, ih = sh.image.size
        except Exception:
            continue
        want = sh.width * ih / iw
        if abs(sh.height / want - 1) <= 0.015:
            continue
        cy = sh.top + sh.height / 2
        if sh.height > want:   # 세로로 늘어났다 → 높이를 줄인다
            sh.height = int(want)
            sh.top = int(cy - want / 2)
        else:                  # 가로로 늘어났다 → 폭을 줄인다
            cx = sh.left + sh.width / 2
            w2 = sh.height * iw / ih
            sh.width = int(w2)
            sh.left = int(cx - w2 / 2)
        n += 1
    return n


def title_gap(slide, key):
    """한 글상자 안에서 큰 글(제목) 다음에 작은 글(설명)이 오면 설명 문단 위에 간격을 둔다 —
    줄간격 100%·문단 간격 0 으로 쌓으면 제목과 설명이 붙어 보인다(시안은 떨어져 있다). 간격 = 제목 크기의 40%"""
    n = 0
    for sh in walk(slide.shapes):
        if sh is key or not getattr(sh, "has_text_frame", False):
            continue
        paras = [p for p in sh.text_frame.paragraphs if p.text.strip() and p.runs]
        for a, b in zip(paras, paras[1:]):
            za = max((r.font.size.pt for r in a.runs if r.font.size), default=0)
            zb = max((r.font.size.pt for r in b.runs if r.font.size), default=0)
            if za and zb and za >= zb * 1.3:
                want = Pt(round(za * 0.4, 1))
                if (b.space_before or 0) < want:
                    b.space_before = want
                    n += 1
    return n


def readability(no, slide, key):
    """가독성 점검(AI 없이) — ① 바탕과 글자의 명도 대비: 3:1 아래면 흰 글자/진한 글자 중 대비가 큰 쪽으로 바꾼다
    ② 글자 크기·굵기 강약: 본문 글이 거의 한 크기·한 굵기면 알린다 ③ 잘라 넣은 그림의 해상도: 150dpi 아래면 알린다. 반환: 고친 글 수"""
    fixed, sizes, levels, low = 0, [], set(), []
    for s in walk(slide.shapes):
        if is_pic(s):
            try:
                px = s.image.size[0]
                crop = 1 - (s.crop_left or 0) - (s.crop_right or 0)
                dpi = px * crop / max(0.05, s.width / IN)
                if dpi < 150 and s.width > int(0.8 * IN):
                    low.append(f"{s.width / IN:.1f}in 폭 그림 {dpi:.0f}dpi")
            except Exception:
                pass
            continue
        if s is key or not getattr(s, "has_text_frame", False) or not s.text_frame.text.strip():
            continue
        bg = None
        try:
            if s.fill.type == 1:   # 단색 채우기 — 글자의 바탕
                bg = _rgb(s.fill.fore_color)
        except Exception:
            pass
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                if not r.text.strip():
                    continue
                if r.font.size:
                    sizes.append(round(r.font.size.pt * 2) / 2)
                lat = r._r.find(f"{A}rPr/{A}latin")
                if lat is not None:
                    levels.add(fontset.LEVEL.get(lat.get("typeface"), -1))
                fg = _rgb(r.font.color)
                if bg and fg and _contrast(fg, bg) < 3.0:
                    best = max(("FFFFFF", "1F2937"), key=lambda c: _contrast(c, bg))
                    r.font.color.rgb = RGBColor.from_string(best)
                    fixed += 1
    if len(sizes) >= 8:
        big, small = max(sizes), min(sizes)
        top = max(set(sizes), key=sizes.count)
        if big < small * 1.3 or sizes.count(top) / len(sizes) > 0.85:
            warn(no, "글자 크기 강약이 약함", f"본문 글이 {small:g}~{big:g}pt 로 거의 한 크기입니다 — 제목·키워드는 크고 굵게, 설명은 작게 차이를 주세요")
        if len(levels - {-1}) <= 1:
            warn(no, "글자 굵기 강약이 없음", "모든 글이 같은 굵기입니다 — 제목·강조어는 굵은 서체(3·4), 설명은 본문 서체(2)")
    if low:
        warn(no, "저화질 그림", ", ".join(low[:3]) + " — 인쇄·큰 화면에서 흐려 보입니다. 시안을 다시 만들거나 고화질 자산으로 바꿔 주세요")
    if fixed:
        warn(no, "글자·바탕 명도 대비를 고침", f"바탕과 대비가 3:1 보다 낮은 글 {fixed}곳을 흰 글자 또는 진한 글자로 바꿨습니다")
    return fixed


_TNUM = re.compile(r"^[\s+\-−~약]*[\d][\d,.]*\s*(?:%|퍼센트|억|조|만|천|원|명|건|개|곳|종|개월|일|시간|배|점|회|MB|GB|TB)*(?:\s*[~\-]\s*[\d][\d,.]*\s*\S{0,3})?$")


def align_tables(slide):
    """표 정렬(장표규칙 20) — 머리 줄 가운데 · 길이가 들쭉날쭉한 문장 열 왼쪽 · 자릿수가 다른 비교 수치 열 오른쪽 · 그 밖(짧고 고른 값) 가운데.
    열 단위로 한 번만 정한다(한 열 안에서 정렬이 섞이지 않게)."""
    n = 0
    for sh in walk(slide.shapes):
        if not (getattr(sh, "has_table", False) and sh.has_table):
            continue
        rows = list(sh.table.rows)
        if len(rows) < 2:
            continue
        ncol = len(sh.table.columns)
        for ci in range(ncol):
            body = []
            for row in rows[1:]:
                try:
                    t = row.cells[ci].text.strip()
                except IndexError:
                    t = ""
                if t:
                    body.append(t)
            if not body:
                want = PP_ALIGN.CENTER
            elif all(_TNUM.match(t) for t in body) and len({len(re.sub(r"\D", "", t)) for t in body}) > 1:
                want = PP_ALIGN.RIGHT   # 자릿수가 다른 수치 — 크기 비교
            elif max(len(t) for t in body) > 15 or sum(len(t) for t in body) / len(body) > 10:
                want = PP_ALIGN.LEFT    # 길이가 들쭉날쭉한 문장
            else:
                want = PP_ALIGN.CENTER  # 짧고 고른 값(코드·날짜·상태·고정 자릿수)
            for ri, row in enumerate(rows):
                try:
                    cell = row.cells[ci]
                except IndexError:
                    continue
                a = PP_ALIGN.CENTER if ri == 0 else want
                for p in cell.text_frame.paragraphs:
                    if p.alignment != a:
                        p.alignment = a
                        n += 1
    return n


def _hsl(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    mx, mn = max(r, g, b), min(r, g, b)
    l = (mx + mn) / 2
    if mx == mn:
        return 0.0, 0.0, l
    d = mx - mn
    s = d / (2 - mx - mn) if l > 0.5 else d / (mx + mn)
    hue = ((g - b) / d % 6) if mx == r else ((b - r) / d + 2) if mx == g else ((r - g) / d + 4)
    return hue * 60, s, l


def _blue_step(h):
    """파랑이면 단계(0 진한 남색 · 1 주 파랑 · 2 중간 · 3 옅은 바탕), 아니면 None"""
    hue, sat, l = _hsl(h)
    if not (200 <= hue <= 245 and sat >= 0.25):
        return None
    return 0 if l < 0.30 else 1 if l < 0.55 else 2 if l < 0.85 else 3


def snap_blues(prs):
    """파랑을 4단계로(장표규칙 9-1) — 덱 안에서 단계마다 가장 많이 쓴 파랑 하나로 모은다. 강조색·그림은 건드리지 않는다.
    디자인 타입마다 팔레트가 달라 고정값이 아니라 그 덱의 대표색을 쓴다."""
    acc = ACCENT.lstrip("#").upper()
    els, cnt = [], {}
    for slide in prs.slides:
        for el in slide._element.iter(A + "srgbClr"):
            v = (el.get("val") or "").upper()
            if len(v) != 6 or v == acc:
                continue
            st = _blue_step(v)
            if st is None:
                continue
            els.append((el, st))
            cnt.setdefault(st, {}).setdefault(v, 0)
            cnt[st][v] += 1
    rep = {st: max(c, key=c.get) for st, c in cnt.items()}
    n = 0
    for el, st in els:
        if el.get("val", "").upper() != rep[st]:
            el.set("val", rep[st])
            n += 1
    return n


def is_pic(s):
    return s.shape_type == 13


def fit_axis(body, lo, hi, a0, a1, pos, size, limit):
    """본문 묶음 [a0,a1] 을 [lo,hi] 에 맞춘다. 그림도 같은 비율로 — 따로 두면 그림 위에 얹은 라벨·글과 어긋난다(7% 안쪽만 허용). 비율이 limit 밖이면 False"""
    k = (hi - lo) / (a1 - a0)
    if not (limit[0] <= k <= limit[1]):
        return False
    for s in body:
        p, z = getattr(s, pos), getattr(s, size)
        setattr(s, pos, int(lo + (p - a0) * k))
        setattr(s, size, int(z * k))
        if getattr(s, "has_table", False) and s.has_table:   # 표는 틀만 바꾸면 안 변한다 — 열·행을 같이
            for c in (s.table.columns if pos == "left" else s.table.rows):
                if pos == "left":
                    c.width = int(c.width * k)
                else:
                    c.height = int(c.height * k)
    return True


prs = Presentation(src)
nkey = nfix = nfit = nswap = nemph = ncon = nicon = nshrink = ngap = naspect = ntable = nblue = 0
deck = []
if "--unify-only" in sys.argv:   # 표지·목차·간지처럼 키메시지·가이드가 없는 장 — 서체 한 벌만 맞춘다
    for part in list(prs.slide_masters) + list(prs.slide_layouts) + list(prs.slides):
        nswap += unify_fonts(part._element)
    nbig = 0
    fkind = sys.argv[sys.argv.index("--front") + 1] if "--front" in sys.argv else ""
    for slide in prs.slides:   # 표지 사업명·목차·간지의 큰 글자(20pt 이상)는 키메시지와 같은 제목 서체
        runs = [(r, r.find(A + "rPr")) for r in slide._element.iter(A + "r")]
        top = max([int(rp.get("sz") or 0) for _, rp in runs if rp is not None] or [0])
        if fkind == "cover":   # 표지 — 사업명(가장 큰 글자)만 제목 서체. 슬로건·'제안설명회'·날짜 같은 작은 글은 본문 서체 3단계
            for r, rPr in runs:
                if rPr is None or int(rPr.get("sz") or 0) >= max(2000, top * 0.7):
                    continue
                for tag in ("latin", "ea", "cs"):
                    el = rPr.find(A + tag)
                    if el is not None:
                        el.set("typeface", fontset.SEMI)
        for r, rPr in runs:
            if rPr is not None and int(rPr.get("sz") or 0) >= (max(2000, top * 0.7) if fkind == "cover" else 2000):
                for tag in ("latin", "ea", "cs"):
                    el = rPr.find(A + tag)
                    if el is not None:
                        el.set("typeface", KEY_FONT)
                if fontset.KEY_BOLD:
                    rPr.set("b", "1")   # 가족 이름 + 굵게로 고르는 제목 서체
                if fontset.KEY_SCALE != 1.0:   # 서체마다 덩어리감 맞춤(앞장은 이 단계를 한 번만 거친다)
                    rPr.set("sz", str(round(int(rPr.get("sz") or 0) * fontset.KEY_SCALE / 50) * 50))
                # 큰 글자는 탄탄하게 — 사용자가 표지 사업명(35.5pt)에 맞춘 값: 문자 간격 좁게 0.6pt, 줄 간격 고정 42pt.
                # 다른 크기에는 같은 비율로 — 자간 = 크기 × −1.7%, 줄 간격 = 크기 × 1.18 (pt 고정)
                z = int(rPr.get("sz") or 0) / 100
                rPr.set("spc", str(-round(z * 1.7) + round(fontset.TITLE_TRACK * z * 100)))   # 1/100 pt — 제목 서체 자간을 조금 더(2026-10-14)
                pPr = r.getparent().find(A + "pPr")
                if pPr is None:
                    pPr = r.getparent().makeelement(A + "pPr", {})
                    r.getparent().insert(0, pPr)
                for old in pPr.findall(A + "lnSpc"):
                    pPr.remove(old)
                ln = pPr.makeelement(A + "lnSpc", {})
                ln.append(ln.makeelement(A + "spcPts", {"val": str(round(z * 1.18 * 100))}))
                pPr.insert(0, ln)
                nbig += 1
    if fkind == "cover":   # 사업명 아래 '제안설명회'·날짜는 사업명 글 끝에서 1cm 아래부터(서로 간격은 그대로)
        for slide in prs.slides:
            tx = [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
            def big(sh):
                return max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
            if not tx:
                continue
            title = max(tx, key=big)
            zt = big(title)
            lines = sum(1 for p in title.text_frame.paragraphs if p.text.strip()) or 1
            tf = title.text_frame
            text_h = int(lines * zt * 1.18 * 12700)   # 줄 간격 고정(크기 × 1.18pt)
            if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
                t_bot = title.top + (title.height + text_h) // 2
            elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
                t_bot = title.top + title.height - (tf.margin_bottom or 0)
            else:
                t_bot = title.top + (tf.margin_top or 0) + text_h
            below = sorted([sh for sh in tx if sh is not title and sh.top >= title.top + title.height // 3 and big(sh) < zt * 0.8
                            and sh.left < title.left + title.width and sh.left + sh.width > title.left], key=lambda sh: sh.top)
            if below:
                first = below[0]
                ftf = first.text_frame
                f_txt_top = first.top + ((first.height - int(big(first) * 1.2 * 12700)) // 2 if ftf.vertical_anchor == MSO_ANCHOR.MIDDLE else (ftf.margin_top or 0))
                d = (t_bot + 360000) - f_txt_top   # 1cm = 360000 EMU
                for sh in below:
                    sh.top = sh.top + d
    try:   # 표지·목차·간지에도 섞인 언어 글(그 글자가 없는 서체) — 그 런만 그 글자가 있는 서체로(fontbundle.fallback). 서체 모으기는 본문 마무리가 한다
        import fontbundle
        for part in list(prs.slides):
            fontbundle.fallback(part._element)
    except Exception as e:
        print("다른 언어 글 서체 검사 실패(무시):", str(e)[:120])
    prs.save(src)
    print(f"서체 {fontset.NAME} 로 {nswap}곳 · 큰 글자 {nbig}곳 {KEY_FONT}")
    sys.exit(0)
if "--polish" in sys.argv:   # 초안 다듬기 — 배치는 그대로 두고 서체 통일 · 큰 글자 제목 서체 · 가독성 점검만(키메시지·가이드 맞춤 없음)
    for part in list(prs.slide_masters) + list(prs.slide_layouts) + list(prs.slides):
        nswap += unify_fonts(part._element)
    for no, slide in enumerate(prs.slides, 1):
        for r in slide._element.iter(A + "r"):   # 20pt 이상 큰 글자(장 제목·표지 제목)는 제목 서체
            rPr = r.find(A + "rPr")
            if rPr is not None and int(rPr.get("sz") or 0) >= 2000:
                fresh = all((rPr.find(A + tag) is None or rPr.find(A + tag).get("typeface") != KEY_FONT) for tag in ("latin", "ea"))
                for tag in ("latin", "ea", "cs"):
                    el = rPr.find(A + tag)
                    if el is not None:
                        el.set("typeface", KEY_FONT)
                if fresh and fontset.KEY_SCALE != 1.0:   # 덩어리감 맞춤 — 처음 바뀌는 글자만(다시 돌거나 앞장이면 곱하지 않는다)
                    rPr.set("sz", str(round(int(rPr.get("sz")) * fontset.KEY_SCALE / 50) * 50))
                if fontset.KEY_BOLD:
                    rPr.set("b", "1")
                rPr.set("spc", str(min(int(rPr.get("spc") or 0), round(fontset.TITLE_TRACK * int(rPr.get("sz"))))))   # 제목 서체 자간 조금 더 좁게 — 다시 돌아도 같은 값(더 좁아지지 않게)
                nkey += 1
        ncon += readability(no, slide, None)
        lines = []
        for sh in walk(slide.shapes):
            if getattr(sh, "has_table", False) and sh.has_table:
                lines += [" | ".join(c.text.strip() for c in row.cells) for row in sh.table.rows]
            elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                lines.append(sh.text_frame.text.strip())
        deck.append({"n": no, "text": chr(10).join(lines)[:2500]})
    prs.save(src)
    if TEXT_OUT:
        json.dump(deck, open(TEXT_OUT, "w", encoding="utf-8"), ensure_ascii=False)
elif not FONTS_ONLY:
    for part in list(prs.slide_masters) + list(prs.slide_layouts) + list(prs.slides):   # 서체 한 벌로 — 글·표·끝 문단 속성까지 전부
        nswap += unify_fonts(part._element)
    for no, slide in enumerate(prs.slides, 1):
        top = list(slide.shapes)
        # ── 키메시지 찾기: 이름 keymsg, 없으면 본문 맨 위쪽의 가장 큰 글(18pt↑, 넓은 상자)
        key = next((s for s in top if (s.name or "").lower().startswith("keymsg")), None)
        if key is None:
            cand = [s for s in top if getattr(s, "has_text_frame", False) and s.text_frame.text.strip() and int(0.72 * IN) <= s.top <= int(1.9 * IN)
                    and max_pt(s) >= 15 and len(s.text_frame.text.strip()) >= 6 and s.width >= prs.slide_width * 0.4]
            if cand:
                key = max(cand, key=max_pt)
                key.name = "keymsg"
        if key is not None:
            for p in key.text_frame.paragraphs:
                p.alignment = KEY_ALIGN
                for r in p.runs:
                    r.font.size = Pt(KEY_PT_EFF)
                    r.font.bold = None
                    set_font(r, KEY_FONT)
                    if fontset.KEY_BOLD:
                        r.font.bold = True
                    r._r.get_or_add_rPr().set("spc", str(round(fontset.TITLE_TRACK * KEY_PT_EFF * 100)))   # 키메시지도 제목 서체 자간
            if color_keywords(key.text_frame, keywords(no)):
                nemph += 1
            pt = fit_pt(key.text_frame.text.strip().replace("\n", " "), KEY_W - 0.2)
            if pt < KEY_PT_EFF:   # 한 줄에 안 들어가는 긴 키메시지 — 넘치지 않게 줄이고 알린다
                for p in key.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(pt)
                warn(no, "키메시지가 길어 글자를 줄임", f"{KEY_PT_EFF}pt 로는 한 줄에 안 들어가 {pt}pt 로 줄였습니다 — 문구를 줄이면 {KEY_PT_EFF}pt 로 맞춰집니다")
            key_top0 = key.top
            key.left, key.top, key.width, key.height = int(KEY_L * IN), int(KEY_TOP * IN), int(KEY_W * IN), int(KEY_H * IN)
            if KEY_ALIGN == PP_ALIGN.LEFT:   # 왼쪽 정렬 — 글의 왼쪽 끝이 본문 가이드(G_L)와 같은 세로선에 오게(상자 안쪽 여백 0)
                key.left, key.width = int(G_L * IN), int((G_R - G_L) * IN)
                key.text_frame.margin_left = 0
            key.text_frame.word_wrap = False
            nkey += 1
        else:
            warn(no, "키메시지를 찾지 못함", f"맨 위 헤드라인이 없어 {KEY_PT}pt·자리 통일을 하지 못했습니다")
        # ── 본문 묶음을 가이드에 맞춘다
        body = [s for s in top if s is not key and s.top is not None and s.width and s.top >= int(0.72 * IN) and s.top + s.height <= int(7.03 * IN) and s.width < prs.slide_width * 0.99]
        if body:
            t0, b0 = min(s.top for s in body), max(s.top + s.height for s in body)
            T, B = int(BODY_TOP * IN), int(BODY_BOT * IN)
            if key is not None and abs(t0 - T) > int(0.02 * IN) or b0 > B + int(0.03 * IN):
                dy = T - t0
                if b0 + dy <= B + int(0.03 * IN) and abs(dy) <= int(0.6 * IN):   # 통째로 옮기면 되는 경우
                    for s in body:
                        s.top = s.top + dy
                    nfit += 1
                elif fit_axis(body, T, B, t0, b0, "top", "height", (0.93, 1.0)):   # 조금 눌러 담는다
                    nfit += 1
                else:
                    if key is not None and t0 < int((KEY_TOP + KEY_H) * IN):
                        key.top = key_top0   # 본문을 못 내리면 키메시지가 겹친다 → 원래 자리로
                    warn(no, "본문 위·아래 여백이 기준과 다름", f"본문 {t0 / IN:.2f}~{b0 / IN:.2f}in (기준 {BODY_TOP}~{BODY_BOT}in) — 차이가 커서 자동으로 맞추지 않았습니다")
            l0, r0 = min(s.left for s in body), max(s.left + s.width for s in body)
            L, R = int(G_L * IN), int(G_R * IN)
            if abs(l0 - L) > int(0.02 * IN) or abs(r0 - R) > int(0.02 * IN):
                if fit_axis(body, L, R, l0, r0, "left", "width", (0.93, 1.07)):
                    nfit += 1
                else:
                    warn(no, "좌우 끝이 가이드와 다름", f"본문 {l0 / IN:.2f}~{r0 / IN:.2f}in (가이드 {G_L}~{G_R}in) — 차이가 커서 자동으로 맞추지 않았습니다")
        # ── 아이콘 + 옆 글: 세로 가운데를 맞춘다(작은 그림 오른쪽에 붙은 한 줄 글 — 글상자를 그림 높이에 맞추고 가운데 정렬)
        pics = [s for s in walk(slide.shapes) if s.shape_type == 13 and s.height and s.height <= int(0.75 * IN) and s.width <= int(0.9 * IN)]
        txts = [s for s in walk(slide.shapes) if getattr(s, "has_text_frame", False) and s is not key and s.shape_type != 13 and s.text_frame.text.strip()
                and chr(10) not in s.text_frame.text.strip() and len(s.text_frame.paragraphs) == 1 and s.height and s.height <= int(0.9 * IN)]
        for pc in pics:
            cy = pc.top + pc.height / 2
            best = None
            for t in txts:
                gap = t.left - (pc.left + pc.width)
                if -int(0.05 * IN) <= gap <= int(0.3 * IN) and t.top < pc.top + pc.height and t.top + t.height > pc.top and t.height <= pc.height * 2.2:
                    if best is None or gap < best[0]:
                        best = (gap, t)
            if best:
                t = best[1]
                # 아이콘 옆에 '제목 + 설명'이 위아래로 쌓인 경우(300일 / 약 10개월…)는 손대지 않는다 —
                # 제목을 아이콘 가운데로 내리면 바로 아래 설명 글과 겹친다. 옮긴 자리가 다른 글과 겹쳐도 그만둔다
                ny = int(cy - t.height / 2)
                stacked = any(o is not t and o.left < t.left + t.width and o.left + o.width > t.left
                              and o.top < pc.top + pc.height + int(0.1 * IN) and o.top + o.height > pc.top - int(0.1 * IN) for o in txts + [x for x in walk(slide.shapes) if getattr(x, "has_text_frame", False) and x is not t and x is not key and x.shape_type != 13 and x.text_frame.text.strip()])
                hits = any(o is not t and o is not key and getattr(o, "has_text_frame", False) and o.shape_type != 13 and o.text_frame.text.strip()
                           and o.left < t.left + t.width and o.left + o.width > t.left and o.top < ny + t.height and o.top + o.height > ny
                           and not (o.top < t.top + t.height and o.top + o.height > t.top) for o in walk(slide.shapes))
                if stacked or hits:
                    continue
                if abs((t.top + t.height / 2) - cy) > int(0.02 * IN) or t.text_frame.vertical_anchor != MSO_ANCHOR.MIDDLE:
                    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    t.text_frame.margin_top = t.text_frame.margin_bottom = 0
                    t.top = int(cy - t.height / 2)
                    nicon += 1
        # ── 도형 안 글
        lines = []
        for s in walk(slide.shapes):
            if getattr(s, "has_table", False) and s.has_table:
                for row in s.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
                continue
            if not getattr(s, "has_text_frame", False) or not s.text_frame.text.strip():
                continue
            tf = s.text_frame
            lines.append(tf.text.strip())
            if s is key:
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                continue
            paras = [p for p in tf.paragraphs if p.text.strip()]
            try:
                filled = s.fill.type is not None and s.fill.type != 5   # 5 = 채우기 없음
            except Exception:
                filled = False
            if len(paras) == 1:
                p = paras[0]
                ls = p.line_spacing
                if ls is not None and not isinstance(ls, float) and s.height <= ls * 2.2:
                    # 한 줄 글에 걸린 고정 pt 줄간격이 글을 위로 민다 → 100% + 세로 가운데(상자 자리는 그대로)
                    p.line_spacing = 1.0
                    p.space_before = p.space_after = Pt(0)
                    tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    tf.word_wrap = False   # 한 줄 상자는 줄바꿈하지 않는다(서체가 바뀌어도 두 줄로 안 깨지게)
                    nfix += 1
                    continue
            if filled and len(paras) <= 2 and s.height <= int(0.75 * IN):   # 낮은 도형만(큰 카드의 위쪽 글은 건드리지 않는다)
                tf.margin_top = tf.margin_bottom = min(tf.margin_top, tf.margin_bottom)
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                nfix += 1
        naspect += keep_aspect(slide)
        nshrink += fit_one_line(no, slide, key)
        ngap += title_gap(slide, key)
        ncon += readability(no, slide, key)
        ntable += align_tables(slide)
        deck.append({"n": no, "text": "\n".join(lines)[:2500]})
    nblue = snap_blues(prs)
    prs.save(src)
    if TEXT_OUT:
        json.dump(deck, open(TEXT_OUT, "w", encoding="utf-8"), ensure_ascii=False)

# ── 섞인 언어 — 글 런마다 서체에 없는 글자가 있으면 그 글자를 다 가진 같은 굵기 서체로(fontbundle.fallback, 2026-10-18)
import fontbundle
try:   # 옛 장표·템플릿에서 따라온 남의 서체(유료·엔진에 없음)를 이 장표의 본문 서체로 — 산출물에 그 이름이 남지 않게(R97)
    _sz = fontbundle.sanitize_pptx(src)
    if _sz:
        print(fontbundle.sanitize_note(_sz))
except Exception as e:
    warn(0, "남의 서체 정리 실패", str(e)[:160])
try:
    _fb = fontbundle.fallback_pptx(src)
    if _fb:
        print("다른 언어 글 " + str(len(_fb)) + "곳 서체 바꿈 — " + " · ".join(dict.fromkeys(f"{a} → {b}" for a, b, _ in _fb)))
except Exception as e:   # 바꾸지 못해도 장표는 그대로 — 알리기만
    warn(0, "다른 언어 글 서체 검사 실패", str(e)[:160])
# ── 쓴 서체 → 산출물 '5_서체'(pptx 에 실제로 적힌 a:latin·a:ea 서체를 모은다 — fontbundle.bundle, 라이선스·설치 안내 포함)
_paths, miss, noship = fontbundle.bundle(src, fonts_dir)
got = {os.path.basename(x) for x in _paths if re.search(r"\.(ttf|otf|ttc)$", x, re.I)}
if noship:
    warn(0, "산출물에 넣지 않은 서체", ", ".join(noship) + " — 재배포가 안 되는 서체라 받는 사람이 따로 설치해야 합니다")
if miss:
    warn(0, "이 PC 에 없는 서체", ", ".join(miss) + " — 서체 파일을 산출물에 넣지 못했습니다")
if fontset.FACES == fontset.ASTA and fontset.BODY_ID == "siwol" and not FONTS_ONLY:   # 옛 목록(a시월구일)일 때만
    warn(0, "a시월구일 서체가 이 엔진 PC 에 없음", "a시월구일1~4 대신 같은 굵기의 Asta Sans 로 만들었습니다 — 엔진 PC 에 a시월구일을 설치하면 다음 제작부터 바뀝니다")
for w in warns:
    print(w)
if "--polish" in sys.argv:
    print(f"큰 글자 {nkey}곳 제목 서체 {KEY_FONT} · 서체 {fontset.NAME} 로 {nswap}곳 · 명도 대비 고침 {ncon}곳 · 서체 파일 {len(got)}개")
else:
    print(f"키메시지 {nkey}장 통일(강조어 색 {nemph}장) · 가이드 맞춤 {nfit}건 · 아이콘 옆 글 {nicon}곳 가운데 · 넘친 한 줄 글 {nshrink}곳 줄임 · 그림 비율 되돌림 {naspect}개 · 표 정렬 {ntable}곳 · 파랑 4단계로 {nblue}곳 ·제목·설명 간격 {ngap}곳 · 도형 글 {nfix}개 정리 · 서체 {fontset.NAME} 로 {nswap}곳 · 서체 파일 {len(got)}개")
