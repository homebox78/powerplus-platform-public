# -*- coding: utf-8 -*-
"""고정 틀(머리 띠·사업명·바닥 선·로고·쪽번호)은 본문 장의 레이아웃(슬라이드 마스터 쪽)에 한 번, 장 이름·절 제목만 장마다 — — 엔진이 그리지 않고 코드가 그려 항상 같은 자리·크기.
  frame.py <pptx> <plan.json> <opts.json>
opts: {project, client_logo, vendor_logo}. 도형 이름은 pd_frame_… — 다시 돌리면 지우고 새로 넣는다.
시안의 틀(design.compose_frame)과 같은 구성: 왼쪽 작은 장 이름 + 굵은 절 제목, 오른쪽 사업명 / 바닥 선·로고·쪽번호."""
import json, os, re, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
import fontset

sys.stdout.reconfigure(encoding="utf-8")
IN = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
HEAD, FOOT, PAD, LOGO_H = 0.70, 0.45, 0.25, 0.19   # 바닥 로고는 작게(10-16 사용자) — 0.24 에서 줄임


def rgb(h):
    return RGBColor.from_string(h)


def text(slide, x, y, w, h, t, size, color, font, align=PP_ALIGN.LEFT, name=""):
    tb = slide.shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = False, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    r.font.size, r.font.color.rgb = Pt(size), rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        rPr.append(rPr.makeelement(A + tag, {"typeface": font}))
    tb.name = "pd_frame_" + name
    return tb


def to_back(slide, shapes):
    """틀은 본문 뒤로 — 본문 도형이 가려지지 않게"""
    tree = slide.shapes._spTree
    for i, sh in enumerate(shapes):
        tree.remove(sh._element)
        tree.insert(2 + i, sh._element)


if sys.argv[1] == "strip":   # 틀만 떼어 낸다: strip <pptx>
    prs = Presentation(sys.argv[2])
    for slide in prs.slides:
        for sh in [sh for sh in slide.shapes if (sh.name or "").startswith("pd_frame")]:
            sh._element.getparent().remove(sh._element)
    for lay in prs.slide_layouts:   # 레이아웃(마스터 쪽)에 넣은 틀도
        for sh in [sh for sh in lay.shapes if (sh.name or "").startswith("pd_frame")]:
            sh._element.getparent().remove(sh._element)
    prs.save(sys.argv[2])
    sys.exit(0)
prs = Presentation(sys.argv[1])
plan = json.load(open(sys.argv[2], encoding="utf-8"))
o = json.load(open(sys.argv[3], encoding="utf-8"))
W, H = prs.slide_width / IN, prs.slide_height / IN
slides = plan.get("slides") or []
project = str(o.get("project") or "").strip()
FR = o.get("frame") or {}   # 디자인 타입의 머리 띠·바닥(reference/styles/styles.json, '{accent}' 는 워커가 바꿔 넘긴다)
FILL = FR.get("fill") or ["1D4ED8", "3B82F6"]
body = [s for s in prs.slides if not any((sh.name or "").startswith("pd_front") for sh in s.shapes)]
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def build_common(slide):
    """모든 본문 장에 같은 틀 — 임시 장에 그린 뒤 레이아웃으로 옮긴다"""
    made = []
    if FR.get("head", "band") == "band":   # 꽉 찬 띠(가로 그라데이션)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Emu(int(HEAD * IN)))
        band.line.fill.background()
        band.shadow.inherit = False
        band.fill.gradient()
        band.fill.gradient_angle = 0
        st = band.fill.gradient_stops
        st[0].color.rgb, st[1].color.rgb = rgb(FILL[0]), rgb(FILL[-1])
        st[0].position, st[1].position = 0, 1
    else:   # 'line' — 흰 바탕 + 머리 아래 가는 색 선
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(int((HEAD - 0.03) * IN)), prs.slide_width, Emu(int(0.03 * IN)))
        band.line.fill.background()
        band.shadow.inherit = False
        band.fill.solid()
        band.fill.fore_color.rgb = rgb(FILL[0])
    band.name = "pd_frame_band"
    made.append(band)
    if project:
        made.append(text(slide, W - PAD - 4.2, 0.22, 4.2, 0.26, project, 10.5, FR.get("project", "DBEAFE"), fontset.BODY, PP_ALIGN.RIGHT, name="project"))
    y0 = H - FOOT
    # 바닥 선 — 연결선은 테마 그림자가 붙어 굵게 보인다. 그림자 없는 얇은 사각형으로
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(PAD * IN)), Emu(int(y0 * IN)), Emu(int((W - 2 * PAD) * IN)), Emu(9525))
    line.line.fill.background()
    line.shadow.inherit = False
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(FR.get("foot_line", "E5E7EB"))
    line.name = "pd_frame_line"
    made.append(line)
    pics = {}
    lim_h, lim_w = int((FOOT - 0.06) * IN), int(2.3 * IN)

    def clamp(pic):
        if pic.height > lim_h:
            k = lim_h / pic.height
            pic.width, pic.height = int(pic.width * k), lim_h
        if pic.width > lim_w:   # 가로로 긴 로고는 폭을 막는다
            k = lim_w / pic.width
            pic.width, pic.height = lim_w, int(pic.height * k)

    for which in ("client_logo", "vendor_logo"):
        p = o.get(which)
        if p and os.path.isfile(p):
            try:
                pic = slide.shapes.add_picture(p, 0, Emu(int((y0 + (FOOT - LOGO_H) / 2) * IN)), height=Emu(int(LOGO_H * IN)))
                # 면적으로 맞춘다(가로형 5:1 로고면 높이 LOGO_H) — 발주기관 로고는 제안사 로고보다 15% 크게
                side = LOGO_H * IN * 5 ** 0.5 * (1.15 if which == "client_logo" else 1.0)
                k = side / (pic.width * pic.height) ** 0.5
                pic.width, pic.height = int(pic.width * k), int(pic.height * k)
                clamp(pic)
                pic.name = "pd_frame_" + which
                pics[which] = pic
            except Exception:
                pass
    # 높이로도 15% 크게(DCK-206) — 면적만 맞추면 가로로 긴 제안사 로고가 더 커 보인다.
    # 발주기관 로고를 키울 수 있으면 키우고, 틀(높이·폭 상한)에 막히면 제안사 로고를 줄인다
    c, v = pics.get("client_logo"), pics.get("vendor_logo")
    if c is not None and v is not None and c.height < v.height * 1.15:
        want = int(v.height * 1.15)
        k = min(want / c.height, lim_h / c.height, lim_w / c.width)
        if k > 1:
            c.width, c.height = int(c.width * k), int(c.height * k)
        if c.height < v.height * 1.15:
            k = c.height / (v.height * 1.15)
            v.width, v.height = int(v.width * k), int(v.height * k)
    for which, right in (("client_logo", False), ("vendor_logo", True)):
        pic = pics.get(which)
        if pic is None:
            continue
        pic.top = Emu(int((y0 + FOOT / 2) * IN - pic.height / 2))
        pic.left = Emu(int((W - PAD) * IN) - pic.width) if right else Emu(int(PAD * IN))
        made.append(pic)
    # 쪽번호 — 장 번호 필드(장을 넣고 빼도 저절로 맞는다)
    pg = text(slide, W / 2 - 0.5, y0 + 0.1, 1.0, 0.26, "#", 9, FR.get("page", "6B7280"), fontset.BODY, PP_ALIGN.CENTER, name="page")
    r = pg.text_frame.paragraphs[0].runs[0]._r
    fld = r.makeelement(A + "fld", {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"})
    for ch in list(r):
        fld.append(ch)
    t = fld.find(A + "t")
    if t is not None:
        t.text = "‹#›"
    r.getparent().replace(r, fld)
    made.append(pg)
    return made


def to_layout(layout, tmp, shapes):
    """임시 장의 도형을 레이아웃(마스터 쪽)으로 옮긴다 — 그림은 레이아웃에 그림 관계를 새로 건다"""
    tree = layout.shapes._spTree
    for sh in [sh for sh in layout.shapes if (sh.name or "").startswith("pd_frame")]:
        tree.remove(sh._element)
    for k, sh in enumerate(shapes):
        el = sh._element
        for blip in el.iter(A + "blip"):
            rid = blip.get(R_NS + "embed")
            part = tmp.part.related_part(rid)
            blip.set(R_NS + "embed", layout.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"))
        el.getparent().remove(el)
        tree.insert(2 + k, el)


done = 0
stray = 0
layouts = []
for slide in body:
    if slide.slide_layout not in layouts:
        layouts.append(slide.slide_layout)
for lay in layouts:   # 본문 장이 쓰는 레이아웃마다 틀을 한 번
    tmp = prs.slides.add_slide(lay)
    made = build_common(tmp)
    to_layout(lay, tmp, made)
    if FR.get("bg"):   # 어두운 타입 — 장 바탕색도 레이아웃에
        lay.background.fill.solid()
        lay.background.fill.fore_color.rgb = rgb(FR["bg"])
    sid = prs.slides._sldIdLst[-1]   # 임시 장 지우기
    prs.part.drop_rel(sid.get(R_NS + "id"))
    prs.slides._sldIdLst.remove(sid)
for i, slide in enumerate(body):   # 장마다 다른 것(장 이름·절 제목)만 장 위에 — 사용자가 장마다 고치는 글
    slide._element.set("showMasterSp", "1")
    for sh in [sh for sh in slide.shapes if (sh.name or "").startswith("pd_frame")]:
        sh._element.getparent().remove(sh._element)
    # 틀은 레이아웃에만(DCK-207) — 장마다 따로 놓인 로고·쪽번호가 남아 있으면 레이아웃 틀과 겹친다.
    # 머리 띠·바닥 안에 통째로 들어간 그림(이름에 logo/로고)과 바닥의 숫자 글상자(손으로 친 쪽번호)만 걷는다
    for sh in list(slide.shapes):
        top, bot = (sh.top or 0) / IN, ((sh.top or 0) + (sh.height or 0)) / IN
        in_head, in_foot = bot <= HEAD + 0.02, top >= H - FOOT - 0.02
        nm = (sh.name or "").lower()
        if sh.shape_type == 13 and (in_head or in_foot) and ("logo" in nm or "로고" in nm or in_foot):
            sh._element.getparent().remove(sh._element)
            stray += 1
        elif in_foot and getattr(sh, "has_text_frame", False) and re.fullmatch(r"\s*[-–]?\s*\d{1,3}\s*[-–]?\s*", sh.text_frame.text or ""):
            sh._element.getparent().remove(sh._element)
            stray += 1
    s = slides[i] if i < len(slides) and len(slides) == len(body) else {}
    made = []
    sec = re.sub(r"\s+", " ", str(s.get("section") or "").replace("**", "")).strip()
    ch = re.sub(r"\s+", " ", str(s.get("chapter") or "").replace("**", "")).strip()
    if ch and sec.startswith(ch):   # 장 이름에 '·'가 있어도('기술·기능 방안') 장 이름째로 뗀다(DCK-106)
        rest = re.sub(r"^\s*[·/|>:\-–—]?\s*", "", sec[len(ch):])
        small, big = (ch, rest) if rest else ("", ch)
    else:   # 장 이름이 따로 없으면 띄어 쓴 구분자(' · ')에서만 나눈다
        parts = [p.strip() for p in re.split(r"\s+[·/|>]\s+", sec, maxsplit=1)] if sec else [""]
        small, big = (parts[0], parts[1]) if len(parts) > 1 and parts[1] else ("", parts[0])
    if small:
        made.append(text(slide, PAD, 0.09, W * 0.6, 0.2, small, 11, FR.get("small", "BFDBFE"), fontset.BODY, name="chapter"))
    if big:
        made.append(text(slide, PAD, 0.30 if small else 0.18, W * 0.62, 0.34, big, 18, FR.get("title", "FFFFFF"), fontset.BOLD, name="section"))
    to_back(slide, made)
    done += 1
prs.save(sys.argv[1])
print(f"고정 틀 {done}장" + (f" · 장마다 따로 놓인 로고·쪽번호 {stray}개를 걷어 레이아웃 틀만 남김" if stray else ""))
