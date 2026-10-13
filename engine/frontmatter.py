# -*- coding: utf-8 -*-
"""표지 · 목차 · 간지를 결과 장표에 넣는다 (AI 없이, 기획의 장·절 정보로).
  frontmatter.py <pptx> <plan.json> <opts.json> <out.pptx>
opts: {cover, toc, divider (bool), title, project, accent("#RRGGBB"), client_logo, vendor_logo, date}
넣은 장의 모든 도형 이름은 pd_front_… 로 시작한다 — 다시 돌리면 예전 것을 지우고 새로 넣는다.
본문 장의 쪽번호는 그대로다(표지·목차·간지는 쪽번호 없음). 간지는 장(章)이 둘 이상일 때만 넣는다.
"""
import copy, datetime, json, os, re, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
import fontset

sys.stdout.reconfigure(encoding="utf-8")
IN = 914400
F_KEY, F_XB, F_SB, F_M = fontset.KEY_FONT, fontset.BOLD, fontset.SEMI, fontset.BODY   # a시월구일 4·3·2 (없으면 Asta Sans)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def is_front(slide):
    return any((s.name or "").startswith("pd_front") for s in slide.shapes)


def kinds(prs):
    """장마다 'body' 또는 표지·목차·간지 종류"""
    out = []
    for s in prs.slides:
        k = next(((sh.name or "")[9:].split("_")[0] for sh in s.shapes if (sh.name or "").startswith("pd_front_")), None)
        out.append(k or "body")
    return out


def drop_fronts(prs):
    lst = prs.slides._sldIdLst
    for sld_id, slide in list(zip(list(lst), list(prs.slides))):
        if is_front(slide):
            prs.part.drop_rel(sld_id.rId)
            lst.remove(sld_id)
    prs.part.rename_slide_parts([x.rId for x in lst])   # 남은 장의 파일 이름을 1..N 으로 — 새로 넣는 장과 이름이 겹치지 않게


class Page:
    def __init__(self, prs, kind, at):
        layout = min(prs.slide_layouts, key=lambda l: len(l.shapes))
        self.s = prs.slides.add_slide(layout)
        for ph in list(self.s.placeholders):
            ph._element.getparent().remove(ph._element)
        self.s._element.set("showMasterSp", "0")   # 마스터의 머리 띠·바닥을 숨긴다
        lst = prs.slides._sldIdLst
        el = lst[-1]
        lst.remove(el)
        lst.insert(at, el)
        self.kind, self.n, self.W, self.H = kind, 0, prs.slide_width / IN, prs.slide_height / IN

    def _name(self, sh):
        self.n += 1
        sh.name = f"pd_front_{self.kind}_{self.n}"
        return sh

    def box(self, x, y, w, h, fill, fill2=None, angle=0, shape=MSO_SHAPE.RECTANGLE):
        sh = self.s.shapes.add_shape(shape, Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
        sh.line.fill.background()
        sh.shadow.inherit = False
        if fill2:
            sh.fill.gradient()
            sh.fill.gradient_angle = angle
            st = sh.fill.gradient_stops
            st[0].color.rgb, st[1].color.rgb = rgb(fill), rgb(fill2)
            st[0].position, st[1].position = 0, 1
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        return self._name(sh)

    def text(self, x, y, w, h, runs, size, color, font, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, spacing=None):
        """runs: 문자열 또는 [(글, 색)] 또는 문단 목록[[(글, 색)], …]"""
        tb = self.s.shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
        tf = tb.text_frame
        tf.word_wrap, tf.vertical_anchor = wrap, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs if isinstance(runs, list) else [(runs, color)]]
        for i, pr in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p._p.get_or_add_pPr().set("eaLnBrk", "0")   # 한글 낱말 중간에서 줄이 끊기지 않게
            if spacing:
                p.line_spacing = spacing
            for t, c in pr:
                r = p.add_run()
                r.text = t
                r.font.size, r.font.color.rgb = Pt(size), rgb(c or color)
                rPr = r._r.get_or_add_rPr()
                for tag in ("latin", "ea", "cs"):
                    el = rPr.makeelement(A + tag, {"typeface": font})
                    rPr.append(el)
        return self._name(tb)

    def logo(self, path, x, y, h, right=False):
        if not path or not os.path.isfile(path):
            return
        try:
            pic = self.s.shapes.add_picture(path, Emu(int(x * IN)), Emu(int(y * IN)), height=Emu(int(h * IN)))
            if right:
                pic.left = Emu(int(x * IN)) - pic.width
            self._name(pic)
        except Exception:
            pass


def tw(t):
    """글 폭(글자 크기 단위) 어림 — 한글·전각 1, 영문·숫자·기호 0.55"""
    return sum(1.0 if ord(c) > 0x2E80 else (0.3 if c == " " else 0.55) for c in str(t or ""))


def fit_size(t, width_in, base, lo):
    """한 줄 폭(인치)에 들어가는 가장 큰 글자 크기(pt). lo 로도 안 되면 lo — 그때는 줄을 바꾼다. 말줄임표는 쓰지 않는다."""
    need = tw(t)
    if need <= 0:
        return base
    return max(lo, min(base, int(width_in * 72 / need * 2) / 2))


def fit_items(items, width_in, base, lo, sep=" · "):
    """여러 항목을 한 줄에 — 줄이지 않고 들어가는 항목만, 남으면 '외 N개'. → (글, 크기)"""
    for size in (base, (base + lo) / 2, lo):
        cap = width_in * 72 / size
        out = []
        for i, t in enumerate(items):
            rest = len(items) - i - 1
            cand = sep.join(out + [t]) + (f"{sep}외 {rest}개" if rest else "")
            if tw(cand) <= cap:
                out.append(t)
            else:
                break
        if len(out) == len(items):
            return sep.join(out), size
        if size == lo:
            return (sep.join(out) + (f"{sep}외 {len(items) - len(out)}개" if out else f"{len(items)}개 항목")), size
    return sep.join(items), lo


def emph(t, color, accent):
    """**강조** → 강조색 런"""
    parts = re.split(r"\*\*(.+?)\*\*", t)
    return [(p, accent if i % 2 else color) for i, p in enumerate(parts) if p]


def clean(t):
    return re.sub(r"\s+", " ", str(t or "").replace("**", "")).strip()


def chapters(slides):
    """기획의 section('장 · 절')으로 장을 묶는다 → [{name, items:[(쪽, 제목)]}]"""
    out = []
    for i, s in enumerate(slides, 1):
        sec = clean(s.get("section"))
        chf = clean(s.get("chapter"))
        if chf:   # 기획이 장(章)을 따로 적었으면 그대로 — '성능·품질' 같은 장 이름이 '·'에서 깨지지 않게
            ch, sub = chf, (sec[len(chf):].lstrip(" ·").strip() if sec.startswith(chf) else ("" if sec == chf else sec))
        else:
            ch, _, sub = sec.partition("·")
            ch, sub = ch.strip(), sub.strip()
        title = sub or clean(s.get("title"))   # 자르지 않는다 — 넣는 곳에서 글 크기·항목 수로 맞춘다(DCK-203)
        if not out or out[-1]["name"] != ch:
            out.append({"name": ch, "items": []})
        out[-1]["items"].append((i, title))
    return out


def build(prs, plan, o):
    drop_fronts(prs)
    slides = plan.get("slides") or []
    body = len(prs.slides)
    accent = o.get("accent") or "#FF2A6D"
    if not re.match(r"^#[0-9A-Fa-f]{6}$", accent):
        accent = "#FF2A6D"
    chs = chapters(slides) if len(slides) == body else []   # 본문 장 수가 기획과 다르면 장 묶음을 믿을 수 없다
    named = [c for c in chs if c["name"]]
    multi = len(named) >= 2
    title = clean(o.get("title") or plan.get("title") or "")
    project = clean(o.get("project") or "")
    made = []

    # ── 간지: 장이 바뀌는 본문 장 앞에 (뒤에서부터 넣어야 앞 번호가 안 밀린다)
    if o.get("divider") and multi:
        for ci in range(len(chs) - 1, -1, -1):
            c = chs[ci]
            pg = Page(prs, "div", c["items"][0][0] - 1)
            pg.box(0, 0, pg.W, pg.H, "1E40AF", "3B82F6", 30)
            pg.text(0.9, 1.25, 4, 1.9, f"{ci + 1:02d}", 120, "60A5FA", F_KEY, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
            pg.box(0.95, 3.35, 0.9, 0.08, accent)
            pg.text(0.9, 3.6, pg.W - 1.8, 1.0, re.sub(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ\dIVX]+\s*[.．]\s*", "", c["name"]) or c["name"], 36, "FFFFFF", F_KEY, anchor=MSO_ANCHOR.TOP)
            rows = [[(f"{n:>2}   ", "93C5FD"), (t, "DBEAFE")] for n, t in c["items"][:8]]
            longest = max((f"{n:>2}   {t}" for n, t in c["items"][:8]), key=tw, default="")
            pg.text(0.95, 4.75, pg.W - 1.9, 2.2, rows, fit_size(longest, pg.W - 1.9, 14, 11), "DBEAFE", F_M, spacing=1.45)
            made.append("간지")

    # ── 목차
    off = 0
    if o.get("toc") and slides and len(slides) == body:
        pg = Page(prs, "toc", 0)
        off = 1
        pg.box(0, 0, pg.W, pg.H, "FFFFFF")
        pg.box(0, 0, 3.3, pg.H, "1D4ED8", "3B82F6", 90)
        pg.text(0.6, 2.75, 2.4, 0.9, "목차", 40, "FFFFFF", F_KEY)
        pg.box(0.62, 3.75, 0.7, 0.07, accent)
        pg.text(0.6, 3.95, 2.4, 0.4, "CONTENTS", 12, "BFDBFE", F_SB)
        x, w, top, bot = 4.0, pg.W - 4.0 - 0.7, 0.9, pg.H - 0.8
        if multi:
            rh = min(1.25, (bot - top) / len(chs))
            y = top + ((bot - top) - rh * len(chs)) / 2
            for ci, c in enumerate(chs):
                pg.text(x, y, 0.8, rh * 0.6, f"{ci + 1:02d}", 24, accent.lstrip("#"), F_KEY, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                nm = c["name"] or "본문"
                pg.text(x + 0.85, y, w - 1.9, rh * 0.6, nm, fit_size(nm, w - 1.9, 18, 13), "0B2E6B", F_KEY, anchor=MSO_ANCHOR.MIDDLE)
                a, b = c["items"][0][0], c["items"][-1][0]
                pg.text(x + w - 1.0, y, 1.0, rh * 0.6, f"{a}" if a == b else f"{a} – {b}", 13, "1D4ED8", F_SB, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, wrap=False)
                # 절 이름은 자르지 않는다(말줄임표 금지, DCK-203) — 들어가는 만큼 온전히, 남으면 '외 N개'
                sub, ssz = fit_items([t for _, t in c["items"]], w - 0.9, 10.5, 9)
                pg.text(x + 0.85, y + rh * 0.58, w - 0.9, rh * 0.34, sub, ssz, "6B7280", F_M, wrap=False)
                pg.box(x, y + rh - 0.02, w, 0.012, "E5E7EB")
                y += rh
        else:
            items = [it for c in chs for it in c["items"]][:14]
            rh = min(0.62, (bot - top) / max(1, len(items)))
            y = top + ((bot - top) - rh * len(items)) / 2
            for n, t in items:
                pg.text(x, y, 0.7, rh, f"{n:02d}", 18, accent.lstrip("#"), F_KEY, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                pg.text(x + 0.8, y, w - 0.8, rh, t, fit_size(t, w - 0.8, 15, 11), "0B2E6B", F_SB, anchor=MSO_ANCHOR.MIDDLE)
                pg.box(x, y + rh - 0.01, w, 0.012, "E5E7EB")
                y += rh
        made.append("목차")

    # ── 엔딩 · 질의응답 — 맨 뒤 세트(시안이 있으면 워커가 시안대로 만든 장으로 갈아 끼운다)
    if o.get("ending") and slides:
        for kind, big, small in (("end", "감사합니다", ""), ("qna", "질의응답", "Q&A")):
            pg = Page(prs, kind, len(prs.slides._sldIdLst))
            pg.box(0, 0, pg.W, pg.H, "1E40AF", "3B82F6", 35)
            pg.text(0.9, pg.H / 2 - 1.0, pg.W - 1.8, 1.2, big, 54, "FFFFFF", F_KEY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, wrap=False)
            if small:
                pg.text(0.9, pg.H / 2 + 0.35, pg.W - 1.8, 0.6, small, 20, "BFDBFE", F_SB, PP_ALIGN.CENTER, MSO_ANCHOR.TOP, wrap=False)
            made.append("엔딩" if kind == "end" else "질의응답")

    # ── 표지
    if o.get("cover"):
        pg = Page(prs, "cover", 0)
        pg.box(0, 0, pg.W, pg.H, "1E40AF", "3B82F6", 35)
        pg.box(0, pg.H - 1.15, pg.W, 1.15, "FFFFFF")
        pg.box(0.92, 2.3, 0.9, 0.08, accent)
        main, _, tail = (project or title).partition(" — ")   # 표지 큰 글 = 사업명(없으면 작업 이름). '제목 — 부제' 는 부제를 작게 한 줄로
        if not tail:
            main, _, tail = title.partition(" - ")
        cap = (pg.W - 1.8) * 72   # 한 줄 폭(pt). 한글 한 자 ≈ 글자 크기
        size = max(26, min(40, int(cap / max(1, len(main)))))
        if len(main) * size > cap:   # 26pt 로도 한 줄이 안 되면 두 줄
            size = max(24, min(34, int(cap * 2 * 0.9 / len(main))))
        lines = 1 if len(main) * size <= cap else 2
        pg.text(0.9, 2.6, pg.W - 1.8, 0.75 * lines + 0.2, main, size, "FFFFFF", F_KEY, spacing=1.2)
        if tail:
            pg.text(0.9, 2.6 + size / 72 * 1.35 * lines + 0.25, pg.W - 1.8, 0.5, tail, 20, "BFDBFE", F_SB, wrap=False)
        pg.text(0.9, pg.H - 1.75, 4, 0.35, o.get("date") or datetime.date.today().strftime("%Y. %m."), 13, "DBEAFE", F_M, wrap=False)
        pg.logo(o.get("client_logo"), 0.9, pg.H - 0.8, 0.45)
        pg.logo(o.get("vendor_logo"), pg.W - 0.9, pg.H - 0.8, 0.45, right=True)
        made.insert(0, "표지")
    return made


if __name__ == "__main__":
    if sys.argv[1] == "drop":   # 표지·목차·간지를 떼고 본문만 남긴다: drop <in> <out>
        prs = Presentation(sys.argv[2])
        drop_fronts(prs)
        prs.save(sys.argv[3])
        sys.exit(0)
    if sys.argv[1] == "kinds":   # 장마다 종류 — 워커가 렌더를 본문 쪽번호에 맞춰 올릴 때 쓴다
        print(json.dumps(kinds(Presentation(sys.argv[2]))))
        sys.exit(0)
    if sys.argv[1] == "tag":   # 시안대로 만든 앞장을 갈아 끼운 뒤 그 장의 도형 이름을 pd_front_<종류>_N 으로: tag <pptx> <장번호>=<종류> …
        prs = Presentation(sys.argv[2])
        for a in sys.argv[3:]:
            i, k = a.split("=", 1)
            s = prs.slides[int(i) - 1]
            s._element.set("showMasterSp", "0")
            shapes = list(s.shapes)
            if not shapes:   # 빈 장이면 표시용 점 하나 — 종류를 알아볼 수 있게
                shapes = [s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(9525), Emu(9525))]
                shapes[0].fill.background()
                shapes[0].line.fill.background()
            for n, sh in enumerate(shapes, 1):
                sh.name = f"pd_front_{k}_{n}"
        prs.save(sys.argv[2])
        sys.exit(0)
    src, plan_p, opt_p, dst = sys.argv[1:5]
    prs = Presentation(src)
    plan = json.load(open(plan_p, encoding="utf-8"))
    made = build(prs, plan, json.load(open(opt_p, encoding="utf-8")))
    prs.save(dst)
    print("넣은 장: " + (", ".join(f"{k} {made.count(k)}" for k in dict.fromkeys(made)) if made else "없음"))
