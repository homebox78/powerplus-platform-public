# -*- coding: utf-8 -*-
"""merge.py 의 PowerPoint 없는 판(리눅스 엔진) — python-pptx 로 한 장짜리 pptx 들을 합친다. 인자는 merge.py 와 같다.
  merge_pptx.py <out.pptx> <s01.pptx> <s02.pptx> …            순서대로 한 파일로
  merge_pptx.py replace <base.pptx> <쪽번호>=<sNN.pptx> …      base 의 그 장만 갈아 끼운다
장 안의 도형·배경은 그대로 복사하고, 장이 가진 연결(그림·차트·SmartArt·동영상·OLE·태그·하이퍼링크 …)은 새 파일에 다시 넣어 rId 를 바꿔 단다.
- 그림: 같은 그림은 한 번만(get_or_add_image_part)
- 그 밖의 내부 연결: 파트를 통째로, 하위 연결까지 재귀로 복사 — 이름은 새 파일에서 비어 있는 번호(chart1.xml 이 두 장에 다 있어도 chart2.xml)
- 메모(notesSlide)·검토 의견(comments)·다른 장으로 가는 링크는 버린다(새 파일에서 가리킬 곳이 없다)
- 끝내 못 옮긴 연결은 그 도형만 빼고 'WARN: 쪽|제목|설명' 줄을 남긴다 — 합치기 전체를 멈추지 않는다(R79)
레이아웃은 base 의 같은 이름(없으면 첫 레이아웃).
같은 11장을 PowerPoint 합치기와 비교해 PowerPoint 렌더 픽셀이 같았다(2026-10-05).
2026-10-18 차트 든 장에서 '지원하지 않는 연결: …/chart' 로 sys.exit 해 작업 55ae9b33 이 통째로 실패한 뒤 위처럼 바꿨다(R79)."""
import copy, io, os, re, shutil, sys
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
# 새 파일로 옮기지 않는 연결 — 메모·의견은 장 밖(마스터·작성자 목록)을 다시 가리켜 옮기면 깨진다. 다른 장 링크는 갈 곳이 없다
DROP = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.COMMENTS, RT.SLIDE,
        "http://schemas.microsoft.com/office/2018/10/relationships/comments",
        "http://schemas.microsoft.com/office/2017/10/relationships/threadedComment"}
WARN = []


def warn(n, title, detail):
    WARN.append(f"WARN: {n}|{title}|{detail}")


def layout_for(dst, name):
    for m in dst.slide_masters:
        for l in m.slide_layouts:
            if l.name == name:
                return l
    return dst.slide_layouts[0]


def tmpl_of(partname):
    """/ppt/charts/chart1.xml → /ppt/charts/chart%d.xml (숫자가 없으면 이름 끝에 번호 자리)"""
    p = str(partname).replace("%", "%%")
    head, ext = os.path.splitext(p)
    m = re.search(r"\d+$", head)
    return (head[:m.start()] if m else head) + "%d" + ext


def _set_rel(part, rId, reltype, target, external=False):
    """part 에 같은 rId 로 연결을 단다 — 복사한 XML 안의 r:id 를 고치지 않아도 되게"""
    rels = part.rels
    rels._rels[rId] = _Relationship(part.partname.baseURI, rId, reltype, RTM.EXTERNAL if external else RTM.INTERNAL, target)
    rels.__dict__.pop("_rels_by_reltype", None)   # lazyproperty 캐시 — 새 연결이 보이게


def copy_part(src_part, dst_pkg, memo, attach):
    """src_part 를 dst_pkg 에 새 이름으로 복사한다. attach(new_part) 로 먼저 부모에 매달아야
    next_partname 이 그 이름을 '쓰는 중'으로 본다(같은 장에 차트 두 개면 chart2·chart3). 하위 연결은 재귀로, rId 는 원본 그대로."""
    key = id(src_part)
    if key in memo:   # 같은 장 안에서 두 연결이 한 파트를 가리키면(동영상 video+media) 한 번만
        attach(memo[key])
        return memo[key]
    name = dst_pkg.next_partname(tmpl_of(src_part.partname))
    new = PartFactory(PackURI(name), src_part.content_type, dst_pkg, src_part.blob)
    memo[key] = new
    attach(new)
    for r_id, rel in list(src_part.rels.items()):
        if rel.is_external:
            _set_rel(new, r_id, rel.reltype, rel.target_ref, external=True)
        elif rel.reltype == RT.SLIDE or rel.reltype == RT.SLIDE_LAYOUT:   # 하위 파트가 장·레이아웃을 가리키는 일은 없다 — 있으면 버린다
            continue
        else:
            copy_part(rel.target_part, dst_pkg, memo, lambda p, r_id=r_id, rel=rel: _set_rel(new, r_id, rel.reltype, p))
    return new


def drop_refs(nsld, dead, n):
    """옮기지 못한 rId 를 가리키는 도형은 빼고, 링크(hlinkClick 등)는 r:id 만 비운다 — 깨진 참조로 파일이 안 열리는 걸 막는다"""
    tree = nsld.find(qn("p:cSld")).find(qn("p:spTree"))
    for node in list(nsld.iter()):
        for a, v in list(node.attrib.items()):
            if not a.startswith(R_NS) or v not in dead:
                continue
            tag = node.tag.split("}")[-1]
            if tag in ("hlinkClick", "hlinkHover", "hlinkMouseOver", "snd"):
                node.set(a, "")
                continue
            top = node
            while top.getparent() is not None and top.getparent() is not tree:
                top = top.getparent()
            if top.getparent() is tree:
                tree.remove(top)
                warn(n, "옮기지 못한 개체를 뺐습니다", f"{top.tag.split('}')[-1]} ({dead[v]})")
            else:   # 장 밖(배경·전환 소리 등) — 속성만 지운다
                del node.attrib[a]


def copy_rels(s, new, dst_pkg, n=0, base=""):
    """장 s 의 연결을 새 장 new 에 단다(그림은 같은 그림 한 번만, 차트·SmartArt·미디어 등은 파트째 새 이름으로).
    반환 (rid 옛→새, dead 못 옮긴 rId→종류, memo 복사한 파트). 같은 파일 안 복제(vertical.clone_slide)도 이걸 쓴다(R81)"""
    rid, dead, memo = {}, {}, {}
    for r_id, rel in s.part.rels.items():
        if rel.reltype in DROP:
            if rel.reltype != RT.SLIDE_LAYOUT:
                dead[r_id] = rel.reltype.rsplit("/", 1)[-1]
            continue
        try:
            if rel.is_external:
                rid[r_id] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype == RT.IMAGE:
                _, rid[r_id] = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            else:
                box = {}
                copy_part(rel.target_part, dst_pkg, memo,
                          lambda p, box=box, rel=rel: box.setdefault("rId", new.part.relate_to(p, rel.reltype)))
                rid[r_id] = box["rId"]
        except Exception as e:   # 모르는 꼴이라 못 옮겨도 합치기는 이어 간다 — 그 개체만 빠진다
            dead[r_id] = rel.reltype.rsplit("/", 1)[-1]
            warn(n, "장 안의 연결을 옮기지 못함", f"{base}: {rel.reltype.rsplit('/', 1)[-1]} — {type(e).__name__}: {str(e)[:120]}")
    return rid, dead, memo


def remap_refs(nsld, rid, dead, memo, n=0):
    """복사해 온 장 XML 의 r: 속성을 새 rId 로. 못 옮긴 rId 를 가리키는 도형은 먼저 뺀다"""
    if dead:   # 옛 rId 기준으로 먼저 — 바꾼 뒤에 하면 새 rId 가 옛 '못 옮긴' rId 와 같아 멀쩡한 그림을 뺀다
        drop_refs(nsld, dead, n)
    for node in nsld.iter():   # r: 이름공간 속성 전부 — r:embed·r:link·r:id 에 더해 SmartArt 의 r:dm·r:lo·r:qs·r:cs
        for a, v in list(node.attrib.items()):
            if a.startswith(R_NS) and v in rid:
                node.set(a, rid[v])
    for p in memo.values():   # SmartArt 데이터 파트의 dataModelExt relId 는 '장' 의 rId(그린 모양 drawing) — 새 rId 로
        if "diagramData" in (p.content_type or "") and hasattr(p, "_element"):
            for node in p._element.iter():
                v = node.get("relId")
                if v in rid:
                    node.set("relId", rid[v])


def append(dst, src_path, n=0):
    """src 의 첫 장을 dst 끝에 붙이고 그 장을 돌려준다"""
    s = Presentation(src_path).slides[0]
    new = dst.slides.add_slide(layout_for(dst, s.slide_layout.name))
    for ph in list(new.placeholders):   # 레이아웃이 만든 빈 자리는 지운다
        ph._element.getparent().remove(ph._element)
    rid, dead, memo = copy_rels(s, new, dst.part.package, n, os.path.basename(src_path))
    sld, nsld = s._element, new._element
    cs_new = nsld.find(qn("p:cSld"))
    cs_src = sld.find(qn("p:cSld"))
    bg = cs_src.find(qn("p:bg"))
    if bg is not None:
        cs_new.insert(0, copy.deepcopy(bg))
    tree = cs_new.find(qn("p:spTree"))
    for el in list(cs_src.find(qn("p:spTree"))):
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(copy.deepcopy(el))
    for tag in ("p:clrMapOvr", "p:transition", "p:timing", "p:extLst"):   # 색 매핑·전환·애니메이션·확장
        x = sld.find(qn(tag))
        if x is not None:
            old = nsld.find(qn(tag))
            if old is not None:
                nsld.remove(old)
            nsld.append(copy.deepcopy(x))
    remap_refs(nsld, rid, dead, memo, n)
    return new


def move(pres, old_idx, new_idx):
    lst = pres.slides._sldIdLst
    el = list(lst)[old_idx]
    lst.remove(el)
    lst.insert(new_idx, el)


def delete(pres, idx):
    lst = pres.slides._sldIdLst
    el = list(lst)[idx]
    pres.part.drop_rel(el.rId)
    lst.remove(el)


def merge(out, parts):
    shutil.copy2(parts[0], out)
    pres = Presentation(out)
    while len(pres.slides) > 1:
        delete(pres, len(pres.slides) - 1)
    for i, p in enumerate(parts[1:], 2):
        append(pres, p, i)
    pres.save(out)
    return len(pres.slides)


def replace(base, pairs):
    pres = Presentation(base)
    for n, f in pairs:
        if n < 1 or n > len(pres.slides):
            raise ValueError(f"{n}쪽이 없습니다")
        append(pres, f, n)
        move(pres, len(pres.slides) - 1, n)   # 예전 n 뒤로
        delete(pres, n - 1)
    pres.save(base)
    return len(pairs)


KEPT = "만든 장은 남겨 두었다 — 다시 돌리면 합치기부터"   # merge.py(PowerPoint)·worker.build_parallel 과 같은 꼬리(R81)


def _main():
    if sys.argv[1] == "replace":
        base = os.path.abspath(sys.argv[2])
        pairs = sorted((int(a.split("=", 1)[0]), os.path.abspath(a.split("=", 1)[1])) for a in sys.argv[3:])
        print(f"갈아 끼운 장 {replace(base, pairs)}")
    else:
        out, parts = os.path.abspath(sys.argv[1]), [os.path.abspath(p) for p in sys.argv[2:]]
        if not parts:
            raise ValueError("합칠 파일이 없습니다")
        print(f"장 {merge(out, parts)}개를 한 파일로 합쳤습니다")
    for w in WARN:
        print(w)


if __name__ == "__main__":
    try:
        _main()
    except Exception as e:   # 장 파일(sNN.pptx)은 읽기만 했다 — 그대로 남는다
        sys.exit(f"합치기 실패: {str(e)[:300] if isinstance(e, ValueError) else f'{type(e).__name__}: {str(e)[:300]}'} — {KEPT}")
