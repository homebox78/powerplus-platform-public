# -*- coding: utf-8 -*-
"""[Windows · PowerPoint 있는 PC 에서 한 번] 서체마다 PowerPoint 가 글줄을 놓는 수치를 재서 fonts/ppt_metrics.json 에 둔다.
리눅스 엔진(vcenter.py)은 PowerPoint 없이 이 표와 서체 파일로 'PowerPoint 화면에서 글자가 어디 찍히는지'를 계산한다
(LibreOffice 렌더는 도형 안 글줄 위치가 PowerPoint 와 달라 세로 가운데 실측에 쓸 수 없다 — 2026-10-05 비교).
  asc  = 글 상자 안쪽 위 → 첫 줄 기준선 (pt 당, 줄 간격 100%)
  lh   = 줄 간격 100% 에서 한 줄 높이 (pt 당)
  k    = 줄 간격을 늘렸을 때 첫 기준선이 내려가는 몫 (첫 기준선 = asc + k·(배수-1)) — lh 배수와 따로 잰다
서체 파일 경로는 fontset.INSTALLED(이 PC) 에서 찾는다. 인자 없이 실행: py ppt_metrics.py [서체 이름 …]"""
import json, os, sys, tempfile, time
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fonts", "ppt_metrics.json")
import fontset

TXT = "가나다 Ag1"
PT = 40
PXW = 4000           # 10in 장 → 400px/in
SPC = (1.0, 1.5)


def faces():
    fl = json.load(open(os.path.join(HERE, "fonts", "fonts.json"), encoding="utf-8"))
    out = [t["face"] for t in fl["title"]]
    for b in fl["body"]:
        out += b["faces"]
    return list(dict.fromkeys(out + list(fontset.ASTA) + ["맑은 고딕"]))


def txt_for(face):
    """재는 글 — 한글이 없는 서체(Noto Sans JP·SC, 2026-10-18 언어별 서체)는 한자로. 없는 글자를 PowerPoint 가 다른 서체로 그리면 줄 높이가 그 서체 것이 된다"""
    import fontcover
    path = fontset.INSTALLED.get(face.lower())
    return "漢字体 Ag1" if path and ord("가") not in fontcover.cmap(path) else TXT


def box(slide, face, spc, top_in):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(top_in), Inches(9), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    bp = tf._txBody.find(qn("a:bodyPr"))
    for x in list(bp):
        bp.remove(x)
    from lxml import etree
    etree.SubElement(bp, qn("a:noAutofit"))
    for i in range(2):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spc
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = txt_for(face)
        r.font.size = Pt(PT)
        rpr = r._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            etree.SubElement(rpr, qn(tag)).set("typeface", face)


def bands(im, x0, x1, y0, y1):
    """잉크 줄 묶음 [(위, 아래)] — 빈 줄로 나뉜 덩어리"""
    px = im.load()
    rows = [y for y in range(y0, y1) if any(px[x, y] < 128 for x in range(x0, x1, 2))]
    out = []
    for y in rows:
        if out and y - out[-1][1] <= 3:
            out[-1][1] = y
        else:
            out.append([y, y])
    return [tuple(b) for b in out if b[1] - b[0] > 3]


def main():
    want = sys.argv[1:] or faces()
    import win32com.client as win32
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    for f in want:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box(s, f, SPC[0], 0.5)
        box(s, f, SPC[1], 3.5)
    tmp = tempfile.mkdtemp(prefix="pd_metrics_")
    src = os.path.join(tmp, "m.pptx")
    prs.save(src)
    app = win32.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(src, ReadOnly=True, WithWindow=False)
    try:
        for i in range(1, pres.Slides.Count + 1):
            pres.Slides(i).Export(os.path.join(tmp, f"m{i:02d}.png"), "PNG", PXW, int(PXW * 0.75))
    finally:
        pres.Close()
        time.sleep(0.3)
    K = PXW / 720.0   # px / pt
    res = json.load(open(OUT, encoding="utf-8")) if os.path.isfile(OUT) else {}
    for i, f in enumerate(want, 1):
        path = fontset.INSTALLED.get(f.lower())
        if not path:
            print("서체 파일 없음(건너뜀):", f)
            continue
        im = Image.open(os.path.join(tmp, f"m{i:02d}.png")).convert("L")
        font = ImageFont.truetype(path, int(round(PT * K)))
        btop = font.getbbox(txt_for(f), anchor="ls")[1]   # 기준선 위 잉크 끝(음수 px)
        m = {}
        for spc, top_in in zip(SPC, (0.5, 3.5)):
            y0 = int(top_in * 72 * K)
            b = bands(im, int(0.5 * 72 * K), int(9.5 * 72 * K), y0, y0 + int(2.2 * 72 * K))
            if len(b) < 2:
                print("잉크 줄을 못 찾음:", f, spc, b)
                break
            base1 = b[0][0] - btop
            base2 = b[1][0] - btop
            m[spc] = ((base1 - y0) / K / PT, (base2 - base1) / K / PT)
        if len(m) < 2:
            continue
        asc, lh = m[1.0]
        asc15, lh15 = m[1.5]
        res[f] = {"asc": round(asc, 4), "lh": round(lh, 4), "k": round((asc15 - asc) / 0.5, 4), "lhk": round(lh15 / lh / 1.5, 4),
                  "file": os.path.basename(path)}
        print(f, res[f])
    json.dump(res, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("→", OUT)


if __name__ == "__main__":
    main()
