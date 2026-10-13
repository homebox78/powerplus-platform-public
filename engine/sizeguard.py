# -*- coding: utf-8 -*-
"""초안 다듬기 결과를 원본과 견줘 '하면 안 되는 변경'을 코드로 되돌린다 — 지침서 문장만으로는 AI 가 안 지킨다.
(2026-10-12 실제 다듬기 작업: 본문 12pt 를 7~8pt 로 줄이고 기획 메모 도형까지 바꿨다.)

  sizeguard.py <원본.pptx> <결과.pptx> <저장.pptx> [--dry]

1. 기획 메모(노란 채움·'메모' 이름·장 밖 도형)는 원본 그대로 되돌린다(서체·크기·문구·위치).
2. 같은 도형(도형 번호가 같고 글이 같은 문단)에서 글자가 원본보다 작아졌으면 원본 크기로 되돌린다.
   되돌린 뒤 상자를 넘치면 0.25pt 씩 다시 줄인다 — AI 가 정한 크기 밑으로는 내려가지 않는다.
   자동 맞춤(normAutofit fontScale)으로 줄인 것도 원본 비율로.
3. 장 아래쪽이 25% 넘게 비면 WARN 줄로 알린다(worker 가 확인 목록에 싣는다).
출력: 사람이 읽는 줄 + 'WARN: 장|제목|설명' + 'GUARD: 오류slug|장|개수'(known_errors.json 의 slug — 엔진이 개선 체크리스트에 쌓는다)."""
import copy, os, re, sys
sys.argv, _ARGV = sys.argv[:1], sys.argv
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vcenter as V
import textlines as T
import cardfit as C
sys.argv = _ARGV
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

sys.stdout.reconfigure(encoding="utf-8")
EMU = V.EMU
_WS = re.compile(r"\s+")


def _sid(el):
    nv = el.find(".//" + qn("p:cNvPr"))
    return nv.get("id") if nv is not None else None


def _norm(s):
    return _WS.sub("", s or "")


def _shapes(slide):
    return {s: sh for sh in V.flat(slide.shapes._spTree) for s in [_sid(sh.el)] if s}


def _key_texts(m):
    return {k for k, sh in m.items() if sh.kind == "sp" and _norm("".join(x.text or "" for x in sh.el.iter(qn("a:t"))))}


def _match(orig, out):
    """결과 장마다 원본 장 하나 — 같은 장 수면 순서대로, 아니면 도형 번호·글이 가장 많이 겹치는 장"""
    om = [_shapes(s) for s in orig.slides]
    res = []
    same = len(orig.slides) == len(out.slides)
    for i, s in enumerate(out.slides):
        m = _shapes(s)
        if same:
            res.append((m, om[i]))
            continue
        best, bi = -1, None
        for j, o in enumerate(om):
            sc = sum(1 for k in _key_texts(m) if k in o and _norm(_txt(o[k].el))[:20] == _norm(_txt(m[k].el))[:20])
            if sc > best:
                best, bi = sc, j
        res.append((m, om[bi] if bi is not None and best > 0 else {}))
    return res


def _txt(el):
    return "".join(x.text or "" for x in el.iter(qn("a:t")))


def _box(sh):
    b = T.Box(sh)
    return b if b.ok else None


def _overflow(box):
    return any(k in ("넘침", "가로넘침") for k, *_ in T.problems(box))


def _restore_scale(ob, nb):
    """결과에 더 작은 자동 맞춤 비율이 붙었으면 원본 비율로"""
    if ob is None or nb is None or ob.bp is None or nb.bp is None or nb.fscale >= ob.fscale - 0.001:
        return False
    na = nb.bp.find(qn("a:normAutofit"))
    if ob.fscale >= 0.999:
        na.attrib.pop("fontScale", None)
        na.attrib.pop("lnSpcReduction", None)
    else:
        na.set("fontScale", str(int(round(ob.fscale * 100000))))
    return True


def guard(src, out, dst, dry):
    orig, res = Presentation(src), Presentation(out)
    SW, SH = res.slide_width / EMU, res.slide_height / EMU
    memo_n = size_n = scale_n = 0
    warns, guards = [], []
    for sno, (m, om) in enumerate(_match(orig, res), 1):
        changed = []
        m0, s0, c0 = memo_n, size_n, scale_n
        for k, sh in list(m.items()):
            o = om.get(k)
            if o is None or sh.kind != "sp" or o.kind != "sp":
                continue
            # 1. 기획 메모 — 통째로 원본 그대로
            if C._memo(o, orig.slide_width / EMU, orig.slide_height / EMU):
                if etree.tostring(o.el) != etree.tostring(sh.el):
                    memo_n += 1
                    if not dry:
                        sh.el.getparent().replace(sh.el, copy.deepcopy(o.el))
                continue
            ob, nb = _box(o), _box(sh)
            if ob is None or nb is None or _norm(ob.text()) != _norm(nb.text()):
                continue
            if _restore_scale(ob, nb) and not dry:
                scale_n += 1
                nb = _box(sh)
            if len(ob.paras) != len(nb.paras):
                continue
            touched = []
            for op, np_ in zip(ob.paras, nb.paras):
                if _norm(op.text) != _norm(np_.text) or len(op.runs) != len(np_.runs):
                    continue
                prev = T._sizes(np_)
                up = False
                for orr, nr in zip(op.runs, np_.runs):
                    ov = V._sz(orr.find(qn("a:rPr")), op.end, op.ldef)
                    nv = V._sz(nr.find(qn("a:rPr")), np_.end, np_.ldef)
                    if nv < ov - 0.3 and T._size_known(op):
                        up = True
                        if not dry:
                            rp = nr.find(qn("a:rPr"))
                            if rp is None:
                                rp = nr.makeelement(qn("a:rPr"), {})
                                nr.insert(0, rp)
                            rp.set("sz", str(int(round(ov * 100))))
                if up:
                    touched.append((np_, prev))
            if not touched:
                continue
            size_n += 1
            changed.append(sh.Name)
            if dry:
                continue
            # 넘치면 되돌린 문단만 0.25pt 씩 — AI 가 정한 크기 밑으로는 안 내려간다
            for _ in range(40):
                b = _box(sh)
                if b is None or not _overflow(b):
                    break
                moved = False
                for np_, prev in touched:
                    pa = next((p for p in b.paras if p.p is np_.p), None)
                    if pa is None:
                        continue
                    floor = [int(s) / 100 if s else 0 for s in prev]
                    cur = [V._sz(r.find(qn("a:rPr")), pa.end, pa.ldef) for r in pa.runs]
                    if all(c - 0.25 >= f - 0.01 for c, f in zip(cur, floor)):
                        T._set_size(pa, -0.25)
                        moved = True
                if not moved:
                    break
        for slug, d in (("memo-changed", memo_n - m0), ("font-shrink", size_n - s0), ("autofit-scale", scale_n - c0)):
            if d:
                guards.append(f"GUARD: {slug}|{sno}|{d}")
        if changed:
            print(f"{sno}장 글자 크기 원본으로: " + ", ".join(changed[:8]) + (" …" if len(changed) > 8 else ""))
        # 3. 아래쪽 빈 공간
        bots = [s.Top + s.Height for s in V.flat(res.slides[sno - 1].shapes._spTree)
                if not C._memo(s, SW, SH) and s.Top + s.Height <= SH + 1
                and not (s.Width > SW * 0.9 and s.Height > SH * 0.4)        # 배경 판
                and not (s.Top > SH * 0.88 and s.Height < SH * 0.12)]       # 바닥 띠·쪽 번호
        if bots:
            gap = (SH * 0.88 - max(bots)) / SH
            if gap > 0.25:
                guards.append(f"GUARD: bottom-empty|{sno}|1")
                warns.append(f"WARN: {sno}|장 아래쪽이 비었습니다|내용이 장 높이의 {100 - int(gap * 100) - 12}%까지만 차 있습니다 — 카드·도식을 키우거나 아래로 펼쳐 주세요")
    if not dry:
        res.save(dst)
    print(f"기획 메모 되돌림 {memo_n}개 · 글자 크기 되돌림 {size_n}개 · 자동 맞춤 비율 되돌림 {scale_n}개{' (미리보기)' if dry else ''}")
    for w in warns + guards:
        print(w)


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(a) < 3:
        print(__doc__)
        sys.exit(2)
    guard(os.path.abspath(a[0]), os.path.abspath(a[1]), os.path.abspath(a[2]), "--dry" in sys.argv)
