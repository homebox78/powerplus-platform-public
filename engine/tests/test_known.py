# -*- coding: utf-8 -*-
"""한 번 나온 오류(engine/known_errors.json)의 회귀 시험 — 합성 pptx 로 게이트가 지금도 막는지 본다.
known_errors.json 의 test 값이 이 파일의 함수 이름이다. 새 오류를 고치면 그 줄과 함께 시험을 추가한다.

  python3 engine/tests/test_known.py        (전부 통과면 0, 하나라도 실패면 1 — deploy-engine.sh 가 배포 전에 돌린다)
PowerPoint·LibreOffice·AI 호출 없음."""
import io, json, os, subprocess, sys, tempfile, contextlib

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
ROOT = os.path.dirname(ENGINE)
sys.path.insert(0, ENGINE)
_ARGV = sys.argv
sys.argv = sys.argv[:1]
import textlines as T  # noqa: E402
import cardfit as C  # noqa: E402
import sizeguard as G  # noqa: E402
sys.argv = _ARGV

from pptx import Presentation  # noqa: E402
from pptx.util import Pt, Inches  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

TMP = tempfile.mkdtemp(prefix="known-")


def _path(name):
    return os.path.join(TMP, name)


def _deck():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p, p.slides.add_slide(p.slide_layouts[6])


def _text(sl, x, y, w, h, s, pt=None, fill=None, name=None):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    tf = sh.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = s
    if pt:
        r.font.size = Pt(pt)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    return sh


def _sizes(path):
    out = []
    for sl in Presentation(path).slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        out.append((sh.name, r.font.size.pt if r.font.size else None))
    return out


def _guard(src, out):
    dst = _path("guarded.pptx")
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        G.guard(src, out, dst, False)
    return dst, buf.getvalue()


def _pair(orig_pt, new_pt, text="원본 본문 글자는 줄이지 않는다", memo=False):
    """같은 도형 번호·같은 글로 원본/결과 두 판"""
    a, b = _path("orig.pptx"), _path("out.pptx")
    for path, pt in ((a, orig_pt), (b, new_pt)):
        p, sl = _deck()
        _text(sl, 1, 1, 6, 1.5, text, pt, fill="FFFF99" if memo else None, name="메모" if memo else "본문")
        _text(sl, 1, 5.8, 11, 0.8, "아래까지 채우는 띠", 12)     # 아래 빈 공간 경고 안 나게
        p.save(path)
    return a, b


# ── 시험 ─────────────────────────────────────────

def test_font_shrink_restored():
    a, b = _pair(12, 8)
    dst, log = _guard(a, b)
    got = dict(_sizes(dst))["본문"]
    assert got == 12, f"8pt 로 줄인 글이 원본 12pt 로 안 돌아옴: {got}"
    assert "GUARD: font-shrink|1|" in log, "GUARD 줄이 안 나옴 — 체크리스트에 안 쌓인다"


def test_autofit_scale_restored():
    a, b = _pair(12, 12)
    p = Presentation(b)
    sh = next(s for s in p.slides[0].shapes if s.name == "본문")
    bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
    for c in list(bp):
        bp.remove(c)
    na = bp.makeelement(qn("a:normAutofit"), {"fontScale": "50000"})
    bp.append(na)
    p.save(b)
    dst, log = _guard(a, b)
    bp = next(s for s in Presentation(dst).slides[0].shapes if s.name == "본문").text_frame._txBody.find(qn("a:bodyPr"))
    na = bp.find(qn("a:normAutofit"))
    assert na is None or na.get("fontScale") is None, f"fontScale 그대로: {na.get('fontScale')}"
    assert "GUARD: autofit-scale|1|" in log


def test_memo_restored():
    a, b = _pair(11, 7, text="디자인: 여기 표를 넣어 주세요", memo=True)
    dst, log = _guard(a, b)
    got = dict(_sizes(dst))["메모"]
    assert got == 11, f"메모 도형이 원본으로 안 돌아옴: {got}"
    assert "GUARD: memo-changed|1|" in log


def test_unknown_size_untouched():
    p, sl = _deck()
    sh = _text(sl, 1, 1, 4, 1, "크기 없는 글", None)
    box = T.Box(next(s for s in __import__("vcenter").flat(sl.shapes._spTree) if s.Name == sh.name))
    assert box.paras and not T._size_known(box.paras[0]), "크기를 모르는 글을 아는 것으로 봄 — 18pt 를 박는다"


def _cards(heights=(4.0, 4.0)):
    p, sl = _deck()
    for i, h in enumerate(heights):
        x = 1 + i * 4
        c = sl.shapes.add_shape(1, Inches(x), Inches(1.2), Inches(3.5), Inches(h))
        c.name = f"카드{i + 1}"
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFB)
        _text(sl, x + 0.2, 1.4, 3.1, 0.6, f"카드 {i + 1} 제목", 14, name=f"글{i + 1}")
    src = _path("cards.pptx")
    p.save(src)
    return src


def _card_h(path):
    return {s.name: s.height / 914400 for s in Presentation(path).slides[0].shapes if s.name.startswith("카드")}


def test_card_fit_shrinks_siblings():
    src, dst = _cards(), _path("cards-fit.pptx")
    with contextlib.redirect_stdout(io.StringIO()):
        C.fit(src, dst, False, None, 10, 18)
    h = _card_h(dst)
    assert h["카드1"] < 4.0 - 0.1, f"빈 공간을 안 줄임: {h}"
    assert abs(h["카드1"] - h["카드2"]) < 0.01, f"형제 카드 높이가 제각각: {h}"


def test_card_fit_keeps_60():
    src, dst = _cards(), _path("cards-fit60.pptx")
    with contextlib.redirect_stdout(io.StringIO()):
        C.fit(src, dst, False, None, 10, 18)
    h = _card_h(dst)
    assert h["카드1"] >= 4.0 * 0.6 - 0.01, f"원래의 60% 아래로 줄임: {h}"


def test_double_space():
    p, sl = _deck()
    _text(sl, 1, 1, 6, 1, "AI  가  심사한다", 12, name="본문")
    _text(sl, 1, 3, 6, 1, "디자인: 메모  그대로", 12, fill="FFFF99", name="메모")
    src, dst = _path("sp.pptx"), _path("sp-out.pptx")
    p.save(src)
    with contextlib.redirect_stdout(io.StringIO()):
        C.spaces(src, dst, False, None)
    t = {s.name: s.text_frame.text for s in Presentation(dst).slides[0].shapes}
    assert t["본문"] == "AI 가 심사한다", t["본문"]
    assert t["메모"] == "디자인: 메모  그대로", "메모 도형의 공백까지 고침"


def test_bottom_empty_warned():
    a, b = _path("be-a.pptx"), _path("be-b.pptx")
    for path in (a, b):
        p, sl = _deck()
        _text(sl, 1, 0.5, 6, 1, "위쪽에만 있는 글", 14, name="본문")
        p.save(path)
    _, log = _guard(a, b)
    assert "GUARD: bottom-empty|1|1" in log and "WARN: 1|" in log, log


def test_binary_attrs():
    for f in ("x.jpg", "x.png", "x.ttf", "x.woff2", "x.pptx"):
        r = subprocess.run(["git", "-C", ROOT, "check-attr", "binary", "--", "engine/skill/" + f],
                           capture_output=True, text=True)
        if r.returncode != 0:      # 저장소 밖(서버 배포본)에서는 건너뜀
            return
        assert r.stdout.strip().endswith("set"), f"engine 아래 {f} 가 binary 가 아님 — 줄끝 규칙이 깨뜨린다: {r.stdout}"


def test_reference_guide_present():
    d = os.path.join(ENGINE, "reference-guide")
    files = [f for f in os.listdir(d)] if os.path.isdir(d) else []
    assert len(files) >= 3, f"engine/reference-guide 지침 파일이 없음: {files}"


def test_engine_syntax():
    """엔진 파이썬 전부 문법 검사 — 시험이 안 읽는 파일(worker.py 등)의 들여쓰기 깨짐도 잡는다(2026-10-12 실제로 잡힘)"""
    import ast
    for dp, _, fs in os.walk(ENGINE):
        if ".venv" in dp or "__pycache__" in dp:
            continue
        for f in fs:
            if f.endswith(".py"):
                p = os.path.join(dp, f)
                ast.parse(open(p, encoding="utf-8").read(), p)


def test_json_quote_repair():
    """AI 가 쓴 JSON 의 글 속 따옴표·빠진 쉼표 — 40장 개요가 두 번 다 이것으로 실패했다(2026-10-13)"""
    import re
    src = open(os.path.join(ENGINE, "design.py"), encoding="utf-8").read()
    ns = {"re": re, "json": json}
    exec(src[src.index("def _loads(t):"):src.index("def pptx_text(path):")], ns)
    pj = ns["parse_json"]
    assert pj('{"title": "식품 "통합" 구축", "core": "x"}')["title"] == '식품 "통합" 구축'
    assert len(pj('{"a": [{"t": "1"} {"t": "2"}], "b": "c"}')["a"]) == 2
    assert pj('{"a": "x", "b": "그는 "좋다", 했다", "c": 1}')["c"] == 1
    assert pj('{"ch": [{"name": "사업 "이해"", "s": []}]}')["ch"][0]["name"] == '사업 "이해"'
    assert pj('{"a": "ok", "b": ["x", "y"]}') == {"a": "ok", "b": ["x", "y"]}
    assert pj('{"a": "x" "b": "그는 "좋다" 했다", "c": 1}') == {"a": "x", "b": '그는 "좋다" 했다', "c": 1}   # 닫는 따옴표 뒤 쉼표 빠짐
    # 완성된 객체 뒤에 객체가 더 붙음(Extra data) — B안 와이어 묶음 4건이 이것으로 버려졌다
    assert [x["n"] for x in pj('{"slides": [{"n": 1}]} {"slides": [{"n": 2}]}')["slides"]] == [1, 2]
    assert [x["n"] for x in pj('{"n": 3, "wire": {}}\n{"n": 4, "wire": {}}')["slides"]] == [3, 4]
    assert pj('{"slides": [{"n": 1}], "t": "a"}\n설명 {"x": 1}')["t"] == "a"

def test_divider_memo_tone():
    """간지 한 줄 — 약속(promise) 우선, 메모 말투 goal 은 버린다"""
    import re as _re
    src = open(os.path.join(ENGINE, "design.py"), encoding="utf-8").read()
    ns = {"re": _re}
    for name in ("MEMO_TONE", "divider_line"):
        pre = name + r"\b" if name == "MEMO_TONE" else "def " + name + r"\("
        m = _re.search(r"^(" + pre + r".*?)(?=^\S)", src, _re.S | _re.M)
        exec(m.group(1), ns)
    dl = ns["divider_line"]
    assert dl({"promise": "정확하게 **알고** 있습니다!", "goal": "x를 보인다"}) == "정확하게 **알고** 있습니다!"
    assert dl({"goal": "자료 품질 기준을 먼저 못 박는다."}) == ""
    assert dl({"goal": "AI 판단 체계로 실행 확신을 준다"}) == ""
    assert dl({"goal": "치밀하게 준비했습니다!"}) == "치밀하게 준비했습니다!"
    assert "promise" in src.split("OUTLINE_SYSTEM = ")[1][:3000]


def test_registry_consistent():
    """known_errors.json 의 모든 test 가 이 파일에 있고, slug 가 겹치지 않는다"""
    items = json.load(open(os.path.join(ENGINE, "known_errors.json"), encoding="utf-8"))["items"]
    slugs = [x["slug"] for x in items]
    assert len(slugs) == len(set(slugs)), "slug 중복"
    for x in items:
        for k in ("slug", "area", "title", "rule", "gate", "commit"):
            assert x.get(k), f"{x.get('slug')}: {k} 비어 있음"
        if x.get("test"):
            assert x["test"] in globals(), f"{x['slug']}: 시험 함수 {x['test']} 없음"


def main():
    tests = [(k, v) for k, v in globals().items() if k.startswith("test_") and callable(v)]
    bad = 0
    for k, fn in tests:
        try:
            fn()
            print(f"통과  {k}")
        except Exception as e:      # noqa: BLE001
            bad += 1
            print(f"실패  {k}: {type(e).__name__}: {e}")
    print(f"{len(tests) - bad}/{len(tests)} 통과")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
