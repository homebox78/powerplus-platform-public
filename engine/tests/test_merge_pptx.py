# -*- coding: utf-8 -*-
"""리눅스 합치기(merge_pptx.py) 시험 — 차트·표·그림·SmartArt·메모·링크·태그가 든 한 장짜리 pptx 를 합쳐
장 수·차트 수·차트 데이터(카테고리·값)·내장 엑셀이 원본과 같은지, 모든 r:id 가 살아 있는지 본다(R79).

  python engine/tests/test_merge_pptx.py          (전부 통과면 0)
  PD_MERGE_REAL_DIR=<폴더> python …               그 폴더의 s01·s02·s05·s07.pptx(운영 작업 55ae9b33 실물)로도 시험 — 파일은 저장소에 넣지 않는다
PowerPoint·LibreOffice·AI 호출 없음."""
import io, os, shutil, subprocess, sys, tempfile, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)
import merge_pptx as M  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.opc.package import Part  # noqa: E402
from pptx.opc.packuri import PackURI  # noqa: E402
from pptx.oxml import parse_xml  # noqa: E402
from pptx.util import Inches  # noqa: E402

TMP = tempfile.mkdtemp(prefix="merge-")
R_NS = M.R_NS
PNG = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                    "1f15c4890000000d49444154789c6360f8cfc0f01f0005fe02fea7d9a1"
                    "f60000000049454e44ae426082")


def _p(name):
    return os.path.join(TMP, name)


def _blank():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def make_line():
    prs, s = _blank()
    cd = CategoryChartData()
    cd.categories = [f"{y}-{m:02d}" for y in (2023, 2024, 2025) for m in range(1, 13)]   # R69 의 월별 36개월 꺾은선
    cd.add_series("전화", [100 + i * 3 for i in range(36)])
    cd.add_series("온라인", [50 + (i % 7) * 11 for i in range(36)])
    s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1), Inches(8), Inches(4.5), cd)
    s.shapes.add_picture(io.BytesIO(PNG), Inches(10), Inches(1), Inches(1), Inches(1))
    s.notes_slide.notes_text_frame.text = "발표 메모"   # 메모는 버려져도 파일은 열려야 한다
    prs.save(_p("line.pptx"))
    return _p("line.pptx")


def make_bar():
    prs, s = _blank()
    cd = CategoryChartData()
    cd.categories = ["가", "나", "다"]
    cd.add_series("2024", (1.5, 2.25, 3))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), cd)
    cd2 = CategoryChartData()
    cd2.categories = ["A", "B"]
    cd2.add_series("x", (7, 9))
    s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(7), Inches(1), Inches(5), Inches(4), cd2)   # 한 장에 차트 두 개
    t = s.shapes.add_table(2, 2, Inches(1), Inches(5.5), Inches(4), Inches(1)).table
    t.cell(0, 0).text = "구분"
    t.cell(1, 1).text = "값"
    tb = s.shapes.add_textbox(Inches(7), Inches(5.5), Inches(4), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "누리집"
    r.hyperlink.address = "https://example.com/"   # 외부 링크
    prs.save(_p("bar.pptx"))
    return _p("bar.pptx")


DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"


def make_smartart():
    """python-pptx 는 SmartArt 를 못 만든다 — 파트 다섯 개(data·layout·quickStyle·colors·drawing)를 손으로 단다"""
    prs, s = _blank()
    pkg = prs.part.package
    rt = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    kinds = [("data", "diagramData", "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml", "dataModel"),
             ("layout", "diagramLayout", "application/vnd.openxmlformats-officedocument.drawingml.diagramLayout+xml", "layoutDef"),
             ("quickStyle", "diagramQuickStyle", "application/vnd.openxmlformats-officedocument.drawingml.diagramStyle+xml", "styleDef"),
             ("colors", "diagramColors", "application/vnd.openxmlformats-officedocument.drawingml.diagramColors+xml", "colorsDef")]
    ids = {}
    drawing = Part(PackURI("/ppt/diagrams/drawing1.xml"), "application/vnd.ms-office.drawingml.diagramDrawing+xml", pkg,
                   f'<dsp:drawing xmlns:dsp="{DSP}" xmlns:dgm="{DGM}"><dsp:spTree/></dsp:drawing>'.encode())
    d_rid = s.part.relate_to(drawing, "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing")
    img = Part(PackURI("/ppt/media/dimage1.png"), "image/png", pkg, PNG)   # drawing 안 그림 — 하위 연결 재귀
    drawing.relate_to(img, rt + "image")
    for nm, rel, ct, root in kinds:
        body = f'<dgm:{root} xmlns:dgm="{DGM}" xmlns:dsp="{DSP}">'
        if nm == "data":
            body += f'<dgm:extLst><dsp:dataModelExt relId="{d_rid}" minVer="http://schemas.openxmlformats.org/drawingml/2006/diagram"/></dgm:extLst>'
        body += f"</dgm:{root}>"
        p = Part(PackURI(f"/ppt/diagrams/{nm}1.xml"), ct, pkg, body.encode())
        ids[nm] = s.part.relate_to(p, rt + rel)
    gf = parse_xml(
        '<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:nvGraphicFramePr><p:cNvPr id="50" name="SmartArt 1"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="914400" y="914400"/><a:ext cx="4572000" cy="2743200"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{DGM}"><dgm:relIds xmlns:dgm="{DGM}" '
        f'r:dm="{ids["data"]}" r:lo="{ids["layout"]}" r:qs="{ids["quickStyle"]}" r:cs="{ids["colors"]}"/>'
        '</a:graphicData></a:graphic></p:graphicFrame>')
    s.shapes._spTree.append(gf)
    tag = Part(PackURI("/ppt/tags/tag1.xml"), "application/vnd.openxmlformats-officedocument.presentationml.tags+xml", pkg,
               b'<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tag name="PD" val="1"/></p:tagLst>')
    s.part.relate_to(tag, rt + "tags")   # 모르는 내부 연결 꼴 — 통째로 복사
    prs.save(_p("smart.pptx"))
    return _p("smart.pptx")


def chart_data(prs):
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                pl = sh.chart.plots[0]
                xlsx = sh.chart.part.chart_workbook.xlsx_part
                out.append((sh.chart.chart_type, list(pl.categories), [(se.name, list(se.values)) for se in pl.series],
                            xlsx.blob if xlsx is not None else None))
    return out


def check_integrity(path):
    """모든 장의 r: 속성이 그 장의 연결에 있고, 같은 이름 파트가 두 번 들어가지 않았다"""
    names = zipfile.ZipFile(path).namelist()
    assert len(names) == len(set(names)), "zip 안 이름 중복"
    prs = Presentation(path)
    for i, s in enumerate(prs.slides, 1):
        for node in s._element.iter():
            for a, v in node.attrib.items():
                if a.startswith(R_NS) and v:
                    assert v in s.part.rels, f"{i}쪽 {node.tag.split('}')[-1]}@{a.split('}')[-1]}={v} 연결 없음"
    return prs


def test_merge_charts_tables_images():
    files = [make_line(), make_bar(), make_line(), make_smartart()]
    want = sum((chart_data(Presentation(f)) for f in files), [])
    out = _p("merged.pptx")
    n = M.merge(out, files)
    assert n == 4, f"장 수 {n}"
    prs = check_integrity(out)
    got = chart_data(prs)
    assert len(got) == len(want) == 4, f"차트 수 {len(got)} (원본 {len(want)})"
    for g, w in zip(got, want):
        assert g[:3] == w[:3], f"차트 데이터 다름: {g[:2]} / {w[:2]}"
        assert g[3] == w[3], "내장 엑셀이 원본과 다름"
    charts = [x for x in zipfile.ZipFile(out).namelist() if x.startswith("ppt/charts/chart")]
    assert len(charts) == 4, f"차트 파트 {charts}"   # chart1.xml 이 여러 장에 있어도 이름이 겹치지 않음
    tables = sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)
    assert tables == 1, f"표 {tables}"
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
    assert pics == 2, f"그림 {pics}"
    links = [r.hyperlink.address for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert links == ["https://example.com/"], f"링크 {links}"


def test_merge_smartart_and_tags():
    out = _p("merged_sa.pptx")
    M.merge(out, [make_bar(), make_smartart()])
    prs = check_integrity(out)
    s = prs.slides[1]
    rels = {r.reltype.rsplit("/", 1)[-1]: r for r in s.part.rels.values()}
    for k in ("diagramData", "diagramLayout", "diagramQuickStyle", "diagramColors", "diagramDrawing", "tags"):
        assert k in rels, f"SmartArt·태그 연결 {k} 없음: {sorted(rels)}"
    gf = [sh for sh in s.shapes if sh.name == "SmartArt 1"]
    assert gf, "SmartArt 도형이 빠짐"
    ri = next(gf[0]._element.iter("{%s}relIds" % DGM))
    assert s.part.rels[ri.get(R_NS + "dm")].reltype.endswith("diagramData")
    assert s.part.rels[ri.get(R_NS + "cs")].reltype.endswith("diagramColors")
    data = rels["diagramData"].target_part
    ext = next(parse_xml(data.blob).iter("{%s}dataModelExt" % DSP))
    assert s.part.rels[ext.get("relId")].reltype.endswith("diagramDrawing"), "dataModelExt relId 가 drawing 을 안 가리킴"
    drw = rels["diagramDrawing"].target_part
    assert any(r.reltype.endswith("/image") for r in drw.rels.values()), "drawing 안 그림 연결이 빠짐"


def test_replace_keeps_chart():
    base = _p("base.pptx")
    M.merge(base, [make_bar(), make_line(), make_smartart()])
    M.replace(base, [(2, make_bar())])
    prs = check_integrity(base)
    assert len(prs.slides) == 3
    kinds = [len([sh for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]) for s in prs.slides]
    assert kinds == [2, 2, 0], f"장마다 차트 수 {kinds}"


def test_unmovable_rel_warns_not_exits():
    """끝내 못 옮기는 연결은 그 도형만 빼고 WARN — 전체를 멈추지 않는다(작업 55ae9b33 은 여기서 sys.exit)"""
    orig = M.PartFactory

    def boom(partname, content_type, package, blob):
        if "chart" in content_type and "drawingml.chart+xml" in content_type:
            raise RuntimeError("시험용 실패")
        return orig(partname, content_type, package, blob)
    M.PartFactory, M.WARN[:] = boom, []
    try:
        out = _p("merged_warn.pptx")
        n = M.merge(out, [make_smartart(), make_line(), make_bar()])
    finally:
        M.PartFactory = orig
    assert n == 3
    prs = check_integrity(out)
    assert not chart_data(prs), "실패한 차트가 깨진 채 남음"
    assert any("옮기지 못" in w for w in M.WARN), M.WARN
    assert sum(1 for sh in prs.slides[1].shapes if sh.shape_type == 13) == 1, "같은 장의 그림까지 빠짐"
    M.WARN[:] = []


def test_cli_merge_prints_and_exits_0():
    out = _p("cli.pptx")
    r = subprocess.run([sys.executable, os.path.join(ENGINE, "merge_pptx.py"), out, make_bar(), make_line()],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stderr[-300:]
    assert "장 2개" in r.stdout, r.stdout


def test_real_job_55ae9b33():
    """운영 작업 55ae9b33 의 실물 4장(2쪽 꺾은선 차트) — PD_MERGE_REAL_DIR 이 있을 때만"""
    d = os.environ.get("PD_MERGE_REAL_DIR")
    if not d:
        print("    (PD_MERGE_REAL_DIR 없음 — 실물 시험 건너뜀)")
        return
    src = [os.path.join(d, f"s{n:02d}.pptx") for n in (1, 2, 5, 7)]
    files = []
    for f in src:   # 실물은 읽기만 — 복사본으로
        shutil.copy2(f, _p("real_" + os.path.basename(f)))
        files.append(_p("real_" + os.path.basename(f)))
    want = sum((chart_data(Presentation(f)) for f in files), [])
    out = _p("real_merged.pptx")
    assert M.merge(out, files) == 4
    got = chart_data(check_integrity(out))
    assert len(got) == len(want) == 1, f"차트 수 {len(got)} (원본 {len(want)})"
    assert got[0][:3] == want[0][:3] and got[0][3] == want[0][3], "차트 데이터가 원본과 다름"


def main():
    tests = [(k, v) for k, v in globals().items() if k.startswith("test_") and callable(v)]
    bad = 0
    for k, fn in tests:
        try:
            fn()
            print(f"통과  {k}")
        except Exception as e:      # noqa: BLE001
            bad += 1
            print(f"실패  {k}: {type(e).__name__}: {e}")
    print(f"{len(tests) - bad}/{len(tests)} 통과")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
