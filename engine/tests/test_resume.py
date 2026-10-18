# -*- coding: utf-8 -*-
"""이어 만들기 시험(R81) — 다시 돌리면 끝난 일을 처음부터 하지 않는지, 일부 실패가 앞 결과를 버리지 않는지.
  ① 기획 단계: 기획 뒤 시안이 0장이거나 도중에 예외가 나도, 다시 돌리면 기획 AI 를 다시 부르지 않는다(시작 옵션 '샘플 1장'도 지킨다)
  ② 엔진 하나(build_parallel=0): 엔진 뒤(렌더·올리기)에서 실패하면 다시 돌릴 때 엔진을 부르지 않는다 · 끝까지 성공한 뒤 다시 돌리면 새로 만든다
  ② 세로 제안서: 조립에서 실패하면 다시 돌릴 때 목차·본문 기획을 다시 짜지 않고, 실패했던 장의 본문만 다시 · 성공 뒤 다시 돌리면 새로
  ③ vertical.clone_slide: 견본의 그림·차트 연결(rId)을 새 장에 옮긴다(깨진 r:embed·r:id 0, 차트는 새 파트)
  ④ merge.py(PowerPoint) 실패 글 꼬리 '만든 장은 남겨 두었다 — 다시 돌리면 합치기부터', 장 파일은 그대로

  python -X utf8 engine/tests/test_resume.py              (전부 통과면 0)
  PD_ENGINE_DIR=<고치기 전 engine 사본> python …          음성 대조 — 고치기 전 코드는 ①②③④ 가 실패해야 한다
가짜 서버(req)·가짜 엔진(run_claude)·가짜 AI(make_plan·outline…)만 쓴다. 실제 AI 호출·운영 서버·PowerPoint 없음."""
import glob, io, json, os, shutil, subprocess, sys, tempfile, types

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.abspath(os.environ.get("PD_ENGINE_DIR") or os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="resume-")
ENG = os.path.join(TMP, "engine")
os.makedirs(ENG)
for pat in ("*.py", "*.json", "*.md"):
    for f in glob.glob(os.path.join(SRC, pat)):
        shutil.copy2(f, ENG)
os.makedirs(os.path.join(ENG, "fonts"))
for f in glob.glob(os.path.join(SRC, "fonts", "*.json")) or glob.glob(os.path.join(os.path.dirname(HERE), "fonts", "*.json")):   # 서체 목록(파일은 빼고)
    shutil.copy2(f, os.path.join(ENG, "fonts"))
if os.path.isdir(os.path.join(SRC, "templates")):
    shutil.copytree(os.path.join(SRC, "templates"), os.path.join(ENG, "templates"))
JOBS = os.path.join(TMP, "jobs")
json.dump({"api_base": "http://127.0.0.1:9/api/v1/deck", "worker_key": "test", "name": "test", "jobs_dir": JOBS, "ai": "openrouter",
           "build_parallel": 0, "image_parallel": 2, "max_minutes": 1}, open(os.path.join(ENG, "config.json"), "w", encoding="utf-8"))
sys.path.insert(0, ENG)
os.chdir(TMP)

import worker as WK  # noqa: E402
import design as D  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PNG = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                    "1f15c4890000000d49444154789c6360f8cfc0f01f0005fe02fea7d9a1"
                    "f60000000049454e44ae426082")


class Server:
    """가짜 서버 — 작업 한 건의 plan 칸·보고를 흉내(Jobs::savePlan·finish·finishPlan 과 같은 규칙)"""
    def __init__(self):
        self.plan, self.finish, self.calls = None, [], []

    def req(self, method, path, body=None, raw=None, timeout=60):
        self.calls.append(path)
        if path.endswith("/plan") and isinstance((body or {}).get("plan"), dict):
            self.plan = json.dumps(body["plan"], ensure_ascii=False)
        if path.endswith("/finish"):
            self.finish.append(body)
            if isinstance(body.get("plan"), dict) and body.get("status") in ("review", "failed"):
                self.plan = json.dumps(body["plan"], ensure_ascii=False)
        return {}


def patch(mod, **kw):
    for k, v in kw.items():
        setattr(mod, k, v)


def server_into(sv):
    patch(WK, req=sv.req, put_file=lambda *a, **k: {}, download=_no_download)
    D.W.update({"req": sv.req, "put_file": lambda *a, **k: {}, "download": _no_download, "log": lambda *a: None})


def _no_download(path, dest):
    raise RuntimeError("404")


def blank_pptx(path, chart=False):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(1), Inches(1), Inches(1))
    if chart:
        cd = CategoryChartData()
        cd.categories = ["가", "나"]
        cd.add_series("값", (1, 2))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(2), Inches(4), Inches(3), cd)
    prs.save(path)
    return path


# ───────── ① 기획 단계
def _plan_job(jid, opts):
    return {"id": jid, "kind": "design", "phase": "plan", "status": "running", "brief": "자료 없는 직접 쓴 글", "note": "", "theme": "#2468FF",
            "opts": json.dumps(opts), "plan": None, "spent_usd": 0}


def _plan_mocks(calls, mock_ok, front_raise=False):
    def make_plan(job, *a, **k):
        calls["plan"] += 1
        return {"title": "T", "slides": [{"n": i, "title": f"{i}쪽", "mood": f"m{i}", "wire": []} for i in (1, 2, 3)]}

    def make_mockup(job, s, *a, **k):
        calls["mock"].append(s["n"])
        if not mock_ok:
            raise RuntimeError("이미지 모델 실패(가짜)")
        return PNG

    def save_mockup(job, n, png, s=None, project=""):
        p = os.path.join(job.dir, "_mockup", f"s{n:02d}.jpg")
        os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, "wb").write(png)
        return p

    def front_specs(*a, **k):
        if front_raise:
            raise RuntimeError("앞장 규칙 읽기 실패(가짜 예외)")
        return []
    patch(D, load_context=lambda *a, **k: "", make_plan=make_plan, make_wires=lambda *a, **k: None, fact_check=lambda *a, **k: 0,
          make_mockup=make_mockup, check_mockup=lambda job, s, png, make: (png, None), save_mockup=save_mockup, front_specs=front_specs,
          lessons_text=lambda j: "", parse_plan_md=lambda b: None)


def _run_plan(j):
    job = WK.Job(j)
    try:
        D._plan_phase(job, j)
        return None
    except Exception as e:   # noqa: BLE001 — process() 가 failed 로 보고하는 자리
        return e


def test_plan_kept_when_all_mockups_fail():
    sv = Server(); server_into(sv)
    calls = {"plan": 0, "mock": []}
    j = _plan_job("p1aaaaaaaaaaaaaa", {"mockups": "sample", "sample": 1})
    _plan_mocks(calls, mock_ok=False)
    _run_plan(j)
    assert sv.finish and sv.finish[-1]["status"] == "failed", f"첫 실행은 시안 0장 실패여야 함: {sv.finish}"
    _plan_mocks(calls, mock_ok=True)
    _run_plan(dict(j, plan=sv.plan))
    assert calls["plan"] == 1, f"다시 돌렸더니 기획 AI 를 {calls['plan']}번 불렀다(1번이어야)"
    assert sv.finish[-1]["status"] == "review", f"둘째 실행 보고: {sv.finish[-1].get('status')}"
    assert calls["mock"][1:] == [1], f"시작 옵션(샘플 1장)을 지켜야 함 — 둘째 실행 시안: {calls['mock'][1:]}"
    assert "pre_mockup" not in (sv.finish[-1].get("plan") or {}), "확인 대기(review) 기획에 '시안 전' 표시가 남음"


def test_plan_kept_when_exception_after_plan():
    sv = Server(); server_into(sv)
    calls = {"plan": 0, "mock": []}
    j = _plan_job("p2aaaaaaaaaaaaaa", {"mockups": "all"})
    _plan_mocks(calls, mock_ok=True, front_raise=True)
    e = _run_plan(j)
    assert e is not None, "첫 실행은 도중 예외여야 함"
    _plan_mocks(calls, mock_ok=True)
    _run_plan(dict(j, plan=sv.plan))
    assert calls["plan"] == 1, f"도중 예외 뒤 다시 돌렸더니 기획 AI 를 {calls['plan']}번 불렀다(1번이어야)"
    assert sorted(calls["mock"]) == [1, 2, 3], f"'전체' 옵션 — 시안 3장: {calls['mock']}"


# ───────── ② 엔진 하나(build_parallel=0)
def _single_job(jid):
    j = {"id": jid, "kind": "design", "phase": "build", "status": "running", "title": "T", "note": "", "theme": "#2468FF", "opts": "{}",
         "plan": json.dumps({"title": "T", "slides": [{"n": 1, "title": "가"}]})}
    d = os.path.join(JOBS, jid, "_mockup")
    os.makedirs(d, exist_ok=True)
    json.dump(json.loads(j["plan"]), open(os.path.join(d, "plan.json"), "w", encoding="utf-8"))
    open(os.path.join(d, "s01.jpg"), "wb").write(PNG)
    return j


def _single_mocks(calls, fail_upload):
    def run_claude(job, prompt):
        calls["engine"] += 1
        blank_pptx(os.path.join(job.dir, "결과.pptx"))
        json.dump({"result": "결과.pptx", "summary": "엔진 결과", "checklist": []}, open(os.path.join(job.dir, "result.json"), "w", encoding="utf-8"))
        return 0, {"turns": 3, "cost": 0.5, "text": ""}

    def upload_after(job, out, only=None):
        if fail_upload[0]:
            fail_upload[0] = False
            raise RuntimeError("결과 렌더 실패(가짜)")

    class R:
        returncode, stdout, stderr = 0, "", ""
    patch(WK, run_claude=run_claude, upload_after=upload_after, text_map=lambda *a, **k: None, front_matter=lambda *a, **k: "",
          subprocess=types.SimpleNamespace(run=lambda *a, **k: R(), PIPE=subprocess.PIPE))
    patch(D, prepare_build=lambda job, j: None, lessons_text=lambda j: "")
    WK.CFG["build_parallel"] = 0


def test_single_engine_resume():
    sv = Server(); server_into(sv)
    calls = {"engine": 0}
    j = _single_job("s1aaaaaaaaaaaaaa")
    _single_mocks(calls, [True])
    WK.process(j)
    assert sv.finish[-1]["status"] == "failed", f"첫 실행은 렌더 실패여야 함: {sv.finish[-1]}"
    WK.process(j)
    assert sv.finish[-1]["status"] == "done", f"둘째 실행: {sv.finish[-1].get('status')} {sv.finish[-1].get('error', '')[:200]}"
    assert calls["engine"] == 1, f"엔진 뒤 실패를 다시 돌렸더니 엔진을 {calls['engine']}번 불렀다(1번이어야)"
    WK.process(j)   # 끝까지 성공한 작업을 다시 돌림 = 새로 만들어 달라는 뜻
    assert calls["engine"] == 2, f"성공 뒤 다시 돌리기는 새로 만들어야 함(엔진 {calls['engine']}번)"


# ───────── ② 세로 제안서
def _vertical_mocks(calls, fail_b, fail_build):
    import vertical as V

    def outline(job, brief, note, ctx, raw, pages):
        calls["outline"] += 1
        return {"project": {"name": "가짜 사업", "client": "", "year": "2026", "doc": "기술제안서"}, "_reqs": 0,
                "chapters": [{"no": "Ⅰ", "title": "가", "sections": []}, {"no": "Ⅱ", "title": "나", "sections": []}, {"no": "※", "title": "별첨", "sections": []}]}

    def chapter_pages(job, o, ch, ctx, raw):
        calls["pages"].append(ch["title"])
        if ch["title"] == "나" and fail_b[0]:
            raise RuntimeError("본문 AI 실패(가짜)")
        return [{"title": ch["title"] + " 본문", "blocks": []}]

    class Build:
        def __init__(self, *a, **k):
            self.trimmed, self.pages = [], []

        def build(self, chapters_pages, out):
            calls["build"].append(sorted(len(v) for v in chapters_pages.values()))
            if fail_build[0]:
                fail_build[0] = False
                raise RuntimeError("조립 실패(가짜)")
            blank_pptx(out)
            return 3
    patch(V, _ctx=lambda job, j, brief: ("요약", "원문"), outline=outline, chapter_pages=chapter_pages, resolve_assets=lambda *a, **k: None,
          backgrounds=lambda *a, **k: ({}, ["가짜 배경"]), client_logo=lambda o: None, Build=Build, build_guide=lambda *a, **k: None,
          font_paths=lambda: ([], []), audit_align=lambda prs: [], kind_mix=lambda cp: {})
    return V


def _run_vertical(V, j):
    job = WK.Job(j)
    job.spent = 0.0
    try:
        V.run(job, j, {"purpose": "submit", "slides": 6}, j["brief"], "")
        return None
    except Exception as e:   # noqa: BLE001
        return e


def test_vertical_resume():
    sv = Server(); server_into(sv)
    calls = {"outline": 0, "pages": [], "build": []}
    j = {"id": "v1aaaaaaaaaaaaaa", "kind": "design", "phase": "plan", "brief": "세로 제안서 글", "note": "", "opts": "{}"}
    fail_b, fail_build = [True], [True]
    V = _vertical_mocks(calls, fail_b, fail_build)
    e = _run_vertical(V, j)
    assert e is not None, "첫 실행은 조립 실패여야 함"
    fail_b[0] = False
    e = _run_vertical(V, j)
    assert e is None, f"둘째 실행 실패: {e}"
    assert calls["outline"] == 1, f"다시 돌렸더니 목차 AI 를 {calls['outline']}번 불렀다(1번이어야)"
    assert calls["pages"].count("가") == 1, f"앞 실행에서 만든 '가' 장 본문을 다시 만듦: {calls['pages']}"
    assert calls["pages"].count("나") == 2, f"실패했던 '나' 장만 다시 만들어야 함: {calls['pages']}"
    assert sv.finish[-1]["status"] == "done"
    _run_vertical(V, j)   # 성공 뒤 다시 돌리기 = 새로
    assert calls["outline"] == 2, f"성공 뒤 다시 돌리기는 목차부터 새로(목차 {calls['outline']}번)"


# ───────── ③ vertical.clone_slide
def _broken_refs(prs):
    bad = []
    for i, s in enumerate(prs.slides, 1):
        rels = s.part.rels
        for node in s._element.iter():
            for a, v in node.attrib.items():
                if a.startswith(R_NS) and v and v not in rels:
                    bad.append(f"{i}쪽 {node.tag.split('}')[-1]} {v}")
    return bad


def test_clone_slide_moves_picture_and_chart():
    import vertical as V
    src_p = blank_pptx(os.path.join(TMP, "clone_src.pptx"), chart=True)
    prs = Presentation(src_p)
    src = prs.slides[0]
    src.shapes[0].click_action.hyperlink.address = "https://example.com"   # 바깥 링크(r:id)
    for _ in range(3):   # 같은 파일 안에서 여러 번 — 그 사이 rId 번호가 엇갈리게
        prs.slides.add_slide(prs.slide_layouts[5])
        V.clone_slide(prs, src)
    out = os.path.join(TMP, "clone_out.pptx")
    prs.save(out)
    back = Presentation(out)
    bad = _broken_refs(back)
    assert not bad, f"깨진 연결 {len(bad)}곳: {bad[:4]}"
    charts = [sh.chart.part.partname for s in back.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 4 and len(set(charts)) == 4, f"차트는 장마다 새 파트여야 함: {charts}"
    pics = [sh.image.blob for s in back.slides for sh in s.shapes if sh.shape_type == 13]
    assert len(pics) == 4 and all(b == PNG for b in pics), f"그림 {len(pics)}개(4개·원본과 같아야)"
    vals = [list(sh.chart.plots[0].series[0].values) for s in back.slides for sh in s.shapes if sh.has_chart]
    assert all(v == [1.0, 2.0] for v in vals), f"복제한 차트 값: {vals}"


# ───────── ④ merge.py(PowerPoint) 실패 글
def test_merge_win_failure_message_keeps_slides():
    fake = os.path.join(TMP, "fakecom", "win32com")
    os.makedirs(fake, exist_ok=True)
    open(os.path.join(fake, "__init__.py"), "w").write("")
    open(os.path.join(fake, "client.py"), "w", encoding="utf-8").write(
        "class _Slides:\n    Count = 1\n    def InsertFromFile(self, *a):\n        raise RuntimeError('(-2147352567, 예외가 발생했습니다.)')\n"
        "    def __call__(self, i):\n        raise RuntimeError('없음')\n"
        "class _Pres:\n    Slides = _Slides()\n    def Save(self): pass\n    def Close(self): pass\n"
        "class _Ps:\n    def Open(self, *a, **k): return _Pres()\n"
        "class _App:\n    Presentations = _Ps()\n"
        "def Dispatch(name): return _App()\n")
    parts = [blank_pptx(os.path.join(TMP, f"m{i}.pptx")) for i in (1, 2)]
    before = [open(p, "rb").read() for p in parts]
    env = dict(os.environ, PYTHONPATH=os.path.join(TMP, "fakecom"), PYTHONIOENCODING="utf-8")
    code = ("import os, runpy, sys; os.name = 'nt'; sys.argv = sys.argv[1:]; runpy.run_path(sys.argv[0], run_name='__main__')")
    r = subprocess.run([sys.executable, "-c", code, os.path.join(ENG, "merge.py"), os.path.join(TMP, "m_out.pptx")] + parts,
                       capture_output=True, text=True, encoding="utf-8", errors="replace", env=env, timeout=60)
    msg = (r.stderr or r.stdout).strip()
    assert r.returncode != 0, "PowerPoint 오류인데 성공으로 끝남"
    assert msg.endswith("만든 장은 남겨 두었다 — 다시 돌리면 합치기부터"), f"실패 글 꼬리가 다름: …{msg[-160:]}"
    assert "Traceback" not in msg, "실패 글이 Traceback 으로 나감"
    assert all(os.path.isfile(p) and open(p, "rb").read() == b for p, b in zip(parts, before)), "실패 뒤 장 파일이 바뀌거나 지워짐"


def main():
    tests = [(k, v) for k, v in globals().items() if k.startswith("test_") and callable(v)]
    bad = 0
    for k, fn in tests:
        try:
            fn()
            print(f"통과  {k}")
        except Exception as e:      # noqa: BLE001
            bad += 1
            print(f"실패  {k}: {type(e).__name__}: {e}")
    print(f"{len(tests) - bad}/{len(tests)} 통과  (엔진 코드: {SRC})")
    return 1 if bad else 0


if __name__ == "__main__":
    rc = main()
    os.chdir(HERE)
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(rc)
