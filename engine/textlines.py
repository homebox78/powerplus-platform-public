# -*- coding: utf-8 -*-
"""PowerPoint 없이 '도형 안 글이 몇 줄로, 어디서 끊기는가'를 계산한다 — 리눅스 엔진의 TextRange.Lines 대신.
줄 나누기는 PowerPoint 한글 규칙을 따른다: 한글·한자·가나는 글자 사이 어디서나(그래서 어절이 잘린다),
영문·숫자는 단어째, 공백 뒤에서 끊고, 줄 머리 금지 문장부호(, . ) ! ? 등)는 앞 글자에 붙인다.
글자 폭은 서체 파일(PIL getlength)로, 줄 높이는 PowerPoint 글줄 수치표(fonts/ppt_metrics.json, vcenter.py 와 같은 표)로 잰다.

  textlines.py scan <pptx> [--slides 3,5] [--json out.json]
      장·도형·문단마다 어절 끊김 · 고아 줄(마지막 줄 1~2글자) · 넘침(글이 상자보다 큼)을 보고 — line_scan.py · tone_overflow.py 대신
  textlines.py fix <src> <dst> [--slides …] [--report out.json] [--no-shrink]
      fit_breaks.py · orphan_line.py · nobreak_korean.py 대신: 1순위 문단 글자 0.25pt 씩 줄이기(최대 -2pt, 7pt 미만 금지),
      2순위 잘린 어절 앞 공백을 줄바꿈(<a:br/>)으로, 고아 줄은 윗줄 마지막 어절을 함께 내림. 못 푼 것은 보고.
  textlines.py lines <pptx> <장 번호> [도형 이름 일부]
      도형마다 계산한 줄을 그대로 출력(눈 확인용)
제외: 노란 기획 메모 · '디자인' 말풍선 · 줄바꿈 끈 글상자 · 세로 쓰기 · 돌린 도형 · 표(셀은 따로 보지 않는다).
한계: PowerPoint 와 1~2% 폭 차이가 날 수 있다 — 그래서 폭을 3% 여유(TL_SAFE) 두고 계산하며, 최종 확인은 렌더(lo.py)로 한다."""
import copy, json, os, re, sys
sys.argv, _ARGV = sys.argv[:1], sys.argv   # vcenter 가 가져올 때 인자를 읽지 않게
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vcenter as V
sys.argv = _ARGV
from pptx import Presentation
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
EMU = V.EMU
MEMO = {"FFFF00", "FFFFCC", "FFFF99", "FFFF87"}
NO_HEAD = set(",.!?)]}%:;·…」』〉》’”、。，．！？）］｝")   # 줄 머리에 못 오는 글자
NO_TAIL = set("([{〈《「『‘“（［｛")                     # 줄 끝에 못 오는 글자
WORD = re.compile(r"[0-9A-Za-z가-힣]")
SP = "  \t"
SAFE = float(os.environ.get("TL_SAFE", "0.97"))   # 폭 여유 — LibreOffice·PowerPoint 가 PIL 계산보다 2~3% 넓게 그린다


def cjk(c):
    o = ord(c)
    return (0xAC00 <= o <= 0xD7A3 or 0x1100 <= o <= 0x11FF or 0x3130 <= o <= 0x318F or 0x3040 <= o <= 0x30FF
            or 0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF or 0xF900 <= o <= 0xFAFF or 0xFF00 <= o <= 0xFFEF)


_CW = {}


def cwidth(face, c, pt):
    k = (face, c)
    if k not in _CW:
        f = V._font(face)
        _CW[k] = (f.getlength(c) / 200) if f else (1.0 if cjk(c) else 0.55)
    return _CW[k] * pt


def _defrpr(lst, lvl):
    if lst is None:
        return None
    lv = lst.find(qn(f"a:lvl{lvl + 1}pPr"))
    return lv.find(qn("a:defRPr")) if lv is not None else lst.find(".//" + qn("a:defRPr"))


class Para:
    """문단 하나 — 글자 목록 [(글자, 폭, run 요소|None, run 안 위치)], 줄은 lines()"""
    def __init__(self, p, ldef, fscale, ea_break):
        self.p, self.ldef, self.fscale, self.ea = p, ldef, fscale, ea_break
        self.end = p.find(qn("a:endParaRPr"))
        ppr = p.find(qn("a:pPr"))
        self.ppr = ppr
        self.mar = int(ppr.get("marL", "0")) / EMU if ppr is not None else 0.0
        self.ind = int(ppr.get("indent", "0")) / EMU if ppr is not None else 0.0
        self.chars, self.hard = [], set()    # hard: 그 위치 앞에서 강제 줄바꿈(<a:br/>)
        self.runs = []
        for el in p:
            if el.tag in (qn("a:r"), qn("a:fld")):
                rp = el.find(qn("a:rPr"))
                face = V._face(rp, self.end, ldef)
                pt = V._sz(rp, self.end, ldef) * fscale
                spc = int(rp.get("spc", "0")) / 100 if rp is not None and rp.get("spc") else 0.0
                t = el.findtext(qn("a:t")) or ""
                self.runs.append(el)
                for i, c in enumerate(t):
                    self.chars.append((c, cwidth(face, c, pt) + spc, el, i))
            elif el.tag == qn("a:br"):
                self.hard.add(len(self.chars))
        self.text = "".join(c[0] for c in self.chars)

    def biggest(self):
        rs = [r.find(qn("a:rPr")) for r in self.runs] or [None]
        rp = max(rs, key=lambda r: V._sz(r, self.end, self.ldef))
        return V._sz(rp, self.end, self.ldef) * self.fscale, V._face(rp, self.end, self.ldef)

    def can_break(self, i):
        """i 번째 글자 앞에서 끊을 수 있는가"""
        if i <= 0 or i >= len(self.chars):
            return False
        a, b = self.chars[i - 1][0], self.chars[i][0]
        if b in NO_HEAD or a in NO_TAIL:
            return False
        if a in SP:
            return True
        if b in SP:
            return False
        return self.ea and (cjk(a) or cjk(b))

    def lines(self, width):
        """[(시작, 끝)] — 끝은 포함 안 함, 줄 끝 공백 포함"""
        out, n, s = [], len(self.chars), 0
        if n == 0:
            return [(0, 0)]
        while s < n:
            w, last_ok, i = 0.0, None, s
            avail = (width - (self.ind if not out else 0)) * SAFE
            while i < n:
                if i in self.hard and i > s:
                    break
                c, cw = self.chars[i][0], self.chars[i][1]
                if i > s and self.can_break(i):
                    last_ok = i
                if c not in SP and w + cw > avail + 0.01 and i > s:
                    i = last_ok if last_ok and last_ok > s else i
                    break
                w += cw
                i += 1
            out.append((s, i))
            s = i
        return out

    def seg_text(self, a, b):
        return self.text[a:b]

    def width_of(self, a, b):
        t = self.chars[a:b]
        while t and t[-1][0] in SP:
            t = t[:-1]
        return sum(c[1] for c in t)


def is_memo(el):
    pr = el.find(qn("p:spPr"))
    if pr is not None:
        f = pr.find(qn("a:solidFill"))
        c = f.find(qn("a:srgbClr")) if f is not None else None
        if c is not None and c.get("val", "").upper() in MEMO:
            return True
    nv = el.find(".//" + qn("p:cNvPr"))
    name = nv.get("name", "") if nv is not None else ""
    return "말풍선" in name and "디자인" in "".join(x.text or "" for x in el.iter(qn("a:t")))


class Box:
    """글 상자 하나 — 문단들과 줄, 높이"""
    def __init__(self, sh):
        self.sh, self.el = sh, sh.el
        tx = self.el.find(qn("p:txBody"))
        self.tx = tx
        self.ok = tx is not None
        if not self.ok:
            return
        bp = tx.find(qn("a:bodyPr"))
        self.bp = bp
        g = lambda k, d: int(bp.get(k)) / EMU if bp is not None and bp.get(k) else d
        li, ri, ti, bi = g("lIns", 7.2), g("rIns", 7.2), g("tIns", 3.6), g("bIns", 3.6)
        self.wrap = bp is None or bp.get("wrap") != "none"
        self.vert = bp is not None and bp.get("vert") not in (None, "horz")
        self.autofit = bp is not None and bp.find(qn("a:spAutoFit")) is not None
        na = bp.find(qn("a:normAutofit")) if bp is not None else None
        self.fscale = int(na.get("fontScale", "100000")) / 100000 if na is not None else 1.0
        self.lred = int(na.get("lnSpcReduction", "0")) / 100000 if na is not None else 0.0
        self.box_w = sh.Width / sh.sx - li - ri
        self.box_h = sh.Height / sh.sy - ti - bi
        lst = tx.find(qn("a:lstStyle"))
        self.lst = lst
        self.paras = []
        for p in tx.findall(qn("a:p")):
            ppr = p.find(qn("a:pPr"))
            lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
            lp = lst.find(qn(f"a:lvl{lvl + 1}pPr")) if lst is not None else None
            ea = not ((ppr is not None and ppr.get("eaLnBrk") == "0") or (lp is not None and lp.get("eaLnBrk") == "0"))
            self.paras.append(Para(p, _defrpr(lst, lvl), self.fscale, ea))

    def text(self):
        return "\n".join(p.text for p in self.paras)

    def layout(self):
        """[(문단, [(시작,끝)])], 글 높이(pt)"""
        res, total, first = [], 0.0, True
        for pa in self.paras:
            w = self.box_w - pa.mar if self.wrap else 1e9
            ls = pa.lines(w)
            pt, face = pa.biggest()
            m = V._metric(face)
            spc = 1.0
            ls_el = pa.ppr.find(qn("a:lnSpc")) if pa.ppr is not None else None
            if ls_el is not None and ls_el.find(qn("a:spcPct")) is not None:
                spc = int(ls_el.find(qn("a:spcPct")).get("val")) / 100000
            elif ls_el is not None and ls_el.find(qn("a:spcPts")) is not None:
                spc = int(ls_el.find(qn("a:spcPts")).get("val")) / 100 / max(1.0, m["lh"] * pt)
            spc = max(0.1, spc - self.lred)
            lh = m["lh"] * pt * spc * m.get("lhk", 1.0)
            if not first:
                total += V._spc_pts(pa.ppr, "a:spcBef")
            total += lh * len(ls) + V._spc_pts(pa.ppr, "a:spcAft")
            first = False
            res.append((pa, ls))
        return res, total


def problems(box):
    """[(종류, 문단 번호, 줄 번호, 설명)]"""
    out = []
    lay, total = box.layout()
    for pi, (pa, ls) in enumerate(lay):
        if not pa.text.strip():
            continue
        for k in range(len(ls) - 1):
            a, b = ls[k], ls[k + 1]
            if b[0] in pa.hard or a[1] <= a[0] or b[1] <= b[0]:
                continue
            x, y = pa.text[a[1] - 1], pa.text[b[0]]
            if x in SP or y in SP:
                continue
            if WORD.match(x) and WORD.match(y):
                out.append(("어절", pi, k, pa.seg_text(*a).strip()[-8:] + "|" + pa.seg_text(*b).strip()[:8]))
        if len(ls) >= 2 and len(pa.text.strip()) > 6 and ls[-1][0] not in pa.hard:
            tail = pa.seg_text(*ls[-1]).strip()
            if 0 < len(tail) <= 2:
                out.append(("고아", pi, len(ls) - 1, tail))
        if box.wrap:
            for k, l in enumerate(ls):
                room = box.box_w - pa.mar - (pa.ind if k == 0 else 0)
                w = pa.width_of(*l)
                if room > 0 and w > room * SAFE + 0.5:
                    out.append(("가로넘침", pi, k, f"{pa.seg_text(*l).strip()[:12]} {w:.1f}/{room:.1f}pt"))
    if not box.autofit and box.box_h > 0 and total > box.box_h + 1.0:
        out.append(("넘침", -1, -1, f"글 {total:.1f}pt > 상자 {box.box_h:.1f}pt"))
    if not box.wrap:
        for pi, (pa, ls) in enumerate(lay):
            if ls and pa.width_of(0, len(pa.chars)) > box.box_w + 1:
                out.append(("가로넘침", pi, 0, pa.text[:16]))
    return out


def boxes(slide):
    for sh in V.flat(slide.shapes._spTree):
        if sh.kind != "sp" or (abs(sh.rot) > 1 and abs(sh.rot - 360) > 1):
            continue
        b = Box(sh)
        if not b.ok or b.vert or is_memo(sh.el) or not b.text().strip():
            continue
        yield b


# ── 고치기 ────────────────────────────────────────────────

def _set_size(pa, delta):
    for r in pa.runs:
        rp = r.find(qn("a:rPr"))
        if rp is None:
            rp = r.makeelement(qn("a:rPr"), {})
            r.insert(0, rp)
        cur = V._sz(rp, pa.end, pa.ldef)
        rp.set("sz", str(int(round((cur + delta) * 100))))


def _size_known(pa):
    """글자 크기가 이 상자 안(글자·문단 끝·목록 기본)에 적혀 있나 — 없으면 마스터에서 물려받는 크기라 모른다.
    모르는 채로 줄이면 기본값 18pt 에서 빼서 오히려 키운 크기를 박는다(2026-10-12 한 장 시험에서 카드 글이 겹침)."""
    for r in pa.runs:
        rp = r.find(qn("a:rPr"))
        if not ((rp is not None and rp.get("sz")) or (pa.end is not None and pa.end.get("sz"))
                or (pa.ldef is not None and pa.ldef.get("sz"))):
            return False
    return True


def _sizes(pa):
    return [(r.find(qn("a:rPr")).get("sz") if r.find(qn("a:rPr")) is not None else None) for r in pa.runs]


def _restore(pa, ss):
    for r, s in zip(pa.runs, ss):
        rp = r.find(qn("a:rPr"))
        if rp is None:
            continue
        if s is None:
            rp.attrib.pop("sz", None)
        else:
            rp.set("sz", s)


def _break_at_space(pa, idx):
    """paragraph 글자 idx(공백)를 <a:br/> 로 바꾼다 — run 을 둘로 나눈다"""
    c, _, run, ri = pa.chars[idx]
    if c not in SP or run is None or run.tag != qn("a:r"):
        return False
    t = run.find(qn("a:t"))
    s = t.text or ""
    before, after = s[:ri], s[ri + 1:]
    t.text = before
    br = run.makeelement(qn("a:br"), {})
    rp = run.find(qn("a:rPr"))
    if rp is not None:
        br.append(copy.deepcopy(rp))
    run.addnext(br)
    if after:
        nr = copy.deepcopy(run)
        nr.find(qn("a:t")).text = after
        br.addnext(nr)
    if not before:
        run.getparent().remove(run)
    return True


def _rebuild(box, pi):
    p = box.paras[pi].p
    ppr = p.find(qn("a:pPr"))
    lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
    box.paras[pi] = Para(p, _defrpr(box.lst, lvl), box.fscale, box.paras[pi].ea)
    return box.paras[pi]


def _para_probs(box, pi):
    return [x for x in problems(box) if x[1] == pi and x[0] in ("어절", "고아")]


def fix_box(box, shrink=True, log=None):
    log = log if log is not None else []
    for pi in range(len(box.paras)):
        pa = box.paras[pi]
        if not pa.text.strip() or not _para_probs(box, pi):
            continue
        base = _sizes(pa)
        pt0, _ = pa.biggest()
        done = False
        if shrink and _size_known(pa):
            for step in range(1, 9):
                _restore(pa, base)
                _set_size(pa, -0.25 * step)
                pa = _rebuild(box, pi)
                if pa.biggest()[0] < 7:
                    break
                if not _para_probs(box, pi) and not any(x[0] == "넘침" for x in problems(box)):
                    log.append(("축소", pa.text.strip()[:30], -0.25 * step))
                    done = True
                    break
            if not done:
                _restore(pa, base)
                pa = _rebuild(box, pi)
        if done:
            continue
        # 줄바꿈: 뒤쪽 끊김부터 — 그 줄의 마지막 공백을 줄바꿈으로
        for _ in range(6):
            pr = _para_probs(box, pi)
            if not pr:
                break
            kind, _, k, _ = pr[-1]
            lay, _ = box.layout()
            ls = lay[pi][1]
            line = k if kind == "어절" else k - 1      # 고아 = 윗줄 마지막 어절을 함께 내린다
            a, b = ls[line]
            seg = pa.text[a:b].rstrip(SP)
            cut = max(seg.rfind(" "), seg.rfind(" "))
            if cut <= 0 or not _break_at_space(pa, a + cut):
                log.append(("남김", pa.text.strip()[:30], "어절이 줄보다 김" if kind == "어절" else "고아 줄"))
                break
            pa = _rebuild(box, pi)
            log.append(("줄바꿈", pa.text.replace("\n", " ").strip()[:30], kind))
    return log


def _slides_opt():
    if "--slides" in sys.argv:
        return {int(x) for x in sys.argv[sys.argv.index("--slides") + 1].split(",") if x}
    return None


def main():
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    only = _slides_opt()
    for k in ("--slides", "--json", "--report"):
        if k in sys.argv:
            v = sys.argv[sys.argv.index(k) + 1]
            if v in a:
                a.remove(v)
    if not a:
        print(__doc__)
        return
    cmd = a[0]
    pres = Presentation(a[1])
    if cmd == "lines":
        sno, key = int(a[2]), (a[3] if len(a) > 3 else "")
        for b in boxes(pres.slides[sno - 1]):
            if key and key not in b.sh.Name:
                continue
            lay, total = b.layout()
            print(f"■ {b.sh.Name}  폭 {b.box_w:.1f}pt 높이 {b.box_h:.1f}pt 글 {total:.1f}pt")
            for pa, ls in lay:
                for s, e in ls:
                    print("   │" + pa.text[s:e] + "│")
        return
    if cmd == "scan":
        rows = []
        for sno, sl in enumerate(pres.slides, 1):
            if only and sno not in only:
                continue
            for b in boxes(sl):
                for kind, pi, li, desc in problems(b):
                    rows.append(dict(slide=sno, shape=b.sh.Name, kind=kind, para=pi, line=li, desc=desc,
                                     text=b.text().replace("\n", " ")[:40]))
        cnt = {}
        for r in rows:
            cnt[r["kind"]] = cnt.get(r["kind"], 0) + 1
            print(f"{r['slide']:>3}쪽 {r['kind']:<4} {r['shape'][:18]:<18} {r['desc']}")
        print("합계", " · ".join(f"{k} {v}" for k, v in cnt.items()) or "문제 없음")
        if "--json" in sys.argv:
            json.dump(rows, open(sys.argv[sys.argv.index("--json") + 1], "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        return
    if cmd == "fix":
        dst = a[2]
        log = []
        for sno, sl in enumerate(pres.slides, 1):
            if only and sno not in only:
                continue
            for b in boxes(sl):
                for x in fix_box(b, shrink="--no-shrink" not in sys.argv):
                    log.append((sno, b.sh.Name) + x)
        pres.save(dst)
        cnt = {}
        for x in log:
            cnt[x[2]] = cnt.get(x[2], 0) + 1
        print(" · ".join(f"{k} {v}" for k, v in cnt.items()) or "고칠 것 없음")
        for x in log:
            if x[2] == "남김":
                print("  남김", x)
        if "--report" in sys.argv:
            json.dump(log, open(sys.argv[sys.argv.index("--report") + 1], "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        return
    print(__doc__)


if __name__ == "__main__":
    main()
