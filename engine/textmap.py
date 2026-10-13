# -*- coding: utf-8 -*-
"""웹에서 장표 글을 바로 고치기 위한 도구 (AI 없이).
  map   <pptx> <out.json>              장마다 글이 든 도형·표 칸의 위치(슬라이드 기준 %)와 글
  apply <pptx> <edits.json> <out.pptx> edits = [{slide, id, text}] 를 그 도형에 넣는다(서식은 그대로)
id: 도형은 "s<shape_id>", 표 칸은 "s<shape_id>:<행>:<열>". 줄바꿈("\n") = 문단 또는 줄 나눔.
"""
import copy, json, sys
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def walk(shapes, tf=(0, 0, 1.0, 1.0)):
    """(도형, 절대 left, top, width, height) — 그룹 안 도형은 그룹 변환을 풀어 슬라이드 좌표로"""
    ox, oy, kx, ky = tf
    for s in shapes:
        try:
            l, t, w, h = ox + (s.left or 0) * kx, oy + (s.top or 0) * ky, (s.width or 0) * kx, (s.height or 0) * ky
        except Exception:
            continue
        if s.shape_type == 6:
            x = s._element.grpSpPr.find(A + "xfrm")
            ch_off, ch_ext = (x.find(A + "chOff"), x.find(A + "chExt")) if x is not None else (None, None)
            if ch_off is None or ch_ext is None or not int(ch_ext.get("cx")) or not int(ch_ext.get("cy")):
                continue
            gx, gy = w / int(ch_ext.get("cx")), h / int(ch_ext.get("cy"))
            yield from walk(s.shapes, (l - int(ch_off.get("x")) * gx, t - int(ch_off.get("y")) * gy, gx, gy))
        else:
            yield s, l, t, w, h


def body_slides(prs):
    """(본문 쪽번호, 장) — 표지·목차·간지(도형 이름 pd_front…)는 세지 않는다. 화면의 시안·결과 번호와 같은 번호"""
    n = 0
    for slide in prs.slides:
        if any((s.name or "").startswith("pd_front") for s in slide.shapes):
            continue
        n += 1
        yield n, slide


def lines_of(tf):
    return "\n".join(p.text.replace("\v", "\n") for p in tf.paragraphs)


def build(prs):
    W, H = prs.slide_width, prs.slide_height
    pct = lambda v, k: round(v / k * 100, 2)
    out = []
    for no, slide in body_slides(prs):
        items = []
        for s, l, t, w, h in walk(slide.shapes):
            if getattr(s, "has_table", False) and s.has_table:
                tb, y = s.table, t
                tw, th = sum(c.width for c in tb.columns) or 1, sum(r.height for r in tb.rows) or 1
                for ri, row in enumerate(tb.rows):
                    x, rh = l, row.height / th * h
                    for ci, col in enumerate(tb.columns):
                        cw = col.width / tw * w
                        cell = tb.cell(ri, ci)
                        if not cell.is_spanned and cell.text.strip():
                            items.append({"id": f"s{s.shape_id}:{ri}:{ci}", "x": pct(x, W), "y": pct(y, H), "w": pct(cw, W), "h": pct(rh, H), "text": lines_of(cell.text_frame)})
                        x += cw
                    y += rh
                continue
            if getattr(s, "has_text_frame", False) and s.text_frame.text.strip() and w > 0 and h > 0:
                items.append({"id": f"s{s.shape_id}", "x": pct(l, W), "y": pct(t, H), "w": pct(w, W), "h": pct(h, H), "text": lines_of(s.text_frame),
                              "key": (s.name or "").lower().startswith("keymsg")})
        out.append({"n": no, "items": items})
    return out


def set_para(p, new):
    """문단 글을 바꾸되 서식을 지킨다. 한 런 안에서만 달라졌으면 그 런만 고친다(강조색 낱말 등 유지)"""
    old = p.text
    if old == new:
        return
    runs = list(p.runs)
    if runs and "\v" not in old and "\v" not in new and "".join(r.text for r in runs) == old:
        a = 0
        while a < min(len(old), len(new)) and old[a] == new[a]:
            a += 1
        b = 0
        while b < min(len(old), len(new)) - a and old[-1 - b] == new[-1 - b]:
            b += 1
        end, pos = len(old) - b, 0
        for r in runs:
            n = len(r.text)
            if pos <= a and end <= pos + n and (a < pos + n or r is runs[-1] or a == end):
                r.text = r.text[:a - pos] + new[a:len(new) - b] + r.text[end - pos:]
                return
            pos += n
    rpr = copy.deepcopy(runs[0]._r.find(A + "rPr")) if runs and runs[0]._r.find(A + "rPr") is not None else None
    p.text = new   # "\v" 는 줄 나눔(a:br)으로 들어간다
    if rpr is not None:
        for r in p.runs:
            cur = r._r.find(A + "rPr")
            if cur is not None:
                r._r.remove(cur)
            r._r.insert(0, copy.deepcopy(rpr))


def set_text(tf, text):
    new = [x.rstrip() for x in text.replace("\r", "").split("\n")]
    while len(new) > 1 and not new[-1]:
        new.pop()
    paras = list(tf.paragraphs)
    units = [len(p.text.split("\v")) for p in paras]   # 문단마다 줄(줄 나눔 포함) 수
    if sum(units) == len(new):   # 줄 수가 같으면 문단·줄 나눔 구조 그대로
        i = 0
        for p, k in zip(paras, units):
            set_para(p, "\v".join(new[i:i + k]))
            i += k
        return
    for i, p in enumerate(paras):   # 줄 수가 달라졌으면 줄 = 문단. 남는 줄은 마지막 문단에, 남는 문단은 지운다
        if i < len(new):
            set_para(p, new[i] if i < len(paras) - 1 else "\v".join(new[i:]))
        else:
            p._p.getparent().remove(p._p)


def apply(prs, edits):
    by = {}
    for e in edits:
        by.setdefault(int(e["slide"]), {})[str(e["id"])] = str(e.get("text", ""))
    done = 0
    for no, slide in body_slides(prs):
        want = by.get(no)
        if not want:
            continue
        for s, *_ in walk(slide.shapes):
            if getattr(s, "has_table", False) and s.has_table:
                for ri, row in enumerate(s.table.rows):
                    for ci in range(len(s.table.columns)):
                        k = f"s{s.shape_id}:{ri}:{ci}"
                        if k in want:
                            set_text(s.table.cell(ri, ci).text_frame, want[k])
                            done += 1
            elif f"s{s.shape_id}" in want and getattr(s, "has_text_frame", False):
                set_text(s.text_frame, want[f"s{s.shape_id}"])
                done += 1
    return done


if __name__ == "__main__":
    cmd, src = sys.argv[1], sys.argv[2]
    prs = Presentation(src)
    if cmd == "map":
        json.dump(build(prs), open(sys.argv[3], "w", encoding="utf-8"), ensure_ascii=False)
        print("글 상자", sum(len(s["items"]) for s in build(prs)))
    elif cmd == "apply":
        n = apply(prs, json.load(open(sys.argv[3], encoding="utf-8")))
        prs.save(sys.argv[4])
        print(f"고친 곳 {n}")
