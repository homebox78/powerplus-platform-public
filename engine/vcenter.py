# -*- coding: utf-8 -*-
"""도형 안 글자 세로 가운데 맞춤 (전 장) — PowerPoint 없는 판(리눅스 엔진).
ppt-pipeline/vcenter_text.py(PowerPoint COM)와 같은 일을 python-pptx + PowerPoint 글줄 수치표로 한다(렌더 없음).
1) 글자가 든 도형마다 '보이는 바탕 띠'를 찾는다 — 바탕 = 글 줄 가운데를 품는 채운 도형(자기 자신 포함).
   그 위(z 위)에 겹친 채운 도형이 띠를 가리면 가려진 곳까지만 띠로 본다. 띠 높이는 한두 줄 글 크기여야 한다.
2) 글 경계와 글자 잉크의 위·아래 끝은 렌더하지 않고 계산한다 — PowerPoint 가 서체마다 글줄을 놓는 수치
   (fonts/ppt_metrics.json, ppt_metrics.py 가 PowerPoint 로 잼) + 서체 파일의 글자 모양(PIL getbbox).
   LibreOffice 렌더는 도형 안 글줄 위치가 PowerPoint 와 달라(배지에서 3pt 넘게) 쓰지 않는다 — 2026-10-05 비교.
3) 잉크 가운데를 띠 가운데로: 글 상자가 따로면 상자를 옮기고, 자기 자신이 바탕이면 위·아래 안쪽 여백을 조정.
인자: <src> <dst> [--dry] [--slides 4,7] [--work 폴더(쓰지 않음 — 예전 인자 호환)]"""
import json, os, sys
from pptx import Presentation
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
args = [a for a in sys.argv[1:] if not a.startswith("--")]
opt = lambda k, d=None: sys.argv[sys.argv.index(k) + 1] if k in sys.argv else d
DRY = "--dry" in sys.argv
ONLY = [int(x) for x in opt("--slides").split(",")] if opt("--slides") else None
IN = 72.0
EMU = 12700.0   # 1pt
HERE = os.path.dirname(os.path.abspath(__file__))


def _alpha(fill_el):
    a = fill_el.find(".//" + qn("a:alpha"))
    return int(a.get("val")) / 100000 if a is not None else 1.0


def _ref_on(sp, tag):
    st = sp.find(qn("p:style"))
    ref = st.find(qn(tag)) if st is not None else None
    return ref is not None and ref.get("idx") not in (None, "0")


def _visible_fill(sp):
    """채움이 보이는가(불투명도 50% 넘게). 명시가 없으면 도형 스타일(fillRef)을 본다"""
    pr = sp.find(qn("p:spPr"))
    if pr is None:
        return False
    for tag in ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill"):
        f = pr.find(qn(tag))
        if f is not None:
            return tag != "a:solidFill" or _alpha(f) > 0.5
    if pr.find(qn("a:noFill")) is not None or pr.find(qn("a:grpFill")) is not None:
        return False
    return _ref_on(sp, "a:fillRef")


def _visible_line(sp):
    """선이 보이는가(0.5pt 이상, 불투명도 50% 넘게)"""
    pr = sp.find(qn("p:spPr"))
    ln = pr.find(qn("a:ln")) if pr is not None else None
    if ln is None:
        return _ref_on(sp, "a:lnRef")
    if ln.find(qn("a:noFill")) is not None:
        return False
    w = int(ln.get("w")) / EMU if ln.get("w") else (0.75 if _ref_on(sp, "a:lnRef") else 0)
    f = ln.find(qn("a:solidFill"))
    if f is not None:
        return w >= 0.5 and _alpha(f) > 0.5
    return w >= 0.5 and _ref_on(sp, "a:lnRef")


class Sh:
    """평평하게 편 도형 하나 — 장 위 절대 좌표(pt). (sx, sy) 는 그룹 배율(자식 좌표 1pt 가 장 위 몇 pt)"""
    def __init__(self, el, l, t, w, h, sx, sy, rot):
        self.el, self.Left, self.Top, self.Width, self.Height, self.sx, self.sy, self.rot = el, l, t, w, h, sx, sy, rot
        self.kind = el.tag.split("}")[1]
        nv = el.find(".//" + qn("p:cNvPr"))
        self.Name = nv.get("name") if nv is not None else ""


def _xfrm(el):
    pr = el.find(qn("p:spPr"))
    x = pr.find(qn("a:xfrm")) if pr is not None else None
    if x is None:
        return None
    off, ext = x.find(qn("a:off")), x.find(qn("a:ext"))
    if off is None or ext is None:
        return None
    return x, int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def flat(tree, map_=None, out=None):
    """spTree/grpSp 안 도형을 z 순서대로 평평하게 — map_ 은 자식 EMU → 장 pt 변환 (ax, bx, ay, by).
    돌린 그룹은 건너뛴다(COM 판이 그룹 안 그룹 오류를 건너뛰는 것과 같은 몫)"""
    map_ = map_ or (1 / EMU, 0.0, 1 / EMU, 0.0)
    out = [] if out is None else out
    ax, bx, ay, by = map_
    for el in tree:
        tag = el.tag.split("}")[1]
        if tag == "grpSp":
            gp = el.find(qn("p:grpSpPr"))
            x = gp.find(qn("a:xfrm")) if gp is not None else None
            if x is None or x.find(qn("a:chOff")) is None or int(x.get("rot", "0")) % 21600000:
                continue
            off, ext, cho, che = (x.find(qn(t)) for t in ("a:off", "a:ext", "a:chOff", "a:chExt"))
            kx = int(ext.get("cx")) / max(1, int(che.get("cx")))
            ky = int(ext.get("cy")) / max(1, int(che.get("cy")))
            # 자식 c → 부모 좌표 off + (c - chOff)·k → 장 pt
            flat(el, (ax * kx, bx + ax * (int(off.get("x")) - int(cho.get("x")) * kx),
                      ay * ky, by + ay * (int(off.get("y")) - int(cho.get("y")) * ky)), out)
        elif tag in ("sp", "cxnSp", "pic"):
            g = _xfrm(el)
            if g is None:
                continue
            x, ox, oy, cx, cy = g
            out.append(Sh(el, ax * ox + bx, ay * oy + by, ax * cx, ay * cy, ax * EMU, ay * EMU, int(x.get("rot", "0")) / 60000))
    return out


def filled(s, SW, SH):
    return s.kind == "sp" and (_visible_fill(s.el) or _visible_line(s.el)) and s.Width > 4 and s.Height > 4 \
        and not (s.Width > SW * 0.9 and s.Height > SH * 0.4)


# ── PowerPoint 글줄 모형 (TextRange.Bound* · 화면 실측 대신)
# 첫 기준선 = 글 위 + (asc + k·(배수-1))·pt, 줄 높이 = lh·pt·배수·lhk
import fontset
_MP = os.path.join(HERE, "fonts", "ppt_metrics.json")
METRICS = {k.lower(): v for k, v in (json.load(open(_MP, encoding="utf-8")) if os.path.isfile(_MP) else {}).items()}
_DEF = {"asc": 0.96, "lh": 1.2, "k": 0.78, "lhk": 1.0}
_F = {}


def _metric(face):
    return METRICS.get((face or "").lower(), _DEF)


def _font(face):
    """서체 파일(200px) — 설치된 것 → 수치표의 파일 이름으로 worker/fonts·리눅스 서체 폴더 → 본문 서체"""
    key = (face or "").lower()
    if key not in _F:
        import osdep
        from PIL import ImageFont
        cands = [fontset.INSTALLED.get(key)]
        fn = METRICS.get(key, {}).get("file")
        if fn:
            for d in (os.path.join(HERE, "fonts"), osdep.LX_FONT_DIR):
                cands += [os.path.join(d, fn), os.path.join(d, fn.replace("pd_", ""))]
        cands.append(fontset.file(1))
        _F[key] = None
        for p in cands:
            if p and os.path.isfile(p):
                try:
                    _F[key] = ImageFont.truetype(p, 200)
                    break
                except Exception:
                    continue
    return _F[key]


def _sz(*els):
    for el in els:
        if el is not None and el.get("sz"):
            return int(el.get("sz")) / 100
    return 18.0


def _face(*rprs):
    for rp in rprs:
        if rp is None:
            continue
        for tag in ("a:ea", "a:latin"):
            x = rp.find(qn(tag))
            if x is not None and x.get("typeface") and not x.get("typeface").startswith("+"):
                return x.get("typeface")
    return fontset.BODY


def _spc_pts(ppr, tag):
    x = ppr.find(qn(tag)) if ppr is not None else None
    p = x.find(qn("a:spcPts")) if x is not None else None
    return int(p.get("val")) / 100 if p is not None else 0.0


def _width(f, t, z):
    return (f.getlength(t) / 200 * z) if f else sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in t) * z


def layout(s):
    """PowerPoint 화면에서의 글 — dict(bl, bw, bt, bh, ink_top, ink_bot) 장 위 pt. 글이 없거나 세로 쓰기면 None"""
    tx = s.el.find(qn("p:txBody"))
    if tx is None:
        return None
    bp = tx.find(qn("a:bodyPr"))
    if bp is not None and bp.get("vert") not in (None, "horz"):
        return None
    g = lambda k, d: int(bp.get(k)) / EMU if bp is not None and bp.get(k) else d
    li, ri, ti, bi = g("lIns", 7.2), g("rIns", 7.2), g("tIns", 3.6), g("bIns", 3.6)
    wrap = bp is None or bp.get("wrap") != "none"
    na = bp.find(qn("a:normAutofit")) if bp is not None else None
    fscale = int(na.get("fontScale", "100000")) / 100000 if na is not None else 1.0
    lred = int(na.get("lnSpcReduction", "0")) / 100000 if na is not None else 0.0
    sx, sy = s.sx, s.sy
    box_w = s.Width / sx - li - ri          # 자식 좌표(pt) 기준 글 폭
    lst = tx.find(qn("a:lstStyle"))
    ldef = lst.find(".//" + qn("a:defRPr")) if lst is not None else None
    lines, widths, aligns = [], [], []      # lines: (기준선 y, 글, 서체, pt) — 글 위에서 잰 값
    y, first = 0.0, True
    for p in tx.findall(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        end = p.find(qn("a:endParaRPr"))
        runs, segs = [], [[]]   # segs: <a:br/> 로 나뉜 조각마다 run 들
        for r in p:
            if r.tag == qn("a:r"):
                runs.append((r.findtext(qn("a:t")) or "", r.find(qn("a:rPr"))))
                segs[-1].append(runs[-1])
            elif r.tag == qn("a:br"):
                segs.append([])
        txt = "".join(t for t, _ in runs)
        big = max(runs, key=lambda r: _sz(r[1], end, ldef)) if runs else ("", None)
        pt, face = _sz(big[1], end, ldef) * fscale, _face(big[1], end, ldef)
        m = _metric(face)
        spc = 1.0
        ls = ppr.find(qn("a:lnSpc")) if ppr is not None else None
        if ls is not None and ls.find(qn("a:spcPct")) is not None:
            spc = int(ls.find(qn("a:spcPct")).get("val")) / 100000
        spc = max(0.1, spc - lred)
        n, w, segtxt = 0, 0.0, []
        for sg in segs:
            sw = sum(_width(_font(_face(rp, end, ldef)), t, _sz(rp, end, ldef) * fscale) for t, rp in sg)
            st = "".join(t for t, _ in sg)
            k = max(1, -(-int(sw) // max(1, int(box_w)))) if wrap and box_w > 0 and st.strip() else 1
            n += k
            w = max(w, sw)
            segtxt += [st] * k
        lh = m["lh"] * pt * spc * m.get("lhk", 1.0)
        if not first:
            y += _spc_pts(ppr, "a:spcBef")
        base0 = y + (m["asc"] + m["k"] * (spc - 1)) * pt if first else y + lh - (m["lh"] - m["asc"]) * pt
        for i in range(n):
            lines.append((base0 + i * lh, segtxt[i], face, pt))
        y += lh * n + _spc_pts(ppr, "a:spcAft")
        first = False
        widths.append(min(w, box_w) if wrap else w)
        aligns.append((ppr.get("algn") if ppr is not None else None) or "l")
    if not any(x[1].strip() for x in lines):
        return None
    total = y
    anchor = (bp.get("anchor") if bp is not None else None) or "t"
    top0, bot0 = ti, s.Height / sy - bi
    t0 = (top0 + bot0 - total) / 2 if anchor == "ctr" else (bot0 - total if anchor == "b" else top0)

    def ink(line, which):   # which 1 = 잉크 위 끝, 3 = 잉크 아래 끝 (기준선 기준 getbbox)
        base, txt, face, pt = line
        f = _font(face)
        if f is None:
            return base + (-0.78 if which == 1 else 0.12) * pt
        return base + f.getbbox(txt.strip() or "가", anchor="ls")[which] / 200 * pt
    live = [x for x in lines if x[1].strip()]
    bw = max(widths)
    al = aligns[widths.index(bw)]
    bl = li + (box_w - bw) / 2 if al == "ctr" else (li + box_w - bw if al == "r" else li)
    return dict(bl=s.Left + bl * sx, bw=bw * sx, bt=s.Top + t0 * sy, bh=total * sy,
                ink_top=s.Top + (t0 + ink(live[0], 1)) * sy, ink_bot=s.Top + (t0 + ink(live[-1], 3)) * sy)


def band_of(shs, idx, t, bl, bw, bt, bh, SW, SH):
    tc = bt + bh / 2
    best = None
    for j, f in enumerate(shs):
        if not filled(f, SW, SH) or j > idx:
            continue
        if not (f.Left - 1 <= bl + 1 and bl + bw - 1 <= f.Left + f.Width + 1 and f.Top <= tc <= f.Top + f.Height):
            continue
        top, bot = f.Top, f.Top + f.Height
        for g in shs[j + 1:]:
            if g is t or not filled(g, SW, SH):
                continue
            if g.Left > bl + bw or g.Left + g.Width < bl:
                continue
            if g.Left > bl + 2 or g.Left + g.Width < bl + bw - 2:   # 글 폭을 다 덮는 판만
                continue
            if tc < g.Top < bot:
                bot = g.Top
            if top < g.Top + g.Height < tc:
                top = g.Top + g.Height
        h = bot - top
        if h < bh - 2 or h > max(bh * 2.4, 0.5 * IN):
            continue
        if best is None or h < best[1] - best[0]:
            best = (top, bot, f)
    return best


def _has_text(g):
    return g.kind == "sp" and any((x.text or "").strip() for x in g.el.iter(qn("a:t")))


def crowded(shs, t, band):
    """띠 안에 이 글 말고 다른 글·그림이 함께 있다 — 카드처럼 제목·아이콘이 든 판은 띠가 아니다.
    (COM 판은 이런 판에서 다른 잉크까지 재어 본문을 8~12pt 옮겼다 — 2026-10-05 비교)"""
    top, bot, f = band
    j = next(i for i, x in enumerate(shs) if x is f)
    for g in shs[j + 1:]:
        if g is t or not (g.kind == "pic" or _has_text(g)):
            continue
        if g.Left >= f.Left + f.Width or g.Left + g.Width <= f.Left:
            continue
        ov = min(bot, g.Top + g.Height) - max(top, g.Top)
        if ov > 0.3 * min(g.Height, bot - top):
            return True
    return False


def main():
    src, dst = os.path.abspath(args[0]), os.path.abspath(args[1])
    pres = Presentation(src)
    SW, SH = pres.slide_width / EMU, pres.slide_height / EMU
    cands = []
    for sno, sl in enumerate(pres.slides, 1):
        if ONLY and sno not in ONLY:
            continue
        shs = flat(sl.shapes._spTree)
        for i, t in enumerate(shs):
            if t.kind != "sp" or (abs(t.rot) > 1 and abs(t.rot - 360) > 1):
                continue
            L = layout(t)
            if not L or L["bh"] <= 0 or L["bh"] > 0.9 * IN:
                continue
            band = band_of(shs, i, t, L["bl"], L["bw"], L["bt"], L["bh"], SW, SH)
            if not band or crowded(shs, t, band):
                continue
            txt = "".join(x.text or "" for x in t.el.iter(qn("a:t")))
            k2 = 2 / 3200 * SW   # COM 판: 3200px 렌더에서 띠 위·아래 2px 를 빼고 잰다
            it, ib = max(L["ink_top"], band[0] + k2), min(L["ink_bot"], band[1] - k2)
            if ib <= it:
                continue
            cands.append(dict(sno=sno, el=t.el, sh=t, self=band[2] is t, text=txt[:18],
                              d=round((band[0] + band[1]) / 2 - (it + ib) / 2, 2)))
    fixes = [c for c in cands if 0.6 <= abs(c["d"]) <= 0.2 * IN]
    for c in cands:
        if abs(c["d"]) >= 0.6:
            print(c["sno"], c["sh"].Name, c["d"], "자기" if c["self"] else "상자", c["text"], "" if c in fixes else "(건너뜀)")
    if not DRY:
        for c in fixes:
            fix(c)
        pres.save(dst)
    print("후보", len(cands), "고침" if not DRY else "고칠 것", len(fixes))


def fix(c):
    el, sh, d = c["el"], c["sh"], c["d"]
    if not c["self"]:   # 글 상자를 d 만큼 아래로(그룹 안이면 그룹 배율로 나눠 자식 좌표로)
        off = _xfrm(el)[0].find(qn("a:off"))
        off.set("y", str(int(round(int(off.get("y")) + d / sh.sy * EMU))))
        return
    bp = el.find(qn("p:txBody")).find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):   # 자동 맞춤 끔(AutoSize=0)
        x = bp.find(qn(tag))
        if x is not None:
            bp.remove(x)
    from lxml import etree
    na = etree.Element(qn("a:noAutofit"))   # 순서: prstTxWarp 다음, scene3d·sp3d·flatTx·extLst 앞
    nxt = next((x for x in bp if x.tag in (qn("a:scene3d"), qn("a:sp3d"), qn("a:flatTx"), qn("a:extLst"))), None)
    bp.insert(list(bp).index(nxt), na) if nxt is not None else bp.append(na)
    g = lambda k, dv: int(bp.get(k)) / EMU if bp.get(k) else dv
    mt, mb = g("tIns", 3.6), g("bIns", 3.6)
    dd = d / sh.sy
    a = bp.get("anchor") or "t"
    if a == "t":
        mt = max(0, mt + dd)
    elif a == "b":
        mb = max(0, mb - dd)
    else:
        mt, mb = mt + dd, mb - dd
        if mt < 0: mb -= mt; mt = 0
        if mb < 0: mt -= mb; mb = 0
    bp.set("tIns", str(int(round(mt * EMU))))
    bp.set("bIns", str(int(round(mb * EMU))))


if __name__ == "__main__":
    main()
